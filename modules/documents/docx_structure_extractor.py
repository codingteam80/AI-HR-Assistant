"""Structure-aware DOCX extraction for company knowledge search.

The company reference documents do not consistently use Word Heading styles.
This extractor therefore preserves the original paragraph/table order and uses
conservative semantic/visual cues only to label reliable section boundaries.
When no reliable boundaries exist, the whole file remains one searchable
section so downstream bounded chunking still operates per file.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re

from docx import Document
from docx.document import Document as _Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass(slots=True)
class DOCXStructuredSection:
    """One ordered semantic section from a DOCX file."""

    heading: str
    text: str


@dataclass(slots=True)
class DOCXStructuredDocument:
    """Searchable DOCX text plus conservative semantic sections."""

    full_text: str
    sections: list[DOCXStructuredSection]
    reliable_heading_count: int


@dataclass(slots=True)
class _HeadingSignal:
    label: str
    inline_body: bool = False


class DOCXStructureExtractor:
    """Extract DOCX blocks without depending on Word Heading 1/2/3 styles."""

    _ARTICLE_RE = re.compile(
        r"^(?:addendum\s*[-:]\s*)?(?:article|chapter|part)\s+[ivxlcdm\w.-]+\b.*$",
        re.IGNORECASE,
    )
    _SECTION_RE = re.compile(
        r"^(section\.?\s+[\w.-]+)\s*[.:\-]?\s*(.*)$",
        re.IGNORECASE,
    )
    _DECIMAL_RE = re.compile(r"^(\d+(?:\.\d+)+)\s*[.)\-]?\s*(.*)$")
    _SINGLE_NUMBER_RE = re.compile(r"^\d+[.)]\s+")
    _LETTER_LIST_RE = re.compile(r"^\(?[a-z]\)[.)]?\s+", re.IGNORECASE)
    _GENERIC_LABELS = {
        "policy statement",
        "objective",
        "objectives",
        "criteria",
        "eligibility",
        "example",
        "illustration",
        "flow chart",
        "responsibilities",
    }

    @staticmethod
    def _clean(value: str) -> str:
        value = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()

    @staticmethod
    def _word_count(value: str) -> int:
        return len(re.findall(r"\S+", value or ""))

    @classmethod
    def _is_title_like(cls, text: str) -> bool:
        stripped = cls._clean(text)
        if not stripped or len(stripped) > 140:
            return False
        words = cls._word_count(stripped)
        if words > 16:
            return False
        if stripped.endswith(('.', '!', '?')) and not stripped.endswith(':'):
            return False
        return True

    @staticmethod
    def _bold_ratio(paragraph: Paragraph) -> float:
        total = 0
        bold = 0
        for run in paragraph.runs:
            text = run.text or ""
            meaningful = sum(1 for character in text if not character.isspace())
            if not meaningful:
                continue
            total += meaningful
            if run.bold is True:
                bold += meaningful
        return (bold / total) if total else 0.0

    @staticmethod
    def _upper_ratio(text: str) -> float:
        letters = [character for character in text if character.isalpha()]
        if not letters:
            return 0.0
        return sum(character.isupper() for character in letters) / len(letters)

    @classmethod
    def _heading_signal(cls, paragraph: Paragraph) -> _HeadingSignal | None:
        text = cls._clean(paragraph.text)
        if not text:
            return None

        style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
        if style_name.casefold().startswith("heading"):
            return _HeadingSignal(text[:250], inline_body=False)

        normalized = text.casefold().rstrip(":").strip()
        if normalized in cls._GENERIC_LABELS:
            return _HeadingSignal(text.rstrip(":")[:250], inline_body=False)

        if cls._ARTICLE_RE.match(text):
            return _HeadingSignal(text[:250], inline_body=False)

        section_match = cls._SECTION_RE.match(text)
        if section_match:
            prefix = section_match.group(1).strip().rstrip(".:-")
            remainder = section_match.group(2).strip()
            # A long remainder is the actual rule itself, not merely a title.
            inline_body = cls._word_count(remainder) > 10 or bool(
                re.search(r"[.!?]", remainder)
            )
            label = prefix if inline_body else text
            return _HeadingSignal(label[:250], inline_body=inline_body)

        decimal_match = cls._DECIMAL_RE.match(text)
        if decimal_match and cls._bold_ratio(paragraph) >= 0.60:
            prefix = decimal_match.group(1)
            remainder = decimal_match.group(2).strip()
            if not remainder:
                return _HeadingSignal(prefix, inline_body=False)
            inline_body = cls._word_count(remainder) > 12 or bool(
                re.search(r"[.!?]", remainder)
            )
            if inline_body:
                return _HeadingSignal(prefix, inline_body=True)
            return _HeadingSignal(text[:250], inline_body=False)

        # Avoid treating ordinary numbered/lettered list items as headings.
        if cls._SINGLE_NUMBER_RE.match(text) or cls._LETTER_LIST_RE.match(text):
            return None

        if cls._upper_ratio(text) >= 0.88 and cls._is_title_like(text):
            return _HeadingSignal(text.rstrip(":")[:250], inline_body=False)

        if cls._bold_ratio(paragraph) >= 0.80 and cls._is_title_like(text):
            return _HeadingSignal(text.rstrip(":")[:250], inline_body=False)

        if text.endswith(":") and cls._word_count(text) <= 10:
            return _HeadingSignal(text.rstrip(":")[:250], inline_body=False)

        return None

    @staticmethod
    def _iter_blocks(document: _Document):
        """Yield paragraphs and tables in their original XML order."""

        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                yield Paragraph(child, document)
            elif isinstance(child, CT_Tbl):
                yield Table(child, document)

    @classmethod
    def extract(
        cls,
        data: bytes,
        *,
        default_heading: str = "Document",
    ) -> DOCXStructuredDocument:
        """Return loss-resistant ordered text and conservative sections."""

        document = Document(BytesIO(data))
        full_blocks: list[str] = []
        sections: list[DOCXStructuredSection] = []
        current_heading = cls._clean(default_heading) or "Document"
        pending_headings: list[str] = []
        body_blocks: list[str] = []
        reliable_heading_count = 0
        table_number = 0

        def add_pending_to_body() -> None:
            nonlocal pending_headings
            if pending_headings:
                body_blocks.extend(pending_headings)
                pending_headings = []

        def flush() -> None:
            nonlocal body_blocks
            add_pending_to_body()
            text = cls._clean("\n\n".join(body_blocks))
            if text:
                sections.append(
                    DOCXStructuredSection(
                        heading=current_heading[:250],
                        text=text,
                    )
                )
            body_blocks = []

        for block in cls._iter_blocks(document):
            if isinstance(block, Paragraph):
                text = cls._clean(block.text)
                if not text:
                    continue
                full_blocks.append(text)
                signal = cls._heading_signal(block)
                if signal is None:
                    add_pending_to_body()
                    body_blocks.append(text)
                    continue

                reliable_heading_count += 1
                if body_blocks:
                    flush()
                # Keep every consecutive heading-like line until real body
                # content arrives. The reference corpus contains stacked title,
                # addendum, objective, and article labels; dropping older labels
                # would silently remove searchable source text.
                pending_headings.append(text)
                current_heading = signal.label or current_heading
                if signal.inline_body:
                    add_pending_to_body()
            else:
                table_number += 1
                table_lines = [f"Table {table_number}"]
                for row in block.rows:
                    values = [cls._clean(cell.text) for cell in row.cells]
                    if any(values):
                        table_lines.append(" | ".join(values))
                if len(table_lines) == 1:
                    continue
                table_text = "\n".join(table_lines)
                full_blocks.append(table_text)
                add_pending_to_body()
                body_blocks.append(table_text)

        flush()
        full_text = cls._clean("\n\n".join(full_blocks))
        if not full_text:
            return DOCXStructuredDocument("", [], reliable_heading_count)

        # If semantic/visual cues were not reliable enough to create useful
        # boundaries, explicitly preserve one file-level section. Downstream
        # bounded chunking then splits the entire file with overlap.
        if reliable_heading_count == 0 or not sections:
            sections = [
                DOCXStructuredSection(
                    heading=cls._clean(default_heading) or "Document",
                    text=full_text,
                )
            ]

        return DOCXStructuredDocument(
            full_text=full_text,
            sections=sections,
            reliable_heading_count=reliable_heading_count,
        )
