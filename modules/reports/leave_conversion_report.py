"""Leave conversion-to-cash report rows and Excel export."""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = (
    "Employee No./ID",
    "Employee Name",
    "Leave Year",
    "Available Vacation Leave (Last Year)",
    "Credit Vacation Leave (Current Year)",
    "Remaining Sick Leave (Last Year)",
    "VL Converted to Cash",
    "SL Converted to Cash",
    "Total Conversion",
    "Total Conversion Amount",
)
MISSING_NAME_TOKENS = {"n/a", "na", "none", "null", "-", "—"}


def _name_part(value: object) -> str:
    text = str(value or "").strip()
    if not text or text.casefold() in MISSING_NAME_TOKENS:
        return ""
    return text


def report_employee_name(employee) -> str:
    """Return Last Name, First Name Suffix, M. for payroll reports."""

    last_name = _name_part(getattr(employee, "last_name", None))
    first_name = _name_part(getattr(employee, "first_name", None))
    suffix = _name_part(getattr(employee, "suffix", None))
    middle_name = _name_part(getattr(employee, "middle_name", None))
    if not last_name and not first_name:
        return _name_part(getattr(employee, "full_name", None))
    given_name = " ".join(part for part in (first_name, suffix) if part)
    name = ", ".join(part for part in (last_name, given_name) if part)
    if middle_name:
        name = f"{name}, {middle_name[0].upper()}."
    return name


def _days(value: object) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def build_leave_conversion_rows(
    *,
    employees,
    balances,
    leave_year: int,
) -> list[dict[str, object]]:
    """Build one conversion summary row per employee for a selected year."""

    balance_map: dict[tuple[int, str], object] = {}
    for balance in balances:
        code = str(getattr(balance.leave_type, "code", "") or "").strip().upper()
        balance_map[(balance.employee_id, code)] = balance

    rows: list[dict[str, object]] = []
    for employee in employees:
        vacation = balance_map.get((employee.id, "VACATION"))
        sick = balance_map.get((employee.id, "SICK"))
        vl_converted = _days(
            getattr(vacation, "converted_to_cash_days", 0)
        )
        sl_converted = _days(
            getattr(sick, "converted_to_cash_days", 0)
        )
        rows.append(
            {
                "Employee No./ID": employee.employee_number,
                "Employee Name": report_employee_name(employee),
                "Leave Year": int(leave_year),
                "Available Vacation Leave (Last Year)": _days(
                    getattr(vacation, "beginning_credit_days", 0)
                ),
                "Credit Vacation Leave (Current Year)": _days(
                    getattr(vacation, "credit_days", 0)
                ),
                "Remaining Sick Leave (Last Year)": _days(
                    getattr(sick, "beginning_credit_days", 0)
                ),
                "VL Converted to Cash": vl_converted,
                "SL Converted to Cash": sl_converted,
                "Total Conversion": _days(vl_converted + sl_converted),
                # Reserved for a later payroll-rate implementation.
                "Total Conversion Amount": "",
            }
        )
    return rows


def build_leave_conversion_excel(
    rows: list[dict[str, object]],
) -> bytes:
    """Export conversion rows with a payroll-friendly fixed column layout."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leave Conversion to Cash"
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append([row.get(header, "") for header in HEADERS])

    border = Border(
        left=Side(style="thin", color="BFBFBF"),
        right=Side(style="thin", color="BFBFBF"),
        top=Side(style="thin", color="BFBFBF"),
        bottom=Side(style="thin", color="BFBFBF"),
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="10172A")
        cell.fill = PatternFill("solid", fgColor="FFF59D")
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = border
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for index in range(4, 10):
            row[index - 1].number_format = "0.00"

    widths = (20, 34, 14, 30, 30, 28, 22, 22, 20, 26)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 42
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
