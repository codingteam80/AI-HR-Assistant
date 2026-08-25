"""Excel/PDF exports for deterministic Chat Assistant report artifacts.

Both export formats consume the exact report dictionary rendered in chat so
summary, chart, detail rows, period, and permissions cannot drift between UI
and downloaded files.
"""

from __future__ import annotations

from io import BytesIO
from math import ceil
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.graphics.charts.piecharts import Pie
from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


_HEADER_FILL = "E9EDF5"


def _safe_sheet_title(value: str) -> str:
    cleaned = "".join("_" if char in "[]:*?/\\" else char for char in value)
    return (cleaned or "Report")[:31]


def _fit_widths(sheet, rows: list[dict[str, Any]], columns: list[str]) -> None:
    for index, column in enumerate(columns, start=1):
        sample = [str(column)] + [str(row.get(column, "")) for row in rows[:250]]
        width = min(48, max(12, max(len(value) for value in sample) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width


def _add_excel_chart(workbook: Workbook, report: dict[str, Any]) -> None:
    chart_spec = report.get("chart") or {}
    data = list(chart_spec.get("data") or [])
    if not data:
        return

    sheet = workbook.create_sheet("Chart Data")
    sheet.append((chart_spec.get("category_label", "Category"), chart_spec.get("value_label", "Value")))
    for item in data:
        sheet.append((str(item.get("label", "")), float(item.get("value", 0) or 0)))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)

    values = Reference(sheet, min_col=2, min_row=1, max_row=sheet.max_row)
    categories_ref = Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
    chart_type = str(chart_spec.get("type", "bar")).casefold()
    if chart_type == "line":
        chart = LineChart()
        chart.y_axis.title = str(chart_spec.get("value_label", "Value"))
        chart.x_axis.title = str(chart_spec.get("category_label", "Category"))
    elif chart_type == "pie":
        chart = PieChart()
    elif chart_type == "donut":
        chart = DoughnutChart()
        chart.holeSize = 55
    else:
        chart = BarChart()
        chart.y_axis.title = str(chart_spec.get("value_label", "Value"))
        chart.x_axis.title = str(chart_spec.get("category_label", "Category"))
    chart.title = str(chart_spec.get("title", report.get("title", "Report Chart")))
    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories_ref)
    chart.height = 9
    chart.width = 16
    sheet.add_chart(chart, "D2")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 16


def build_chat_report_excel(report: dict[str, Any]) -> bytes:
    """Build an Excel report from the exact chat report artifact."""

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary["A1"] = str(report.get("title", "Chat Assistant Report"))
    summary["A1"].font = Font(size=16, bold=True)
    summary["A3"] = "Period"
    summary["B3"] = str(report.get("period", ""))
    summary["A4"] = "Generated"
    summary["B4"] = str(report.get("generated_at", ""))
    summary["A5"] = "Timezone"
    summary["B5"] = str(report.get("timezone", ""))
    summary["A7"] = "Summary"
    summary["A7"].font = Font(bold=True)
    for row_index, line in enumerate(list(report.get("summary") or []), start=8):
        summary.cell(row=row_index, column=1, value=f"• {line}")
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 32

    rows = list(report.get("rows") or [])
    columns = list(report.get("columns") or [])
    detail = workbook.create_sheet(_safe_sheet_title("Report Data"))
    if columns:
        detail.append(columns)
        for cell in detail[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in rows:
            detail.append([row.get(column, "") for column in columns])
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions
        _fit_widths(detail, rows, columns)
        for row in detail.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    else:
        detail["A1"] = "No matching detail rows."

    _add_excel_chart(workbook, report)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _pdf_chart(report: dict[str, Any]) -> Drawing | None:
    spec = report.get("chart") or {}
    data = list(spec.get("data") or [])
    if not data:
        return None
    labels = [str(item.get("label", "")) for item in data]
    values = [float(item.get("value", 0) or 0) for item in data]
    chart_type = str(spec.get("type", "bar")).casefold()
    drawing = Drawing(700, 250)
    drawing.add(String(10, 232, str(spec.get("title", "Chart")), fontName="Helvetica-Bold", fontSize=11, fillColor=colors.HexColor("#10172A")))

    if chart_type in {"pie", "donut"}:
        chart = Pie()
        chart.x = 245
        chart.y = 25
        chart.width = 210
        chart.height = 190
        chart.data = values
        chart.labels = [f"{label} ({value:g})" for label, value in zip(labels, values)]
        chart.slices.strokeWidth = 0.5
        drawing.add(chart)
        if chart_type == "donut":
            drawing.add(Circle(350, 120, 42, fillColor=colors.white, strokeColor=colors.white))
        return drawing

    max_value = max(values) if values else 0.0
    max_value = max(max_value, 1.0)
    left, bottom, width, height = 60, 52, 600, 150
    drawing.add(Line(left, bottom, left, bottom + height, strokeColor=colors.HexColor("#64748B")))
    drawing.add(Line(left, bottom, left + width, bottom, strokeColor=colors.HexColor("#64748B")))

    if chart_type == "line":
        count = max(1, len(values) - 1)
        points: list[float] = []
        for index, value in enumerate(values):
            x = left + (width * index / count if len(values) > 1 else width / 2)
            y = bottom + (height * value / max_value)
            points.extend((x, y))
            drawing.add(Circle(x, y, 3, fillColor=colors.HexColor("#4F46E5"), strokeColor=None))
            drawing.add(String(x - 12, y + 7, f"{value:g}", fontSize=7, fillColor=colors.HexColor("#10172A")))
            drawing.add(String(x - 22, bottom - 17, labels[index][:12], fontSize=6.5, fillColor=colors.HexColor("#334155")))
        if len(points) >= 4:
            drawing.add(PolyLine(points, strokeColor=colors.HexColor("#4F46E5"), strokeWidth=2))
    else:
        count = max(1, len(values))
        slot = width / count
        bar_width = min(48, slot * 0.62)
        for index, value in enumerate(values):
            bar_height = height * value / max_value
            x = left + index * slot + (slot - bar_width) / 2
            drawing.add(Rect(x, bottom, bar_width, bar_height, fillColor=colors.HexColor("#4F46E5"), strokeColor=None))
            drawing.add(String(x + 2, bottom + bar_height + 5, f"{value:g}", fontSize=7, fillColor=colors.HexColor("#10172A")))
            drawing.add(String(x - 2, bottom - 17, labels[index][:11], fontSize=6.5, fillColor=colors.HexColor("#334155")))
    return drawing


def build_chat_report_pdf(report: dict[str, Any]) -> bytes:
    """Build a PDF with the same summary, chart data, and detail dataset."""

    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=24,
        rightMargin=24,
        topMargin=24,
        bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story: list[Any] = [
        Paragraph(str(report.get("title", "Chat Assistant Report")), styles["Title"]),
        Paragraph(f"Period: {report.get('period', '')}", styles["Normal"]),
        Paragraph(f"Generated: {report.get('generated_at', '')} ({report.get('timezone', '')})", styles["Normal"]),
        Spacer(1, 8),
    ]
    summary_lines = list(report.get("summary") or [])
    if summary_lines:
        story.append(Paragraph("Summary", styles["Heading3"]))
        for line in summary_lines:
            story.append(Paragraph(f"• {line}", styles["Normal"]))
        story.append(Spacer(1, 8))

    chart = _pdf_chart(report)
    if chart is not None:
        story.append(KeepTogether([chart, Spacer(1, 8)]))

    columns = list(report.get("columns") or [])
    rows = list(report.get("rows") or [])
    if columns:
        table_data: list[list[Any]] = [[Paragraph(str(column), styles["BodyText"]) for column in columns]]
        for row in rows:
            table_data.append([
                Paragraph(str(row.get(column, "")), styles["BodyText"])
                for column in columns
            ])
        available_width = landscape(A4)[0] - 48
        column_width = available_width / max(1, len(columns))
        table = Table(table_data, repeatRows=1, colWidths=[column_width] * len(columns))
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EDF5")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.extend([Paragraph("Detail data", styles["Heading3"]), table])
    else:
        story.append(Paragraph("No matching detail rows.", styles["Normal"]))

    document.build(story)
    return output.getvalue()
