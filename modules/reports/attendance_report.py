"""Excel and PDF attendance report generation."""

from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from models.attendance_record import AttendanceRecord
from services.attendance_service import ATTENDANCE_COLORS


REPORT_HEADERS = (
    "Employee Name", "Date", "Sessions", "Status", "Work Hours",
    "Leave Hours", "Undertime Hours", "OT Hours",
)


def _local_time(value: datetime | None, tz: ZoneInfo) -> str:
    if value is None:
        return "—"
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(tz).strftime("%H:%M")


def _row(record: AttendanceRecord, tz: ZoneInfo) -> list[object]:
    sessions = []
    for item in list(getattr(record, "sessions", ()) or ()):
        sessions.append(
            f"{item.work_status} {_local_time(item.rounded_time_in, tz)}-"
            f"{_local_time(item.rounded_time_out, tz)}"
        )
    status = record.work_status or "—"
    if record.leave_duration_code in {"90501", "90502"}:
        half = "AM" if record.leave_duration_code == "90501" else "PM"
        status = f"{status} + {half} Leave"
    return [
        record.employee.full_name,
        record.attendance_date.strftime("%Y/%m/%d"),
        " / ".join(sessions) or "—",
        status,
        float(record.total_hours),
        float(record.leave_hours),
        float(record.undertime_hours),
        float(record.ot_hours),
    ]


def build_attendance_excel(
    records: list[AttendanceRecord], *, timezone_name: str
) -> bytes:
    """Return a styled workbook with detail and employee totals sheets."""

    tz = ZoneInfo(timezone_name)
    workbook = Workbook()
    detail = workbook.active
    detail.title = "Attendance Detail"
    detail.append(REPORT_HEADERS)
    for cell in detail[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9EDF5")
    totals: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for record in records:
        detail.append(_row(record, tz))
        color = ATTENDANCE_COLORS.get(record.work_status or "")
        if color:
            for cell in detail[detail.max_row]:
                cell.fill = PatternFill("solid", fgColor=color.removeprefix("#"))
        totals[record.employee.full_name][0] += float(record.total_hours)
        totals[record.employee.full_name][1] += float(record.ot_hours)
    for width, letter in zip((28, 14, 42, 18, 14, 14, 18, 12), "ABCDEFGH"):
        detail.column_dimensions[letter].width = width
    detail.freeze_panes = "A2"

    summary = workbook.create_sheet("Employee Totals")
    summary.append(("Employee Name", "Total Hours", "OT Hours"))
    for cell in summary[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E9EDF5")
    for employee_name, values in sorted(totals.items()):
        summary.append((employee_name, round(values[0], 2), round(values[1], 2)))
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 16
    summary.column_dimensions["C"].width = 16
    for sheet in (detail, summary):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_attendance_pdf(
    records: list[AttendanceRecord], *, timezone_name: str,
    start_date: date, end_date: date,
) -> bytes:
    """Return a compact landscape PDF detail report."""

    output = BytesIO()
    document = SimpleDocTemplate(
        output, pagesize=landscape(A4),
        leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Attendance / DTR / OT Report", styles["Title"]),
        Paragraph(
            f"Period: {start_date:%Y/%m/%d} to {end_date:%Y/%m/%d}",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]
    tz = ZoneInfo(timezone_name)
    table_data = [list(REPORT_HEADERS)] + [_row(record, tz) for record in records]
    table = Table(
        table_data,
        repeatRows=1,
        colWidths=[112, 66, 155, 72, 54, 54, 67, 52],
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E9EDF5")),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#10172A")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(table)
    document.build(story)
    return output.getvalue()
