"""Admin Shifting Credits workbook matching the approved report structure."""

from __future__ import annotations

import calendar
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from modules.reports.leave_conversion_report import report_employee_name
from services.overtime_service import OvertimeService


GRID = Side(style="thin", color="BFBFBF")
BORDER = Border(left=GRID, right=GRID, top=GRID, bottom=GRID)
TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
SUBHEADER_FILL = PatternFill("solid", fgColor="EAF2F8")
SUMMARY_FILL = PatternFill("solid", fgColor="E2F0D9")
DATE_FORMAT = "yyyy-mm-dd"


def _decimal(value: object) -> Decimal:
    return Decimal(str(value or 0)).quantize(Decimal("0.01"))


def _local_time(value: datetime | None, timezone_value: ZoneInfo) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone_value).strftime("%H:%M")


def _sheet_name(raw: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]", " ", raw or "Employee").strip() or "Employee"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        tail = f" ({suffix})"
        candidate = f"{base[: 31 - len(tail)]}{tail}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _apply_grid(sheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in sheet.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_widths(sheet, widths: dict[int, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def _credit_status_note(credit) -> str:
    if credit is None:
        return ""
    labels = {
        "pending": "PENDING / earned 1-day OB",
        "available": "AVAILABLE / 1-day OB ready for use",
        "reserved": "RESERVED / linked to pending OB leave request",
        "used": "USED / whole-day OB assigned to approved leave",
        "expired": "EXPIRED / source OT restored to Regular OT",
    }
    return labels.get(str(credit.status), str(credit.status).upper())


def build_shifting_credits_excel(
    *,
    employees,
    overtime_requests,
    credits,
    report_year: int,
    cutoff_day: int,
    availability_cutoffs: int,
    expiration_mode: str,
    expiration_month: int,
    expiration_day: int,
    block_hours: Decimal,
    required_blocks: int,
    straight_ot_threshold_hours: Decimal,
    straight_ot_vl_days: Decimal,
    straight_ot_also_payable: bool,
    excluded_positions: tuple[str, ...],
    timezone_name: str,
) -> bytes:
    """Build Summary, Cut-Off Date, and one sheet for each eligible employee."""

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    timezone_value = ZoneInfo(timezone_name)
    employee_map = {item.id: item for item in employees}
    requests_by_employee: dict[int, list] = {item.id: [] for item in employees}
    for request in overtime_requests:
        if request.employee_id in requests_by_employee and request.date_rendered.year == report_year:
            requests_by_employee[request.employee_id].append(request)

    year_credits = [
        credit for credit in credits
        if credit.employee_id in employee_map and credit.cutoff_start.year == report_year
    ]
    credits_by_employee: dict[int, list] = {item.id: [] for item in employees}
    credit_by_group = {}
    for credit in year_credits:
        credits_by_employee[credit.employee_id].append(credit)
        credit_by_group[credit.group_key] = credit

    summary.merge_cells("A1:H1")
    summary["A1"] = f"SHIFTING CREDITS REPORT — {report_year}"
    summary["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    summary["A1"].fill = TITLE_FILL
    summary["A1"].alignment = Alignment(horizontal="center")
    summary.merge_cells("A2:H2")
    summary["A2"] = (
        "Live eligible employees and approved OT records. Same-cutoff Shifting OB and "
        "straight-OT VL conversion are reported separately."
    )
    summary["A2"].alignment = Alignment(wrap_text=True)
    headers = [
        "Employee",
        "Position",
        "Eligible",
        "Shifting OB Earned (Days)",
        "Available OB (Days)",
        "Regular OT Excess (Hrs)",
        "VL Credit",
        "Notes",
    ]
    for col, value in enumerate(headers, 1):
        cell = summary.cell(4, col, value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, employee in enumerate(employees, 5):
        employee_credits = credits_by_employee.get(employee.id, [])
        employee_requests = requests_by_employee.get(employee.id, [])
        earned = sum((_decimal(item.ob_credit_days) for item in employee_credits), Decimal("0"))
        available = sum(
            (_decimal(item.ob_credit_days) for item in employee_credits if item.status == "available"),
            Decimal("0"),
        )
        regular_excess = Decimal("0.00")
        vl_credit = Decimal("0.00")
        for request in employee_requests:
            vl_credit += _decimal(request.additional_vl_days)
            if request.shifting_credit_group:
                regular_excess += max(
                    Decimal("0.00"),
                    _decimal(request.estimated_hours) - _decimal(request.shifting_credit_hours),
                ) + _decimal(request.shifting_credit_restored_hours)
        statuses = {str(item.status) for item in employee_credits}
        notes = ", ".join(
            label for status, label in (
                ("pending", "Pending availability"),
                ("available", "Available OB"),
                ("reserved", "Reserved OB"),
                ("used", "Used OB"),
                ("expired", "Expired/restored"),
            ) if status in statuses
        ) or "No Shifting OB earned in selected year"
        values = [
            report_employee_name(employee),
            employee.job_title or "",
            "Yes",
            earned,
            available,
            regular_excess.quantize(Decimal("0.01")),
            vl_credit.quantize(Decimal("0.01")),
            notes,
        ]
        for col, value in enumerate(values, 1):
            summary.cell(row_index, col, value)
    if employees:
        _apply_grid(summary, 4, 4 + len(employees), 1, 8)
    _set_widths(summary, {1: 28, 2: 24, 3: 10, 4: 21, 5: 20, 6: 24, 7: 12, 8: 34})
    summary.freeze_panes = "A5"

    rule_row = 6 + len(employees)
    summary.cell(rule_row, 1, "Rule").fill = HEADER_FILL
    summary.cell(rule_row, 2, "Current Setting").fill = HEADER_FILL
    summary.cell(rule_row, 3, "Meaning").fill = HEADER_FILL
    for cell in summary[rule_row][:3]:
        cell.font = Font(bold=True)
    expiration_text = (
        "Follow Leave Credit Reset Date"
        if expiration_mode == "follow_leave_reset"
        else f"Custom: {calendar.month_name[expiration_month]} {expiration_day}"
    )
    rules = [
        ("Qualifying Block", f"{block_hours:g} hours", "Exact qualifying block is consumed; excess stays Regular OT."),
        ("Required Blocks / Cutoff", str(required_blocks), f"{required_blocks} qualifying blocks in the same cutoff create 1 whole-day OB."),
        ("Cutoff", f"1–{cutoff_day} / {cutoff_day + 1}–month end", "Calendar dates; weekends do not move the cutoff."),
        ("Availability Delay", f"{availability_cutoffs} completed cutoffs", "Pending earned OB is protected until usable."),
        ("Expiration", expiration_text, "Unused available OB restores its source hours to Regular OT."),
        (
            "Straight OT Conversion",
            f"{straight_ot_threshold_hours:g} continuous OT hours = {straight_ot_vl_days:g} VL",
            (
                "VL credit + OT payable (dinner deduction still applies)."
                if straight_ot_also_payable
                else "Qualifying hours are converted to VL and are not also OT payable; only excess may remain payable."
            ),
        ),
        ("Exclusions", " / ".join(excluded_positions) or "None", "Excluded positions receive no Shifting/OB or straight-OT VL benefit."),
    ]
    for offset, row in enumerate(rules, rule_row + 1):
        for col, value in enumerate(row, 1):
            summary.cell(offset, col, value)
    _apply_grid(summary, rule_row, rule_row + len(rules), 1, 3)

    cutoff_sheet = workbook.create_sheet("Cut-Off Date")
    cutoff_sheet.merge_cells("A1:G1")
    cutoff_sheet["A1"] = f"CUT-OFF DATE REFERENCE — {report_year}"
    cutoff_sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    cutoff_sheet["A1"].fill = TITLE_FILL
    cutoff_sheet["A1"].alignment = Alignment(horizontal="center")
    cutoff_headers = [
        "Period", "Cutoff Start", "Cutoff End", "Availability After",
        "Expiration Basis", "Example Expiration", "Weekend Handling",
    ]
    for col, value in enumerate(cutoff_headers, 1):
        cell = cutoff_sheet.cell(3, col, value)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    row_index = 4
    for month in range(1, 13):
        last = calendar.monthrange(report_year, month)[1]
        split = min(max(int(cutoff_day), 1), last - 1)
        for label, start_day, end_day in (("1st", 1, split), ("2nd", split + 1, last)):
            start = date(report_year, month, start_day)
            end = date(report_year, month, end_day)
            availability = OvertimeService._availability_after_cutoffs(
                end, cutoff_day=cutoff_day, completed_cutoffs=availability_cutoffs
            )
            cutoff_sheet.cell(row_index, 1, f"{calendar.month_name[month]} {report_year} - {label}")
            cutoff_sheet.cell(row_index, 2, start)
            cutoff_sheet.cell(row_index, 3, end)
            cutoff_sheet.cell(row_index, 4, availability)
            cutoff_sheet.cell(row_index, 5, expiration_text)
            custom_example = date(
                report_year,
                expiration_month,
                min(expiration_day, calendar.monthrange(report_year, expiration_month)[1]),
            )
            cutoff_sheet.cell(row_index, 6, custom_example if expiration_mode == "custom_date" else "See Leave Reset")
            cutoff_sheet.cell(row_index, 7, "Included")
            for col in (2, 3, 4, 6):
                if isinstance(cutoff_sheet.cell(row_index, col).value, date):
                    cutoff_sheet.cell(row_index, col).number_format = DATE_FORMAT
            row_index += 1
    _apply_grid(cutoff_sheet, 3, row_index - 1, 1, 7)
    _set_widths(cutoff_sheet, {1: 28, 2: 15, 3: 15, 4: 18, 5: 31, 6: 18, 7: 18})
    cutoff_sheet.freeze_panes = "A4"
    cutoff_sheet.cell(row_index + 1, 1, "NOTE:").font = Font(bold=True)
    cutoff_sheet.cell(row_index + 2, 1, "1.")
    cutoff_sheet.cell(row_index + 2, 2, "Calendar cutoffs include weekends; Saturday/Sunday do not move the 15th or month-end boundary.")
    cutoff_sheet.cell(row_index + 3, 1, "2.")
    cutoff_sheet.cell(row_index + 3, 2, "Availability delay and expiration are separate. Pending earned credit is protected until it first becomes usable.")
    cutoff_sheet.merge_cells(start_row=row_index + 2, start_column=2, end_row=row_index + 2, end_column=7)
    cutoff_sheet.merge_cells(start_row=row_index + 3, start_column=2, end_row=row_index + 3, end_column=7)

    used_sheet_names = {"summary", "cut-off date"}
    for employee in employees:
        sheet = workbook.create_sheet(_sheet_name(report_employee_name(employee), used_sheet_names))
        employee_requests = sorted(
            requests_by_employee.get(employee.id, []),
            key=lambda item: (item.date_rendered, item.id),
        )
        employee_credits = credits_by_employee.get(employee.id, [])
        sheet.merge_cells("A1:Q1")
        sheet["A1"] = f"SHIFTING CREDITS — {report_employee_name(employee)}"
        sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
        sheet["A1"].fill = TITLE_FILL
        sheet["A1"].alignment = Alignment(horizontal="center")
        department_name = getattr(getattr(employee, "department", None), "name", "") or ""
        sheet["A2"] = f"Employee No.: {employee.employee_number}"
        sheet["B2"] = f"Position: {employee.job_title or ''}"
        sheet["C2"] = f"Project / Department: {department_name}"
        main_headers = [
            "Date", "OT Hours", "", "Shifting Credits", "", "Straight 8h OT",
            "Availability Date", "Expiration Date", "Project", "Date of Usage",
            "Remarks", "Regular OT Excess", "VL Credit", "", "Summary", "Value", "Unit",
        ]
        sub_headers = [
            "", "From", "To", "Qualifying Block Hrs", "OB Credit (Days)", "", "", "", "", "", "",
            "Hours", "Days", "", "Total Qualifying Shifting Hours", "", "hrs",
        ]
        for col, value in enumerate(main_headers, 1):
            cell = sheet.cell(3, col, value)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col, value in enumerate(sub_headers, 1):
            cell = sheet.cell(4, col, value)
            cell.font = Font(bold=True)
            cell.fill = SUBHEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        group_items: dict[str, list] = {}
        for request in employee_requests:
            if request.shifting_credit_group:
                group_items.setdefault(request.shifting_credit_group, []).append(request)
        final_request_ids = {
            max(items, key=lambda item: (item.date_rendered, item.id)).id
            for items in group_items.values()
        }
        row_index = 5
        total_regular_excess = Decimal("0.00")
        total_vl = Decimal("0.00")
        total_qualifying = Decimal("0.00")
        for request in employee_requests:
            credit = credit_by_group.get(request.shifting_credit_group or "")
            qualifying = _decimal(request.shifting_credit_hours)
            total_qualifying += qualifying
            is_final = request.id in final_request_ids
            ob_days = _decimal(credit.ob_credit_days) if credit is not None and is_final else Decimal("0.00")
            straight_hours = _decimal(request.estimated_hours) if _decimal(request.additional_vl_days) > 0 else Decimal("0.00")
            vl_days = _decimal(request.additional_vl_days)
            total_vl += vl_days
            regular_excess = Decimal("0.00")
            if request.shifting_credit_group:
                regular_excess = _decimal(request.payable_hours)
            elif vl_days > 0:
                # For straight-OT VL conversion, payable_hours already reflects
                # whether the company enabled “Consider it also as OT Payable”.
                regular_excess = _decimal(request.payable_hours)
            total_regular_excess += regular_excess
            if request.shifting_credit_group and not is_final:
                remarks = "VALID - qualifying block"
            elif request.shifting_credit_group and is_final:
                remarks = _credit_status_note(credit)
            elif vl_days > 0:
                remarks = (
                    f"Straight OT -> {vl_days:.2f} VL + OT payable"
                    if bool(getattr(request, "straight_vl_also_payable", False))
                    else f"Straight OT -> {vl_days:.2f} VL conversion"
                )
            elif _decimal(request.estimated_hours) >= _decimal(block_hours):
                remarks = "UNMATCHED - remains Regular OT"
            else:
                remarks = "Regular OT"
            values = [
                request.date_rendered,
                _local_time(request.ot_time_start, timezone_value),
                _local_time(request.ot_time_end, timezone_value),
                qualifying,
                ob_days,
                straight_hours,
                credit.availability_date if credit is not None and is_final else None,
                credit.expiration_date if credit is not None and is_final else None,
                department_name,
                credit.usage_date if credit is not None and is_final else None,
                remarks,
                regular_excess.quantize(Decimal("0.01")),
                vl_days,
            ]
            for col, value in enumerate(values, 1):
                sheet.cell(row_index, col, value)
            for col in (1, 7, 8, 10):
                if isinstance(sheet.cell(row_index, col).value, date):
                    sheet.cell(row_index, col).number_format = DATE_FORMAT
            row_index += 1

        summary_items = [
            ("Total Qualifying Shifting Hours", total_qualifying, "hrs"),
            ("OB Shifting Credit Earned", sum((_decimal(c.ob_credit_days) for c in employee_credits), Decimal("0")), "days"),
            ("OB Shifting Credit Available", sum((_decimal(c.ob_credit_days) for c in employee_credits if c.status == "available"), Decimal("0")), "days"),
            ("OB Shifting Credit Reserved", sum((_decimal(c.ob_credit_days) for c in employee_credits if c.status == "reserved"), Decimal("0")), "days"),
            ("OB Shifting Credit Used", sum((_decimal(c.ob_credit_days) for c in employee_credits if c.status == "used"), Decimal("0")), "days"),
            ("OB Shifting Credit Expired", sum((_decimal(c.ob_credit_days) for c in employee_credits if c.status == "expired"), Decimal("0")), "days"),
            ("Regular OT Excess / Restored", total_regular_excess.quantize(Decimal("0.01")), "hrs"),
            (f"VL Credit from Straight {straight_ot_threshold_hours:g}h OT", total_vl.quantize(Decimal("0.01")), "days"),
        ]
        for offset, (label, value, unit) in enumerate(summary_items, 4):
            sheet.cell(offset, 15, label)
            sheet.cell(offset, 16, value)
            sheet.cell(offset, 17, unit)
            for col in (15, 16, 17):
                sheet.cell(offset, col).fill = SUMMARY_FILL
                sheet.cell(offset, col).border = BORDER
        max_row = max(row_index - 1, 4)
        _apply_grid(sheet, 3, max_row, 1, 13)
        _set_widths(
            sheet,
            {1: 13, 2: 10, 3: 10, 4: 19, 5: 18, 6: 16, 7: 18, 8: 18, 9: 22,
             10: 16, 11: 34, 12: 19, 13: 14, 14: 3, 15: 34, 16: 14, 17: 10},
        )
        sheet.freeze_panes = "A5"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
