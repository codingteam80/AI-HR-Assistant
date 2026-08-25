"""Excel export for company disciplinary records."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def build_disciplinary_records_excel(rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Disciplinary Records"

    headers = list(rows[0].keys()) if rows else [
        "Case ID", "Employee", "Violation", "Incident Date", "Incident / Case Description",
        "Evidence / Remarks", "Project / Team / Department", "Previous Offense Count",
        "Current Offense", "Suggested Disciplinary Action", "Actual Action Taken", "Issued By",
        "Reviewed / Approved By", "Date Issued", "Employee Acknowledgment", "Case Status", "Notes",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for index, header in enumerate(headers, start=1):
        max_length = len(str(header))
        for cell in sheet[get_column_letter(index)]:
            max_length = max(max_length, min(60, len(str(cell.value or ""))))
        sheet.column_dimensions[get_column_letter(index)].width = min(62, max(12, max_length + 2))

    sheet.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
