"""Combined DTR, overtime, and leave Excel workbook generation."""

from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Iterable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.leave_codes import (
    LEAVE_DURATION_OPTIONS,
    LEAVE_REASON_OPTIONS,
    duration_label,
    reason_label,
)
from models.attendance_record import AttendanceRecord
from models.employee import Employee
from models.leave_request import LeaveRequest
from models.overtime_request import OvertimeRequest
from services.attendance_service import ATTENDANCE_COLORS, WEEKEND_COLOR


HEADER_FILL = "FFFFFF"
OT_HEADER_FILL = "9DC3E6"
OT_ACCENT_FILL = "92D050"
LEAVE_HEADER_FILL = "FFF59D"
GRID_SIDE = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(
    left=GRID_SIDE,
    right=GRID_SIDE,
    top=GRID_SIDE,
    bottom=GRID_SIDE,
)
APPROVED_LEAVE_STATUSES = {
    "scheduled",
    "approved",
    "in_progress",
    "completed",
    "partially_cancelled",
}
MISSING_NAME_TOKENS = {"n/a", "na", "none", "null", "-", "—"}
LEAVE_TYPE_EXPORT_CODES = {
    "90401": "90401",
    "vacation": "90401",
    "vacation leave": "90401",
    "vl": "90401",
    "90402": "90402",
    "sick": "90402",
    "sick leave": "90402",
    "sl": "90402",
    "90404": "90404",
    "bereavement": "90404",
    "bereavement leave": "90404",
    "90405": "90405",
    "honeymoon": "90405",
    "honeymoon leave": "90405",
    "90406": "90406",
    "paternity": "90406",
    "paternity leave": "90406",
    "90407": "90407",
    "maternity": "90407",
    "maternity leave": "90407",
    "90409": "90409",
    "birthday": "90409",
    "birthday leave": "90409",
    "90410": "90410",
    "emergency": "90410",
    "emergency leave": "90410",
    "el": "90410",
}
LEAVE_TYPE_EXPORT_DESCRIPTIONS = {
    "90401": "Vacation",
    "90402": "Sick",
    "90404": "Bereavement",
    "90405": "Honeymoon",
    "90406": "Paternity",
    "90407": "Maternity",
    "90409": "Birthday",
    "90410": "Emergency",
}
def _local_datetime(
    value: datetime | None,
    timezone_value: ZoneInfo,
) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone_value).replace(tzinfo=None)


def _status_fill(status: str | None) -> str | None:
    color = ATTENDANCE_COLORS.get(str(status or "").upper())
    return color.removeprefix("#") if color else None


def _attendance_leave_code(record: AttendanceRecord | None) -> str:
    """Return the short DTR leave marker from the linked request."""

    if record is None:
        return ""
    request = getattr(record, "leave_request", None)
    leave_type = getattr(request, "leave_type", None)
    label = (
        f"{getattr(leave_type, 'code', '')} {getattr(leave_type, 'name', '')}"
    ).upper()
    if "SICK" in label or label.strip().startswith("SL"):
        return "SL"
    if "EMERGENCY" in label or label.strip().startswith("EL"):
        return "EL"
    if request is not None:
        return "VL"
    status = str(getattr(record, "work_status", "") or "").upper()
    return status if status in {"VL", "SL", "EL"} else ""


def _employee_section(employee: Employee) -> str:
    return employee.department.name if employee.department is not None else ""


def _name_part(value: object) -> str:
    text = str(value or "").strip()
    if not text or text.casefold() in MISSING_NAME_TOKENS:
        return ""
    return text


def _report_employee_name(employee: Employee) -> str:
    """Return Last, First Suffix, M. for payroll report templates."""

    last_name = _name_part(getattr(employee, "last_name", None))
    first_name = _name_part(getattr(employee, "first_name", None))
    suffix = _name_part(getattr(employee, "suffix", None))
    middle_name = _name_part(getattr(employee, "middle_name", None))
    if not last_name and not first_name:
        return _name_part(getattr(employee, "full_name", None))
    given_name = " ".join(part for part in (first_name, suffix) if part)
    primary_name = ", ".join(
        part for part in (last_name, given_name) if part
    )
    middle_initial = f"{middle_name[0].upper()}." if middle_name else ""
    return f"{primary_name}, {middle_initial}" if middle_initial else primary_name


def _apply_header(
    sheet,
    *,
    fill: str,
    auto_filter: bool = True,
    row_height: float = 22,
) -> None:
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="10172A")
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER
    sheet.row_dimensions[1].height = row_height
    if auto_filter:
        sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = True


def _set_widths(sheet, widths: Iterable[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _build_dtr_sheet(
    workbook: Workbook,
    *,
    employees: list[Employee],
    records: list[AttendanceRecord],
    dates: list[date],
    scheduled_workdays: dict[date, bool],
    timezone_value: ZoneInfo,
) -> None:
    sheet = workbook.active
    sheet.title = "DTR Logs"
    sheet.append(["Name", "Section", "Property", *dates])

    record_map = {
        (record.employee_id, record.attendance_date): record
        for record in records
    }
    leave_codes = {
        str(record.work_status or "").upper()
        for record in records
        if record.leave_request_id is not None
    } | {"VL", "SL", "EL"}

    for employee in employees:
        section = _employee_section(employee)
        report_name = _report_employee_name(employee)
        start_row = [report_name, section, "Start Time"]
        end_row = [report_name, section, "End Time"]
        for work_date in dates:
            record = record_map.get((employee.id, work_date))
            status = str(record.work_status or "").upper() if record else ""
            leave_code = _attendance_leave_code(record)
            duration_code = (
                str(getattr(record, "leave_duration_code", "") or "")
                if record is not None
                else ""
            )
            if record is not None and leave_code and duration_code not in {
                "90501",
                "90502",
            }:
                start_row.append(leave_code)
                end_row.append(leave_code)
                continue
            sessions = list(getattr(record, "sessions", ()) or ()) if record else []
            if sessions:
                start_values = []
                end_values = []
                for attendance_session in sessions:
                    local_in = _local_datetime(
                        attendance_session.rounded_time_in,
                        timezone_value,
                    )
                    local_out = _local_datetime(
                        attendance_session.rounded_time_out,
                        timezone_value,
                    )
                    start_values.append(
                        f"{local_in:%H:%M} {attendance_session.work_status}"
                    )
                    end_values.append(
                        f"{local_out:%H:%M} {attendance_session.work_status}"
                        if local_out
                        else f"OPEN {attendance_session.work_status}"
                    )
                if duration_code == "90501":
                    start_values.insert(0, f"AM {leave_code or 'LEAVE'}")
                    end_values.insert(0, f"AM {leave_code or 'LEAVE'}")
                elif duration_code == "90502":
                    start_values.append(f"PM {leave_code or 'LEAVE'}")
                    end_values.append(f"PM {leave_code or 'LEAVE'}")
                start_row.append(" / ".join(start_values))
                end_row.append(" / ".join(end_values))
            else:
                local_in = (
                    _local_datetime(record.time_in, timezone_value)
                    if record is not None
                    else None
                )
                local_out = (
                    _local_datetime(record.time_out, timezone_value)
                    if record is not None
                    else None
                )
                start_row.append(local_in.time() if local_in else None)
                end_row.append(local_out.time() if local_out else None)
        sheet.append(start_row)
        sheet.append(end_row)

        for row_index in (sheet.max_row - 1, sheet.max_row):
            for offset, work_date in enumerate(dates, start=4):
                cell = sheet.cell(row=row_index, column=offset)
                record = record_map.get((employee.id, work_date))
                if not scheduled_workdays.get(work_date, work_date.weekday() < 5):
                    fill = WEEKEND_COLOR.removeprefix("#")
                elif record is not None and getattr(
                    record,
                    "leave_duration_code",
                    None,
                ) in {"90501", "90502"}:
                    fill = ATTENDANCE_COLORS["VL"].removeprefix("#")
                elif record is not None:
                    fill = _status_fill(record.work_status)
                else:
                    fill = None
                if fill:
                    cell.fill = PatternFill("solid", fgColor=fill)
                cell.number_format = "h:mm"
                cell.alignment = Alignment(
                    horizontal="right" if cell.value else "center",
                    vertical="center",
                    wrap_text=isinstance(cell.value, str),
                )

    for index, work_date in enumerate(dates, start=4):
        header = sheet.cell(row=1, column=index)
        header.value = work_date
        header.number_format = "yyyy/m/d"
        sheet.column_dimensions[get_column_letter(index)].width = 18
    _set_widths(sheet, (24, 18, 14))
    _apply_header(
        sheet,
        fill=HEADER_FILL,
        auto_filter=False,
        row_height=20,
    )
    sheet.freeze_panes = "D2"

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column <= 3:
                cell.alignment = Alignment(vertical="center")
            cell.border = THIN_BORDER
        sheet.row_dimensions[row[0].row].height = 30


def _ot_segments(
    record: AttendanceRecord,
    *,
    timezone_value: ZoneInfo,
) -> list[tuple[datetime, datetime, Decimal]]:
    """Allocate computed OT backward across completed rounded sessions."""

    remaining_minutes = Decimal(record.ot_hours) * Decimal("60")
    if remaining_minutes <= 0:
        return []
    completed = []
    for item in list(getattr(record, "sessions", ()) or ()):
        local_in = _local_datetime(item.rounded_time_in, timezone_value)
        local_out = _local_datetime(item.rounded_time_out, timezone_value)
        if local_in and local_out and local_out > local_in:
            completed.append((local_in, local_out))
    if not completed:
        local_in = _local_datetime(record.time_in, timezone_value)
        local_out = _local_datetime(record.time_out, timezone_value)
        if local_in and local_out and local_out > local_in:
            completed.append((local_in, local_out))

    output: list[tuple[datetime, datetime, Decimal]] = []
    for local_in, local_out in reversed(completed):
        if remaining_minutes <= 0:
            break
        session_minutes = Decimal(str((local_out - local_in).total_seconds())) / Decimal("60")
        allocated_minutes = min(session_minutes, remaining_minutes)
        segment_start = local_out - timedelta(minutes=float(allocated_minutes))
        output.append(
            (
                segment_start,
                local_out,
                (allocated_minutes / Decimal("60")).quantize(Decimal("0.01")),
            )
        )
        remaining_minutes -= allocated_minutes
    return list(reversed(output))


def _clock_text(value: datetime | None) -> str:
    """Return Excel-template time text without platform-specific strftime."""

    if value is None:
        return ""
    return f"{value.hour % 12 or 12:02d}:{value.minute:02d}"


def _build_ot_sheet(
    workbook: Workbook,
    *,
    records: list[AttendanceRecord],
    overtime_requests: list[OvertimeRequest],
    timezone_value: ZoneInfo,
) -> None:
    sheet = workbook.create_sheet("Overtime File")
    sheet.append(
        [
            "emp_id",
            "emp_name",
            "date_rendered",
            "OT_time_start",
            "",
            "OT_time_end",
            "",
            "estimated_hours",
            "OT_type",
            "OT_purpose",
            "travel_fare",
            "travel_route",
            "dinner_break_flag",
            "",
        ]
    )
    approved_requests = sorted(
        (item for item in overtime_requests if item.status == "approved"),
        key=lambda item: (
            item.date_rendered,
            item.employee.last_name,
            item.employee.first_name,
            item.id,
        ),
    )
    approved_dates: set[tuple[int, date]] = set()
    for request in approved_requests:
        local_start = _local_datetime(request.ot_time_start, timezone_value)
        local_end = _local_datetime(request.ot_time_end, timezone_value)
        approved_dates.add((request.employee_id, request.date_rendered))
        sheet.append(
            [
                request.employee.employee_number,
                _report_employee_name(request.employee),
                request.date_rendered,
                _clock_text(local_start),
                local_start.strftime("%p") if local_start else "",
                _clock_text(local_end),
                local_end.strftime("%p") if local_end else "",
                float(request.estimated_hours),
                request.ot_type,
                request.ot_purpose,
                float(request.travel_fare) if request.travel_fare is not None else None,
                request.travel_route or "",
                1 if request.dinner_break_flag else 0,
                "CC",
            ]
        )

    # A DTR-only row contains only the four facts actually available from
    # attendance. The request-only columns remain blank instead of being
    # guessed. An approved request replaces the DTR-only row for that date.
    overtime_records = sorted(
        (record for record in records if Decimal(record.ot_hours) > 0),
        key=lambda item: (
            item.attendance_date,
            item.employee.last_name,
            item.employee.first_name,
        ),
    )
    for record in overtime_records:
        if (record.employee_id, record.attendance_date) in approved_dates:
            continue
        sheet.append(
            [
                record.employee.employee_number,
                _report_employee_name(record.employee),
                record.attendance_date,
                "",
                "",
                "",
                "",
                float(record.ot_hours),
                "",
                "",
                None,
                "",
                None,
                "",
            ]
        )

    _apply_header(sheet, fill=OT_HEADER_FILL, row_height=20)
    for index in (4, 5, 8, 10, 12):
        sheet.cell(row=1, column=index).fill = PatternFill(
            "solid",
            fgColor=OT_ACCENT_FILL,
        )
    _set_widths(
        sheet,
        (13, 38, 16, 13, 6, 13, 6, 17, 22, 40, 14, 26, 20, 10),
    )
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2):
        row[2].number_format = "yyyy-mm-dd"
        row[7].number_format = "0.##"
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = THIN_BORDER
        sheet.row_dimensions[row[0].row].height = 18


def _approval_date(
    value: datetime | None,
    timezone_value: ZoneInfo,
) -> date | None:
    localized = _local_datetime(value, timezone_value)
    return localized.date() if localized is not None else None


def _filed_date(request: LeaveRequest, timezone_value: ZoneInfo) -> date | None:
    value = getattr(request, "submitted_at", None) or getattr(
        request,
        "created_at",
        None,
    )
    return _approval_date(value, timezone_value)


def _leave_type_export_code(request: LeaveRequest) -> str:
    leave_type = request.leave_type
    for candidate in (leave_type.code, leave_type.name):
        mapped = LEAVE_TYPE_EXPORT_CODES.get(_name_part(candidate).casefold())
        if mapped:
            return mapped
    return ""


def _code_description(code: str, description: str) -> str:
    """Render coded report values in the payroll-readable export format."""

    return f"{code} - {description}" if code and description else ""


def _leave_type_export_value(request: LeaveRequest) -> str:
    code = _leave_type_export_code(request)
    return _code_description(code, LEAVE_TYPE_EXPORT_DESCRIPTIONS.get(code, ""))


def _leave_remarks(request: LeaveRequest) -> str:
    values = []
    requested_days = Decimal(getattr(request, "requested_days", 0) or 0)
    lwop_days = Decimal(getattr(request, "lwop_days", 0) or 0)
    paid_days = max(Decimal("0.00"), requested_days - lwop_days)
    if lwop_days > 0:
        values.append(
            "Credit/LWP Split: "
            f"{paid_days.normalize()} credited day(s), "
            f"{lwop_days.normalize()} LWP day(s)"
        )
    if request.leader_comment:
        values.append(f"Pre-approver: {request.leader_comment.strip()}")
    if request.manager_comment:
        values.append(f"Approver: {request.manager_comment.strip()}")
    if request.cancellation_comment:
        values.append(f"Cancellation: {request.cancellation_comment.strip()}")
    return " | ".join(values)


def _duration_code(request: LeaveRequest) -> str:
    """Return the stored duration code with a safe legacy default."""

    code = str(getattr(request, "duration_code", "") or "90503")
    return code if code in LEAVE_DURATION_OPTIONS else "90503"


def _duration_export_value(request: LeaveRequest) -> str:
    return duration_label(_duration_code(request))


def _build_leave_sheet(
    workbook: Workbook,
    *,
    leave_requests: list[LeaveRequest],
    timezone_value: ZoneInfo,
) -> None:
    sheet = workbook.create_sheet("Leave File")
    sheet.append(
        [
            "Emp ID",
            "Emp Name",
            "Leave Type",
            "Date Filed",
            "Start Date",
            "End Date",
            "Duration",
            "Reason for Leave",
            "Reason for Leave: Others",
            "Pre-approver",
            "Date Pre-approved",
            "Approver",
            "Date Approved",
            "Remarks",
        ]
    )
    approved_requests = sorted(
        (
            request
            for request in leave_requests
            if request.status in APPROVED_LEAVE_STATUSES
        ),
        key=lambda item: (
            item.start_date,
            item.employee.last_name,
            item.employee.first_name,
        ),
    )
    for request in approved_requests:
        pre_approver = (
            _report_employee_name(request.leader_approver)
            if request.leader_reviewed_at is not None
            and request.leader_approver is not None
            else ""
        )
        approver = (
            _report_employee_name(request.manager)
            if (request.approved_at or request.reviewed_at) is not None
            and request.manager is not None
            else ""
        )
        sheet.append(
            [
                request.employee.employee_number,
                _report_employee_name(request.employee),
                _leave_type_export_value(request),
                _filed_date(request, timezone_value),
                request.start_date,
                request.end_date,
                _duration_export_value(request),
                reason_label(str(getattr(request, "reason_code", "") or "0")),
                (
                    getattr(request, "reason_other", None)
                    or getattr(request, "reason", "")
                    if str(getattr(request, "reason_code", "") or "0") == "0"
                    else ""
                ),
                pre_approver,
                _approval_date(request.leader_reviewed_at, timezone_value),
                approver,
                _approval_date(
                    request.approved_at or request.reviewed_at,
                    timezone_value,
                ),
                _leave_remarks(request),
            ]
        )

    _apply_header(sheet, fill=LEAVE_HEADER_FILL)
    _set_widths(
        sheet,
        (14, 38, 24, 16, 16, 16, 20, 28, 44, 22, 19, 22, 18, 44),
    )
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2):
        for index in (3, 4, 5, 10, 12):
            row[index].number_format = "yyyy-mm-dd"
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            cell.border = THIN_BORDER
        sheet.row_dimensions[row[0].row].height = 18


def build_combined_hr_excel(
    *,
    employees: list[Employee],
    attendance_records: list[AttendanceRecord],
    leave_requests: list[LeaveRequest],
    overtime_requests: list[OvertimeRequest] | None = None,
    start_date: date,
    end_date: date,
    scheduled_workdays: dict[date, bool],
    timezone_name: str,
) -> bytes:
    """Return one workbook containing the three requested HR templates."""

    timezone_value = ZoneInfo(timezone_name)
    workbook = Workbook()
    dates = [
        start_date + timedelta(days=offset)
        for offset in range((end_date - start_date).days + 1)
    ]
    ordered_employees = sorted(
        employees,
        key=lambda item: (item.last_name.casefold(), item.first_name.casefold()),
    )
    _build_dtr_sheet(
        workbook,
        employees=ordered_employees,
        records=attendance_records,
        dates=dates,
        scheduled_workdays=scheduled_workdays,
        timezone_value=timezone_value,
    )
    _build_ot_sheet(
        workbook,
        records=attendance_records,
        overtime_requests=overtime_requests or [],
        timezone_value=timezone_value,
    )
    _build_leave_sheet(
        workbook,
        leave_requests=leave_requests,
        timezone_value=timezone_value,
    )
    workbook.properties.title = "Combined DTR, OT, and Leave Report"
    workbook.properties.subject = (
        f"HR operational report from {start_date.isoformat()} "
        f"to {end_date.isoformat()}"
    )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
