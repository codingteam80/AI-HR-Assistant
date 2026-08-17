"""Secure in-app preview dialog for authorized company-form files.

The caller must obtain file bytes through the service layer first. This
component only renders those already-authorized bytes and never reads an
arbitrary client path.
"""

from __future__ import annotations

import base64
from html import escape
from io import BytesIO
import json
from pathlib import Path

import pandas as pd
import streamlit as st
from docx import Document

from ui.components.browser_bridge import render_browser_bridge


_MAX_PREVIEW_ROWS = 250
_MAX_PREVIEW_COLUMNS = 40
_PREVIEW_CONTENT_HEIGHT = 600
_PREVIEW_WIDGET_HEIGHT = 560
_DISMISS_CONTEXT_KEY = "_company_file_preview_dismiss_context"


def _print_document_shell(*, filename: str, body: str) -> str:
    """Wrap escaped/controlled preview markup in a print-ready document."""

    return f"""
        <!doctype html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>{escape(filename)}</title>
            <style>
                @page {{ margin: 14mm; }}
                * {{ box-sizing: border-box; }}
                body {{
                    margin: 0;
                    color: #111827;
                    background: #FFFFFF;
                    font-family: Arial, Helvetica, sans-serif;
                    font-size: 11pt;
                    line-height: 1.45;
                }}
                h1 {{ margin: 0 0 18px; font-size: 17pt; }}
                h2 {{ margin: 22px 0 10px; font-size: 13pt; }}
                p {{ margin: 0 0 10px; white-space: pre-wrap; }}
                pre {{
                    margin: 0;
                    white-space: pre-wrap;
                    overflow-wrap: anywhere;
                    font: 10pt/1.45 Consolas, monospace;
                }}
                table {{
                    width: 100%;
                    margin: 0 0 18px;
                    border-collapse: collapse;
                    page-break-inside: auto;
                }}
                th, td {{
                    padding: 6px 8px;
                    border: 1px solid #9CA3AF;
                    text-align: left;
                    vertical-align: top;
                    overflow-wrap: anywhere;
                }}
                th {{ background: #F3F4F6; font-weight: 700; }}
                tr {{ page-break-inside: avoid; page-break-after: auto; }}
                img {{
                    display: block;
                    width: auto;
                    max-width: 100%;
                    height: auto;
                    max-height: 255mm;
                    margin: 0 auto;
                    object-fit: contain;
                }}
            </style>
        </head>
        <body>
            <h1>{escape(filename)}</h1>
            {body}
        </body>
        </html>
    """


def _build_print_payload(
    *,
    filename: str,
    mime_type: str,
    data: bytes,
) -> dict[str, str]:
    """Create a safe browser-print payload for one authorized file."""

    extension = Path(filename).suffix.lower()
    encoded = base64.b64encode(data).decode("ascii")

    if extension == ".pdf" or mime_type == "application/pdf":
        return {
            "mode": "pdf",
            "mimeType": "application/pdf",
            "base64": encoded,
        }

    if extension in {".png", ".jpg", ".jpeg"} or mime_type.startswith(
        "image/"
    ):
        safe_mime = mime_type if mime_type.startswith("image/") else "image/png"
        body = (
            f'<img src="data:{escape(safe_mime)};base64,{encoded}" '
            f'alt="{escape(filename)}">'
        )
        return {
            "mode": "html",
            "html": _print_document_shell(filename=filename, body=body),
        }

    if extension == ".docx":
        document = Document(BytesIO(data))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(f"<p>{escape(paragraph.text)}</p>")
        for table in document.tables:
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            if rows:
                parts.append(
                    pd.DataFrame(rows).to_html(
                        index=False,
                        header=False,
                        border=0,
                        escape=True,
                    )
                )
        body = "".join(parts) or "<p>No printable document content.</p>"
        return {
            "mode": "html",
            "html": _print_document_shell(filename=filename, body=body),
        }

    if extension == ".xlsx":
        workbook = pd.ExcelFile(BytesIO(data), engine="openpyxl")
        parts = []
        for sheet_name in workbook.sheet_names:
            frame = workbook.parse(sheet_name=sheet_name, header=None)
            clipped = frame.iloc[:_MAX_PREVIEW_ROWS, :_MAX_PREVIEW_COLUMNS]
            parts.append(f"<h2>{escape(str(sheet_name))}</h2>")
            parts.append(
                clipped.to_html(
                    index=False,
                    header=False,
                    border=0,
                    escape=True,
                )
            )
        return {
            "mode": "html",
            "html": _print_document_shell(
                filename=filename,
                body="".join(parts),
            ),
        }

    if extension == ".csv":
        frame = pd.read_csv(BytesIO(data))
        clipped = frame.iloc[:_MAX_PREVIEW_ROWS, :_MAX_PREVIEW_COLUMNS]
        body = clipped.to_html(index=False, border=0, escape=True)
        return {
            "mode": "html",
            "html": _print_document_shell(filename=filename, body=body),
        }

    if extension == ".txt" or mime_type.startswith("text/"):
        body = f"<pre>{escape(_read_text(data))}</pre>"
        return {
            "mode": "html",
            "html": _print_document_shell(filename=filename, body=body),
        }

    return {"mode": "unsupported"}


def _launch_print_dialog(payload: dict[str, str]) -> None:
    """Open the native browser print dialog from a Streamlit button click."""

    if payload["mode"] == "unsupported":
        return

    payload_json = json.dumps(payload).replace("</", "<\\/")

    component = f"""
        <script>
                (() => {{
                    const payload = {payload_json};
                    const parentDocument = window.parent.document;

                    const removeFrame = (
                        frame,
                        objectUrl = null,
                        delay = 1500
                    ) => {{
                        window.setTimeout(() => {{
                            frame.remove();
                            if (objectUrl) URL.revokeObjectURL(objectUrl);
                        }}, delay);
                    }};

                    const printHtml = (html) => {{
                        const frame = parentDocument.createElement('iframe');
                        frame.setAttribute('aria-hidden', 'true');
                        frame.style.cssText = (
                            'position:fixed;width:1px;height:1px;right:0;bottom:0;'
                            + 'border:0;opacity:0;pointer-events:none;'
                        );
                        parentDocument.body.appendChild(frame);
                        const printDocument = frame.contentDocument;
                        printDocument.open();
                        printDocument.write(html);
                        printDocument.close();
                        window.setTimeout(() => {{
                            frame.contentWindow.focus();
                            frame.contentWindow.print();
                            removeFrame(frame);
                        }}, 250);
                    }};

                    const printPdf = (encoded, mimeType) => {{
                        const binary = window.atob(encoded);
                        const bytes = new Uint8Array(binary.length);
                        for (let index = 0; index < binary.length; index += 1) {{
                            bytes[index] = binary.charCodeAt(index);
                        }}
                        const objectUrl = URL.createObjectURL(
                            new Blob([bytes], {{ type: mimeType }})
                        );
                        const frame = parentDocument.createElement('iframe');
                        frame.setAttribute('aria-hidden', 'true');
                        frame.style.cssText = (
                            'position:fixed;width:1px;height:1px;right:0;bottom:0;'
                            + 'border:0;opacity:0;pointer-events:none;'
                        );
                        let printed = false;
                        const launchPrint = () => {{
                            if (printed) return;
                            printed = true;
                            window.setTimeout(() => {{
                                try {{
                                    frame.contentWindow.addEventListener(
                                        'afterprint',
                                        () => removeFrame(frame, objectUrl),
                                        {{ once: true }}
                                    );
                                    frame.contentWindow.focus();
                                    frame.contentWindow.print();
                                    removeFrame(frame, objectUrl, 60000);
                                }} catch (error) {{
                                    frame.remove();
                                    const printWindow = window.parent.open(
                                        objectUrl,
                                        '_blank'
                                    );
                                    if (printWindow) {{
                                        window.setTimeout(
                                            () => printWindow.print(),
                                            900
                                        );
                                    }}
                                    window.setTimeout(
                                        () => URL.revokeObjectURL(objectUrl),
                                        60000
                                    );
                                }}
                            }}, 650);
                        }};
                        frame.addEventListener('load', launchPrint, {{ once: true }});
                        frame.src = objectUrl;
                        parentDocument.body.appendChild(frame);
                        window.setTimeout(launchPrint, 1800);
                    }};

                    if (payload.mode === 'pdf') {{
                        printPdf(payload.base64, payload.mimeType);
                    }} else if (payload.mode === 'html') {{
                        printHtml(payload.html);
                    }}
                }})();
        </script>
    """
    render_browser_bridge(component)


def _dismiss_file_preview() -> None:
    """Clear preview state when X, outside click, or Escape dismisses it."""

    context = st.session_state.pop(_DISMISS_CONTEXT_KEY, None)
    if not isinstance(context, dict):
        return

    preview_state_key = str(context.get("preview_state_key", "")).strip()
    table_version_key = str(context.get("table_version_key", "")).strip()

    if preview_state_key:
        st.session_state[preview_state_key] = None
    if table_version_key:
        st.session_state[table_version_key] = (
            int(st.session_state.get(table_version_key, 0)) + 1
        )


def _read_text(data: bytes) -> str:
    """Decode text safely using common encodings."""

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _render_pdf(data: bytes) -> None:
    """Render PDF bytes with Streamlit's native in-app PDF viewer."""

    # The old PDF-in-an-iframe approach could render a blank white area in
    # Chromium browsers. ``st.pdf`` uses Streamlit's supported PDF component
    # and accepts authorized raw bytes directly.
    st.pdf(data, height=_PREVIEW_WIDGET_HEIGHT)


def _render_docx(data: bytes) -> None:
    """Render DOCX paragraphs and tables in a readable modal layout."""

    document = Document(BytesIO(data))
    has_content = False

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            has_content = True
            st.markdown(text)

    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        has_content = True
        st.caption(f"Table {table_index}")
        st.dataframe(
            pd.DataFrame(rows),
            hide_index=True,
            width="stretch",
            height=min(_PREVIEW_WIDGET_HEIGHT, 38 + (len(rows) * 35)),
        )

    if not has_content:
        st.info("The DOCX file does not contain previewable text or tables.")


def _render_spreadsheet(data: bytes, filename: str) -> None:
    """Render a selectable sheet from a modern Excel workbook."""

    workbook = pd.ExcelFile(BytesIO(data), engine="openpyxl")
    sheet_name = st.selectbox(
        "Worksheet",
        options=workbook.sheet_names,
        key=f"company_file_preview_sheet_{abs(hash(filename))}",
    )
    frame = workbook.parse(sheet_name=sheet_name, header=None)
    clipped = frame.iloc[:_MAX_PREVIEW_ROWS, :_MAX_PREVIEW_COLUMNS]
    st.dataframe(
        clipped,
        hide_index=True,
        width="stretch",
        height=_PREVIEW_WIDGET_HEIGHT,
    )
    if frame.shape != clipped.shape:
        st.caption(
            f"Preview is limited to {_MAX_PREVIEW_ROWS} rows and "
            f"{_MAX_PREVIEW_COLUMNS} columns. Download the file for the "
            "complete workbook."
        )


def _render_csv(data: bytes) -> None:
    """Render a CSV table with a safe preview size."""

    frame = pd.read_csv(BytesIO(data))
    clipped = frame.iloc[:_MAX_PREVIEW_ROWS, :_MAX_PREVIEW_COLUMNS]
    st.dataframe(
        clipped,
        hide_index=True,
        width="stretch",
        height=_PREVIEW_WIDGET_HEIGHT,
    )
    if frame.shape != clipped.shape:
        st.caption(
            f"Preview is limited to {_MAX_PREVIEW_ROWS} rows and "
            f"{_MAX_PREVIEW_COLUMNS} columns."
        )


def _render_image(data: bytes) -> None:
    """Render an uploaded image submission."""

    st.image(data, width="stretch")


def _render_preview(*, filename: str, mime_type: str, data: bytes) -> None:
    """Dispatch one authorized file to the matching preview renderer."""

    extension = Path(filename).suffix.lower()

    if extension == ".pdf" or mime_type == "application/pdf":
        _render_pdf(data)
        return
    if extension == ".docx":
        _render_docx(data)
        return
    if extension == ".xlsx":
        _render_spreadsheet(data, filename)
        return
    if extension == ".csv":
        _render_csv(data)
        return
    if extension == ".txt" or mime_type.startswith("text/"):
        st.text_area(
            "Text Preview",
            value=_read_text(data),
            height=_PREVIEW_WIDGET_HEIGHT,
            disabled=True,
            label_visibility="collapsed",
        )
        return
    if extension in {".png", ".jpg", ".jpeg"} or mime_type.startswith("image/"):
        _render_image(data)
        return

    if extension in {".doc", ".xls"}:
        st.info(
            "Legacy Word (.doc) and Excel (.xls) files cannot be rendered "
            "reliably in the browser. Download the file to open it in the "
            "appropriate desktop application."
        )
        return

    st.info(
        "A browser preview is not available for this file type. Use the "
        "download button below."
    )


@st.dialog(
    "File Preview",
    width="large",
    dismissible=True,
    on_dismiss=_dismiss_file_preview,
)
def render_file_preview_dialog(
    *,
    filename: str,
    mime_type: str,
    data: bytes,
    preview_state_key: str,
    table_version_key: str,
) -> None:
    """Show one authorized file in a balanced viewport-height modal."""

    st.session_state[_DISMISS_CONTEXT_KEY] = {
        "preview_state_key": preview_state_key,
        "table_version_key": table_version_key,
    }

    # Scope the modal treatment through the keyed preview container. The panel
    # uses equal 48-pixel viewport margins so its visible top and bottom space
    # stays balanced, while shorter screens keep an internally scrolling body.
    st.markdown(
        """
        <style>
            [data-testid="stDialog"]:has(
                div[class*="st-key-company_file_preview_content"]
            ) > div {
                height: calc(100vh - 96px) !important;
                height: calc(100dvh - 96px) !important;
                max-height: calc(100vh - 96px) !important;
                max-height: calc(100dvh - 96px) !important;
            }

            div[class*="st-key-company_file_preview_content"] {
                height: clamp(
                    360px,
                    calc(100dvh - 344px),
                    600px
                ) !important;
            }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.subheader(filename)

    with st.container(
        height=_PREVIEW_CONTENT_HEIGHT,
        border=True,
        key="company_file_preview_content",
    ):
        try:
            _render_preview(filename=filename, mime_type=mime_type, data=data)
        except Exception as error:  # A malformed document must not break the page.
            st.warning(
                "This file could not be rendered in the browser preview. "
                "Download it to open the complete file."
            )
            st.caption(f"Preview detail: {error}")

    print_payload = _build_print_payload(
        filename=filename,
        mime_type=mime_type,
        data=data,
    )
    print_column, download_column = st.columns(2)
    with print_column:
        print_requested = st.button(
            "Print Form",
            type="secondary",
            width="stretch",
            disabled=print_payload["mode"] == "unsupported",
            key=f"preview_print_{preview_state_key}_{abs(hash(filename))}",
        )
    with download_column:
        st.download_button(
            "Download File",
            data=data,
            file_name=filename,
            mime=mime_type,
            type="secondary",
            width="stretch",
            key=f"preview_download_{preview_state_key}_{abs(hash(filename))}",
        )

    if print_requested:
        _launch_print_dialog(print_payload)
