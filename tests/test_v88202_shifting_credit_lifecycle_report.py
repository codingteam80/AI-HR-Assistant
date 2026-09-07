"""v8.8.202 Shifting/OB lifecycle, restoration, attendance use, and report checks."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from config.settings import Settings
from database.base import Base
from models.attendance_record import AttendanceRecord
from models.shifting_credit import ShiftingCredit
from modules.reports.shifting_credits_report import build_shifting_credits_excel
from schemas.attendance_schema import AttendanceStatusInput
from schemas.overtime_schema import OvertimeRequestInput, OvertimeReviewInput
from scripts.create_initial_data import seed_initial_data
from services.attendance_service import AttendanceService
from services.overtime_service import OvertimeService


ROOT = Path(__file__).resolve().parents[1]


def _settings() -> Settings:
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        initial_company_code="V88202",
        initial_company_name="Shifting Lifecycle Company",
        initial_admin_username="admin",
        initial_admin_email="admin@example.com",
        initial_admin_password=SecretStr("Temporary123!"),
        initial_admin_employee_number="ADMIN-202",
    )


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _record(seed, rendered: date, hours: Decimal) -> AttendanceRecord:
    return AttendanceRecord(
        company_id=seed["company"].id,
        employee_id=seed["admin_employee"].id,
        attendance_date=rendered,
        time_in=datetime(rendered.year, rendered.month, rendered.day, 8, 0, tzinfo=timezone.utc),
        time_out=datetime(rendered.year, rendered.month, rendered.day, 17, 0, tzinfo=timezone.utc),
        work_status="WFO",
        scheduled_workday=True,
        regular_hours_target=Decimal("8.00"),
        lunch_break_minutes=60,
        total_hours=Decimal("8.00") + hours,
        ot_hours=hours,
    )


def _request(seed, rendered: date, hours: Decimal, *, dinner: bool = False):
    start = datetime(rendered.year, rendered.month, rendered.day, 17, 0)
    end = start + timedelta(hours=float(hours))
    return OvertimeRequestInput(
        company_id=seed["company"].id,
        employee_id=seed["admin_employee"].id,
        requested_by_user_id=seed["admin_user"].id,
        date_rendered=rendered,
        ot_time_start=start,
        ot_time_end=end,
        estimated_hours=hours,
        ot_type="Regular Overtime",
        ot_purpose="Lifecycle test",
        dinner_break_flag=dinner,
    )


def _approve(service: OvertimeService, seed, request):
    return service.review(
        OvertimeReviewInput(
            company_id=seed["company"].id,
            overtime_request_id=request.id,
            reviewed_by_user_id=seed["admin_user"].id,
            reviewer_employee_id=seed["admin_employee"].id,
            clearance=1,
            decision="approved",
        )
    )


def _pair(
    session,
    seed,
    *,
    first_hours=Decimal("4.00"),
    first_dinner=False,
    first_date=date(2026, 8, 24),
    second_date=date(2026, 8, 26),
):
    session.add_all([
        _record(seed, first_date, first_hours),
        _record(seed, second_date, Decimal("4.00")),
    ])
    session.commit()
    service = OvertimeService(session)
    first = service.submit(_request(seed, first_date, first_hours, dinner=first_dinner))
    _approve(service, seed, first)
    second = service.submit(_request(seed, second_date, Decimal("4.00")))
    _approve(service, seed, second)
    session.refresh(first)
    return service, first, second


def test_pair_creates_protected_pending_credit_with_two_cutoff_delay() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        _service, first, second = _pair(session, seed)
        credit = session.scalar(select(ShiftingCredit))
        assert credit is not None
        assert credit.status == "pending"
        assert credit.cutoff_start == date(2026, 8, 16)
        assert credit.cutoff_end == date(2026, 8, 31)
        assert credit.availability_date == date(2026, 10, 1)
        assert credit.expiration_date == date(2026, 12, 31)
        assert Decimal(credit.qualifying_hours) == Decimal("8.00")
        assert Decimal(first.payable_hours) == Decimal("0.00")
        assert Decimal(second.payable_hours) == Decimal("0.00")


def test_unused_expired_ob_restores_regular_ot_without_removing_consumed_marker() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service, first, second = _pair(
            session,
            seed,
            first_hours=Decimal("5.00"),
            first_dinner=True,
        )
        assert Decimal(first.payable_hours) == Decimal("0.25")
        assert Decimal(second.payable_hours) == Decimal("0.00")

        service.refresh_shifting_credit_statuses(
            company_id=seed["company"].id,
            as_of=date(2027, 1, 1),
        )
        session.flush()
        session.refresh(first)
        session.refresh(second)
        credit = session.scalar(select(ShiftingCredit))
        assert credit is not None and credit.status == "expired"
        assert Decimal(credit.restored_regular_ot_hours) == Decimal("8.00")
        assert Decimal(first.shifting_credit_hours) == Decimal("4.00")
        assert Decimal(second.shifting_credit_hours) == Decimal("4.00")
        assert Decimal(first.shifting_credit_restored_hours) == Decimal("4.00")
        assert Decimal(second.shifting_credit_restored_hours) == Decimal("4.00")
        # Dinner deduction remains applied after restoration: 5 - .75 = 4.25.
        assert Decimal(first.payable_hours) == Decimal("4.25")
        assert Decimal(second.payable_hours) == Decimal("4.00")


def test_available_ob_can_be_used_and_released_by_work_status() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service, _first, _second = _pair(
            session,
            seed,
            first_date=date(2026, 6, 2),
            second_date=date(2026, 6, 4),
        )
        service.refresh_shifting_credit_statuses(
            company_id=seed["company"].id,
            as_of=date(2026, 8, 28),
        )
        session.flush()
        credit = session.scalar(select(ShiftingCredit))
        assert credit is not None and credit.status == "available"

        attendance = AttendanceService(session)
        ob_date = date(2026, 8, 28)
        record = attendance.edit_own_work_status(
            AttendanceStatusInput(
                company_id=seed["company"].id,
                employee_id=seed["admin_employee"].id,
                user_id=seed["admin_user"].id,
                attendance_date=ob_date,
                work_status="OB",
            )
        )
        assert record.work_status == "OB"
        session.refresh(credit)
        assert credit.status == "used"
        assert credit.usage_date == ob_date

        attendance.edit_own_work_status(
            AttendanceStatusInput(
                company_id=seed["company"].id,
                employee_id=seed["admin_employee"].id,
                user_id=seed["admin_user"].id,
                attendance_date=ob_date,
                work_status="WFO",
            )
        )
        session.refresh(credit)
        assert credit.status == "available"
        assert credit.usage_date is None


def test_shifting_report_matches_required_workbook_structure() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service, first, second = _pair(
            session, seed, first_hours=Decimal("5.00")
        )
        service.refresh_shifting_credit_statuses(
            company_id=seed["company"].id,
            as_of=date(2026, 10, 1),
        )
        credit = session.scalar(select(ShiftingCredit))
        rules = service.rule_snapshot(seed["company"].id)
        data = build_shifting_credits_excel(
            employees=[seed["admin_employee"]],
            overtime_requests=[first, second],
            credits=[credit],
            report_year=2026,
            cutoff_day=rules.shifting_credit_cutoff_day,
            availability_cutoffs=rules.availability_cutoffs,
            expiration_mode=rules.expiration_mode,
            expiration_month=rules.expiration_month,
            expiration_day=rules.expiration_day,
            block_hours=rules.shifting_credit_block_hours,
            required_blocks=rules.shifting_credit_required_blocks,
            straight_ot_threshold_hours=rules.additional_vl_threshold_hours,
            straight_ot_vl_days=rules.additional_vl_days,
            straight_ot_also_payable=rules.additional_vl_also_payable,
            excluded_positions=rules.excluded_positions,
            timezone_name="Asia/Manila",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        assert workbook.sheetnames[0:2] == ["Summary", "Cut-Off Date"]
        assert len(workbook.sheetnames) == 3
        employee_sheet = workbook[workbook.sheetnames[2]]
        assert employee_sheet["A3"].value == "Date"
        assert employee_sheet["D4"].value == "Qualifying Block Hrs"
        assert employee_sheet["E4"].value == "OB Credit (Days)"
        assert employee_sheet["F3"].value == "Straight 8h OT"
        assert employee_sheet["G3"].value == "Availability Date"
        assert employee_sheet["H3"].value == "Expiration Date"
        assert employee_sheet["J3"].value == "Date of Usage"
        assert employee_sheet["L3"].value == "Regular OT Excess"
        assert employee_sheet["M3"].value == "VL Credit"


def test_v88202_version_markers() -> None:
    assert 'app_version: str = "0.8.8.202"' in (ROOT / "config" / "settings.py").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.202" in (ROOT / ".env").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.202" in (ROOT / ".env.example").read_text(encoding="utf-8")
