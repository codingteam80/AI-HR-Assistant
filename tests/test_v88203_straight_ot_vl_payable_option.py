"""v8.8.203 configurable straight-OT VL + optional OT-payable behavior."""

from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from config.settings import Settings
from database.base import Base
from database.schema_upgrade import upgrade_existing_schema
from models.attendance_record import AttendanceRecord
from modules.reports.shifting_credits_report import build_shifting_credits_excel
from schemas.attendance_schema import CompanyOvertimeRulesInput
from schemas.overtime_schema import OvertimeRequestInput, OvertimeReviewInput
from scripts.create_initial_data import seed_initial_data
from services.attendance_service import AttendanceService
from services.overtime_service import OvertimeService

ROOT = Path(__file__).resolve().parents[1]


def _settings() -> Settings:
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        initial_company_code="V88203",
        initial_company_name="Straight OT Flex Company",
        initial_admin_username="admin",
        initial_admin_email="admin@example.com",
        initial_admin_password=SecretStr("Temporary123!"),
        initial_admin_employee_number="ADMIN-203",
    )


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _record(seed, rendered: date, hours: Decimal) -> AttendanceRecord:
    return AttendanceRecord(
        company_id=seed["company"].id,
        employee_id=seed["admin_employee"].id,
        attendance_date=rendered,
        time_in=datetime(rendered.year, rendered.month, rendered.day, 8, 0, tzinfo=timezone.utc),
        time_out=datetime(rendered.year, rendered.month, rendered.day, 17, 0, tzinfo=timezone.utc),
        work_status="WFO",
        scheduled_workday=True,
        regular_hours_target=Decimal("8.00"),
        lunch_break_minutes=60,
        total_hours=Decimal("8.00") + hours,
        ot_hours=hours,
    )


def _approve_hours(session, seed, rendered: date, hours: Decimal, *, dinner: bool = False):
    session.add(_record(seed, rendered, hours))
    session.commit()
    start = datetime(rendered.year, rendered.month, rendered.day, 17, 0)
    request = OvertimeService(session).submit(
        OvertimeRequestInput(
            company_id=seed["company"].id,
            employee_id=seed["admin_employee"].id,
            requested_by_user_id=seed["admin_user"].id,
            date_rendered=rendered,
            ot_time_start=start,
            ot_time_end=start + timedelta(hours=float(hours)),
            estimated_hours=hours,
            ot_type="Regular Overtime",
            ot_purpose="v8.8.203 rule test",
            dinner_break_flag=dinner,
        )
    )
    return OvertimeService(session).review(
        OvertimeReviewInput(
            company_id=seed["company"].id,
            overtime_request_id=request.id,
            reviewed_by_user_id=seed["admin_user"].id,
            reviewer_employee_id=seed["admin_employee"].id,
            clearance=1,
            decision="approved",
        )
    )


def test_default_is_conversion_only_and_nine_hours_keeps_only_excess() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        assert seed["company"].shifting_credit_additional_vl_also_payable is False
        request = _approve_hours(session, seed, date(2026, 8, 28), Decimal("9.00"))
        assert Decimal(request.additional_vl_days) == Decimal("0.50")
        assert request.straight_vl_also_payable is False
        assert Decimal(request.payable_hours) == Decimal("1.00")


def test_default_conversion_with_dinner_applies_deduction_only_to_payable_excess() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        request = _approve_hours(
            session, seed, date(2026, 8, 28), Decimal("9.00"), dinner=True
        )
        assert Decimal(request.additional_vl_days) == Decimal("0.50")
        assert Decimal(request.payable_hours) == Decimal("0.25")


def test_checkbox_on_grants_vl_and_keeps_ot_payable() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        AttendanceService(session).save_overtime_rules(
            CompanyOvertimeRulesInput(
                company_id=seed["company"].id,
                dinner_break_deduction_hours=Decimal("0.75"),
                shifting_credits_enabled=True,
                shifting_credit_block_hours=Decimal("4.00"),
                shifting_credit_required_blocks=2,
                shifting_credit_cutoff_day=15,
                additional_vl_threshold_hours=Decimal("8.00"),
                additional_vl_days=Decimal("0.50"),
                additional_vl_also_payable=True,
                excluded_positions=["Trainee", "Design Engineer I", "Design Engineer II"],
            )
        )
        request = _approve_hours(
            session, seed, date(2026, 8, 28), Decimal("8.00"), dinner=True
        )
        assert Decimal(request.additional_vl_days) == Decimal("0.50")
        assert request.straight_vl_also_payable is True
        assert Decimal(request.payable_hours) == Decimal("7.25")


def test_report_marks_optional_straight_vl_plus_ot_payable() -> None:
    factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        AttendanceService(session).save_overtime_rules(
            CompanyOvertimeRulesInput(
                company_id=seed["company"].id,
                dinner_break_deduction_hours=Decimal("0.75"),
                shifting_credits_enabled=True,
                shifting_credit_block_hours=Decimal("4.00"),
                shifting_credit_required_blocks=2,
                shifting_credit_cutoff_day=15,
                additional_vl_threshold_hours=Decimal("8.00"),
                additional_vl_days=Decimal("0.50"),
                additional_vl_also_payable=True,
                excluded_positions=["Trainee", "Design Engineer I", "Design Engineer II"],
            )
        )
        request = _approve_hours(
            session, seed, date(2026, 8, 28), Decimal("8.00"), dinner=True
        )
        rules = OvertimeService(session).rule_snapshot(seed["company"].id)
        data = build_shifting_credits_excel(
            employees=[seed["admin_employee"]],
            overtime_requests=[request],
            credits=[],
            report_year=2026,
            cutoff_day=rules.shifting_credit_cutoff_day,
            availability_cutoffs=rules.availability_cutoffs,
            expiration_mode=rules.expiration_mode,
            expiration_month=rules.expiration_month,
            expiration_day=rules.expiration_day,
            block_hours=rules.shifting_credit_block_hours,
            required_blocks=rules.shifting_credit_required_blocks,
            straight_ot_threshold_hours=rules.additional_vl_threshold_hours,
            straight_ot_vl_days=rules.additional_vl_days,
            straight_ot_also_payable=rules.additional_vl_also_payable,
            excluded_positions=rules.excluded_positions,
            timezone_name="Asia/Manila",
        )
        workbook = load_workbook(BytesIO(data), data_only=False)
        employee_sheet = workbook[workbook.sheetnames[2]]
        assert employee_sheet["K5"].value == "Straight OT -> 0.50 VL + OT payable"
        assert Decimal(str(employee_sheet["L5"].value)) == Decimal("7.25")
        assert Decimal(str(employee_sheet["M5"].value)) == Decimal("0.5")


def test_schema_upgrade_preserves_historical_straight_vl_payable_behavior(tmp_path) -> None:
    database_path = tmp_path / "legacy_v88202.db"
    engine = create_engine(f"sqlite+pysqlite:///{database_path}")
    with engine.begin() as connection:
        connection.execute(text(
            "CREATE TABLE overtime_requests ("
            "id INTEGER PRIMARY KEY, estimated_hours NUMERIC(8,2) NOT NULL, "
            "additional_vl_days NUMERIC(8,2) NOT NULL DEFAULT 0)"
        ))
        connection.execute(text(
            "INSERT INTO overtime_requests (id, estimated_hours, additional_vl_days) "
            "VALUES (1, 8.00, 0.50), (2, 5.00, 0.00)"
        ))
    upgrade_existing_schema(engine)
    columns = {item["name"] for item in inspect(engine).get_columns("overtime_requests")}
    assert "straight_vl_also_payable" in columns
    with engine.connect() as connection:
        rows = connection.execute(text(
            "SELECT id, payable_hours, straight_vl_also_payable "
            "FROM overtime_requests ORDER BY id"
        )).all()
    assert Decimal(str(rows[0].payable_hours)) == Decimal("8")
    assert bool(rows[0].straight_vl_also_payable) is True
    assert Decimal(str(rows[1].payable_hours)) == Decimal("5")
    assert bool(rows[1].straight_vl_also_payable) is False


def test_v88203_ui_schema_and_version_markers() -> None:
    admin_source = (ROOT / "ui" / "pages" / "admin" / "attendance_dashboard.py").read_text(encoding="utf-8")
    schema_source = (ROOT / "database" / "schema_upgrade.py").read_text(encoding="utf-8")
    runtime_source = (ROOT / "database" / "runtime_schema.py").read_text(encoding="utf-8")
    assert '"Consider it also as OT Payable"' in admin_source
    assert 'value=additional_vl_also_payable' in admin_source
    assert 'key=f"shifting_additional_vl_also_payable_{current_user.company_id}"' in admin_source
    assert '"additional_vl_also_payable": additional_vl_also_payable_value' in admin_source
    assert '"shifting_credit_additional_vl_also_payable": "BOOLEAN NOT NULL DEFAULT 0"' in schema_source
    assert '"straight_vl_also_payable": "BOOLEAN NOT NULL DEFAULT 0"' in schema_source
    assert '"shifting_credit_additional_vl_also_payable"' in runtime_source
    assert '"straight_vl_also_payable"' in runtime_source
    assert 'app_version: str = "0.8.8.203"' in (ROOT / "config" / "settings.py").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.203" in (ROOT / ".env").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.203" in (ROOT / ".env.example").read_text(encoding="utf-8")
