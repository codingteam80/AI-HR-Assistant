"""v8.8.178 Reminder Excel bulk import regression checks."""

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from config.settings import Settings
from database.base import Base
from schemas.event_reminder_schema import EVENT_REMINDER_CATEGORIES
from scripts.create_initial_data import seed_initial_data
from services.event_reminder_bulk_import_service import (
    EVENT_REMINDER_IMPORT_COLUMNS,
    EventReminderBulkImportService,
)
from services.event_reminder_service import EventReminderService


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _settings(tmp_path: Path) -> Settings:
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        initial_company_code="REMINDERXLSX",
        initial_company_name="Reminder Excel Company",
        initial_admin_username="admin",
        initial_admin_email="admin@reminder.xlsx",
        initial_admin_password=SecretStr("Temporary123!"),
        initial_admin_employee_number="ADMIN-001",
        initial_admin_first_name="System",
        initial_admin_last_name="Administrator",
        announcement_upload_dir=str(tmp_path / "announcement_images"),
        display_timezone="Asia/Manila",
    )


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _completed_template(rows: list[tuple[object, ...]]) -> bytes:
    workbook = load_workbook(BytesIO(EventReminderBulkImportService.build_template()))
    sheet = workbook["Reminder Import Template"]
    sheet.delete_rows(2, max(1, sheet.max_row - 1))
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_template_matches_create_reminder_fields_and_hides_category_list_in_comment() -> None:
    workbook = load_workbook(BytesIO(EventReminderBulkImportService.build_template()))
    sheet = workbook["Reminder Import Template"]

    assert tuple(cell.value for cell in sheet[1]) == EVENT_REMINDER_IMPORT_COLUMNS
    assert EVENT_REMINDER_IMPORT_COLUMNS == (
        "Category",
        "Event Date",
        "Event / Activity Title",
        "Preparation Notes",
    )
    assert sheet["A1"].comment is not None
    comment_text = sheet["A1"].comment.text
    for category in EVENT_REMINDER_CATEGORIES:
        assert category in comment_text

    # The allowed category list is intentionally guidance inside the collapsed
    # Excel comment rather than extra visible rows/columns in the template.
    visible_values = {
        str(cell.value)
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    }
    assert "Allowed Category values:" not in visible_values


def test_preview_normalizes_current_category_and_accepts_excel_date(tmp_path: Path) -> None:
    factory = _factory()
    with factory() as session:
        settings = _settings(tmp_path)
        seed = seed_initial_data(session, settings)
        service = EventReminderBulkImportService(session, settings=settings)
        file_bytes = _completed_template(
            [
                (
                    "company event",
                    date(2027, 3, 15),
                    "Annual Company Event",
                    "Prepare the employee announcement.",
                )
            ]
        )

        preview = service.prepare_preview(
            file_bytes,
            filename="reminders.xlsx",
            company_id=seed["company"].id,
        )

        assert len(preview) == 1
        assert preview[0]["Category"] == "Company Event"
        assert preview[0]["Event Date"] == "2027/03/15"
        assert preview[0]["Validation"] == "Ready"


def test_invalid_category_is_flagged_before_import(tmp_path: Path) -> None:
    factory = _factory()
    with factory() as session:
        settings = _settings(tmp_path)
        seed = seed_initial_data(session, settings)
        service = EventReminderBulkImportService(session, settings=settings)
        preview = service.prepare_preview(
            _completed_template(
                [
                    (
                        "Not A Current Category",
                        "2027/04/01",
                        "Invalid Category Reminder",
                        "This row must fail validation.",
                    )
                ]
            ),
            filename="reminders.xlsx",
            company_id=seed["company"].id,
        )

        assert preview[0]["Validation"] != "Ready"
        assert "current allowed reminder categories" in preview[0]["Validation"]


def test_valid_excel_batch_import_is_atomic_and_uses_existing_schedule(tmp_path: Path) -> None:
    factory = _factory()
    with factory() as session:
        settings = _settings(tmp_path)
        seed = seed_initial_data(session, settings)
        bulk_service = EventReminderBulkImportService(session, settings=settings)
        preview = bulk_service.prepare_preview(
            _completed_template(
                [
                    (
                        "Holiday / Observance",
                        "2027/06/12",
                        "Araw ng Kalayaan",
                        "Prepare the employee announcement.",
                    ),
                    (
                        "Training / Seminar",
                        "2027/07/15",
                        "HR Compliance Seminar",
                        "Confirm facilitator and participant list.",
                    ),
                ]
            ),
            filename="reminders.xlsx",
            company_id=seed["company"].id,
        )

        created = bulk_service.import_preview_rows(
            preview,
            company_id=seed["company"].id,
            actor_user_id=seed["admin_user"].id,
        )

        assert len(created) == 2
        assert [item.public_id for item in created] == ["REM_000001", "REM_000002"]
        assert all(item.status == "planned" for item in created)
        assert all(item.event_start_at.hour == 1 for item in created)  # 9 AM Asia/Manila -> 1 AM UTC
        assert all(item.reminder_at is not None for item in created)
        assert len(EventReminderService(session, settings=settings).list_for_admin(seed["company"].id)) == 2


def test_create_reminder_ui_contains_excel_bulk_flow_without_replacing_manual_flow() -> None:
    source = (PROJECT_ROOT / "ui/pages/admin/announcements_page.py").read_text(
        encoding="utf-8"
    )
    service_source = (
        PROJECT_ROOT / "services/event_reminder_bulk_import_service.py"
    ).read_text(encoding="utf-8")

    assert '"Upload Reminders via Excel"' in source
    assert '"Download Reminder Excel Template"' in source
    assert '"Validate and Preview Reminders"' in source
    assert '"Import Reminders"' in source
    assert 'Create Reminder Manually' in source
    assert '"Save Smart Reminders"' in source
    assert "EventReminderBulkImportService" in source
    assert "Comment(" in service_source
    assert "EVENT_REMINDER_CATEGORIES" in service_source
    assert ".create_many(" in service_source


def test_version_is_v88178() -> None:
    settings_source = (PROJECT_ROOT / "config/settings.py").read_text(encoding="utf-8")
    env_source = (PROJECT_ROOT / ".env").read_text(encoding="utf-8")
    example_source = (PROJECT_ROOT / ".env.example").read_text(encoding="utf-8")

    assert 'app_version: str = "0.8.8.178"' in settings_source
    assert "APP_VERSION=0.8.8.178" in env_source
    assert "APP_VERSION=0.8.8.178" in example_source
