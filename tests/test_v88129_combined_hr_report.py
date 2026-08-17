"""v8.8.129 combined DTR, OT, and leave workbook regression checks."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from modules.reports.combined_hr_report import (
    _leave_type_export_code,
    _leave_type_export_value,
    _report_employee_name,
    build_combined_hr_excel,
)
from ui.pages.admin.reports_page import _selected_employee_ids


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _source(relative_path: str) -> str:
    return (PROJECT_ROOT / relative_path).read_text(encoding="utf-8")


def _employee(
    employee_id: int,
    number: str,
    first: str,
    last: str,
    *,
    middle: str | None = None,
    suffix: str | None = None,
):
    employee = SimpleNamespace(
        id=employee_id,
        employee_number=number,
        first_name=first,
        last_name=last,
        middle_name=middle,
        suffix=suffix,
        department=SimpleNamespace(name="SPI"),
    )
    employee.full_name = f"{first} {last}"
    return employee


def _attendance(employee, work_date: date, **values):
    defaults = {
        "employee_id": employee.id,
        "employee": employee,
        "attendance_date": work_date,
        "time_in": None,
        "time_out": None,
        "work_status": None,
        "leave_request_id": None,
        "scheduled_workday": True,
        "regular_hours_target": Decimal("8.00"),
        "lunch_break_minutes": 60,
        "total_hours": Decimal("0.00"),
        "ot_hours": Decimal("0.00"),
    }
    defaults.update(values)
    return SimpleNamespace(**defaults)


def _leave(employee, code: str, start_date: date):
    return SimpleNamespace(
        employee=employee,
        employee_id=employee.id,
        leave_type=SimpleNamespace(code=code, name=f"{code} Leave"),
        start_date=start_date,
        end_date=start_date,
        requested_days=Decimal("1.00"),
        reason=f"Approved {code} reason",
        status="completed",
        submitted_at=datetime(2026, 8, 9, 1, tzinfo=timezone.utc),
        leader_approver=SimpleNamespace(
            first_name="Pre",
            middle_name="Leader",
            last_name="Approver",
            suffix=None,
            full_name="Pre Leader Approver",
        ),
        leader_reviewed_at=datetime(2026, 8, 10, 1, tzinfo=timezone.utc),
        leader_comment="Checked",
        manager=SimpleNamespace(
            first_name="Final",
            middle_name="Manager",
            last_name="Approver",
            suffix="Sr.",
            full_name="Final Manager Approver Sr.",
        ),
        reviewed_at=datetime(2026, 8, 11, 1, tzinfo=timezone.utc),
        approved_at=datetime(2026, 8, 11, 1, tzinfo=timezone.utc),
        manager_comment="Approved",
        cancellation_comment=None,
    )


def _workbook():
    employee = _employee(
        1,
        "EMP-001",
        "Lander",
        "Garcia",
        middle="Miguel",
        suffix="Jr.",
    )
    dtr_day = date(2026, 8, 11)
    leave_day = date(2026, 8, 12)
    records = [
        _attendance(
            employee,
            dtr_day,
            time_in=datetime(2026, 8, 11, 0, tzinfo=timezone.utc),
            time_out=datetime(2026, 8, 11, 10, tzinfo=timezone.utc),
            work_status="WFO",
            total_hours=Decimal("9.00"),
            ot_hours=Decimal("1.00"),
        ),
        _attendance(
            employee,
            leave_day,
            work_status="VL",
            leave_request_id=1,
        ),
    ]
    leave_requests = [
        _leave(employee, code, date(2026, 8, 12 + index))
        for index, code in enumerate(("VL", "SL", "EL"))
    ]
    data = build_combined_hr_excel(
        employees=[employee],
        attendance_records=records,
        leave_requests=leave_requests,
        start_date=dtr_day,
        end_date=date(2026, 8, 14),
        scheduled_workdays={
            date(2026, 8, 11): True,
            date(2026, 8, 12): True,
            date(2026, 8, 13): True,
            date(2026, 8, 14): True,
        },
        timezone_name="Asia/Manila",
    )
    return load_workbook(BytesIO(data))


def test_combined_workbook_contains_exact_three_sheets() -> None:
    workbook = _workbook()

    assert workbook.sheetnames == ["DTR Logs", "Overtime File", "Leave File"]


def test_dtr_uses_two_employee_rows_dates_times_and_leave_codes() -> None:
    sheet = _workbook()["DTR Logs"]

    assert [sheet.cell(1, index).value for index in range(1, 4)] == [
        "Name",
        "Section",
        "Property",
    ]
    assert sheet["C2"].value == "Start Time"
    assert sheet["C3"].value == "End Time"
    assert sheet["A2"].value == "Garcia, Lander Jr., M."
    assert sheet["A3"].value == "Garcia, Lander Jr., M."
    assert sheet["D2"].value.hour == 8
    assert sheet["D3"].value.hour == 18
    assert sheet["E2"].value == "VL"
    assert sheet["E3"].value == "VL"
    assert sheet["A1"].fill.fgColor.rgb.endswith("FFFFFF")
    assert sheet["E2"].fill.fgColor.rgb.endswith("FFCCFF")
    assert sheet["D1"].number_format == "yyyy/m/d"
    assert sheet.auto_filter.ref is None
    assert sheet.freeze_panes == "D2"


def test_ot_sheet_exports_only_computed_ot_without_invented_fields() -> None:
    sheet = _workbook()["Overtime File"]

    assert sheet.max_row == 2
    assert sheet["A2"].value == "EMP-001"
    assert [sheet.cell(1, index).value for index in range(1, 15)] == [
        "emp_id",
        "emp_name",
        "date_rendered",
        "OT_time_start",
        None,
        "OT_time_end",
        None,
        "estimated_hours",
        "OT_type",
        "OT_purpose",
        "travel_fare",
        "travel_route",
        "dinner_break_flag",
        None,
    ]
    assert sheet["B2"].value == "Garcia, Lander Jr., M."
    assert sheet["D2"].value is None
    assert sheet["E2"].value is None
    assert sheet["F2"].value is None
    assert sheet["G2"].value is None
    assert sheet["H2"].value == 1
    assert sheet["I2"].value is None
    assert sheet["J2"].value in (None, "")
    assert sheet["K2"].value is None
    assert sheet["L2"].value in (None, "")
    assert sheet["M2"].value is None
    assert sheet["N2"].value is None
    assert sheet["D1"].fill.fgColor.rgb.endswith("92D050")
    assert sheet["H1"].fill.fgColor.rgb.endswith("92D050")
    assert sheet["A1"].fill.fgColor.rgb.endswith("9DC3E6")
    assert sheet.auto_filter.ref == "A1:N2"


def test_leave_sheet_combines_vl_sl_and_el_with_approval_details() -> None:
    sheet = _workbook()["Leave File"]

    assert [sheet.cell(1, index).value for index in range(1, 15)] == [
        "Emp ID",
        "Emp Name",
        "Leave Type",
        "Date Filed",
        "Start Date",
        "End Date",
        "Duration",
        "Reason for Leave",
        "Reason for Leave: Others",
        "Pre-approver",
        "Date Pre-approved",
        "Approver",
        "Date Approved",
        "Remarks",
    ]
    assert sheet["B2"].value == "Garcia, Lander Jr., M."
    assert {sheet.cell(row, 3).value for row in range(2, 5)} == {
        "90401 - Vacation",
        "90402 - Sick",
        "90410 - Emergency",
    }
    assert sheet["D2"].value.date() == datetime(2026, 8, 9).date()
    assert sheet["G2"].value == "90503 - Whole Day"
    assert sheet["H2"].value == "0 - OTHERS"
    assert sheet["I2"].value == "Approved VL reason"
    assert sheet["J2"].value == "Approver, Pre, L."
    assert sheet["L2"].value == "Approver, Final Sr., M."
    assert sheet["N2"].value == "Pre-approver: Checked | Approver: Approved"
    assert sheet["A1"].fill.fgColor.rgb.endswith("FFF59D")
    for cell in ("D2", "E2", "F2", "K2", "M2"):
        assert sheet[cell].number_format == "yyyy-mm-dd"
    assert sheet.auto_filter.ref == "A1:N4"


def test_leave_type_export_uses_requested_payroll_codes_and_descriptions() -> None:
    expected = {
        "VACATION": ("90401", "90401 - Vacation"),
        "SICK": ("90402", "90402 - Sick"),
        "BEREAVEMENT": ("90404", "90404 - Bereavement"),
        "HONEYMOON": ("90405", "90405 - Honeymoon"),
        "PATERNITY": ("90406", "90406 - Paternity"),
        "MATERNITY": ("90407", "90407 - Maternity"),
        "BIRTHDAY": ("90409", "90409 - Birthday"),
        "EMERGENCY": ("90410", "90410 - Emergency"),
    }
    for code, (export_code, export_value) in expected.items():
        request = SimpleNamespace(
            leave_type=SimpleNamespace(code=code, name=f"{code} Leave")
        )
        assert _leave_type_export_code(request) == export_code
        assert _leave_type_export_value(request) == export_value


def test_admin_reports_page_is_company_scoped_and_routed() -> None:
    page = _source("ui/pages/admin/reports_page.py")
    layout = _source("ui/layouts/admin_layout.py")

    assert "company_id=current_user.company_id" in page
    assert ".where(Department.company_id == current_user.company_id)" in page
    assert "LeaveRequestRepository(session).list_company(" in page
    assert 'f"{item.employee_number} · {item.full_name}"' in page
    assert "employee_id=employee_options[selected_employee_label]" in page
    assert 'elif page == "Reports":' in layout
    assert "render_admin_reports_page(current_user)" in layout


def test_duplicate_employee_names_are_filtered_by_unique_id() -> None:
    first = _employee(1, "EMP-001", "Same", "Name")
    second = _employee(2, "EMP-002", "Same", "Name")

    assert _selected_employee_ids(
        employees=[first, second],
        employee_id=second.id,
        department_name="All Departments",
    ) == [2]


def test_report_name_omits_unavailable_optional_name_parts() -> None:
    employee = _employee(1, "EMP-001", "Lander", "Garcia")

    assert _report_employee_name(employee) == "Garcia, Lander"


def test_overtime_request_checkpoint_is_based_directly_on_v88133() -> None:
    settings = _source("config/settings.py")
    readme = _source("README.md")

    assert 'app_version: str = "0.8.8.134"' in settings
    assert "Immediate base checkpoint: 0.8.8.133" in settings
    assert "v8.8.129 — Combined DTR, OT, and Leave Workbook" in readme
    assert "v8.8.130 — Reference-Matched Combined Report Templates" in readme
    assert "v8.8.131 — Payroll-Compatible Report Field Mapping" in readme
    assert "v8.8.132 — Middle-Initial Report Name Format" in readme
    assert (
        "v8.8.133 — Integrated Leave Duration, DTR Sessions, and Payroll Reports"
        in readme
    )
    assert "v8.8.134 — Combined Employee Overtime Request" in readme
