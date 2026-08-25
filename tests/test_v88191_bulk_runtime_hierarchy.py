"""v8.8.191 bulk violations, runtime warnings, and hierarchy awareness."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib import error as urllib_error

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session
from sqlalchemy import select

import models  # noqa: F401
from authentication.current_user import AuthenticatedUser
from database.base import Base
from models.audit_event import AuditEvent
from models.company import Company
from models.employee import Employee
from models.role import Role
from models.user import User
from modules.hr_assistant.admin_hr_assistant import AdminHRAssistant
from modules.hr_assistant.hr_assistant import HRAssistant
from modules.smart_ai.portal_ai import SmartPortalAssistant
from services.policy_violation_bulk_import_service import (
    PolicyViolationBulkImportService,
    VIOLATION_IMPORT_COLUMNS,
)
from services.policy_violation_service import PolicyViolationService
from services.audit_context import clear_audit_actor, set_audit_actor
from services.runtime_connection_service import (
    RuntimeServiceUnavailableError,
    classify_runtime_connection_issue,
)


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def _seed(session: Session):
    company = Company(code="TEST", name="TEST Company", is_active=True)
    session.add(company)
    session.flush()
    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role])
    session.flush()
    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username="admin_test",
        email="admin@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    manager_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username="manager_test",
        email="manager@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    member_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username="member_test",
        email="member@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add_all([admin, manager_user, member_user])
    session.flush()
    supervisor = Employee(
        company_id=company.id,
        user_id=manager_user.id,
        employee_number="SUP-001",
        first_name="Maria",
        last_name="Santos",
        job_title="Senior Specialist",
        employment_status="employed",
    )
    member = Employee(
        company_id=company.id,
        user_id=member_user.id,
        employee_number="MEM-001",
        first_name="Ana",
        last_name="Reyes",
        job_title="Team Lead",  # title must not override actual assignment data
        employment_status="employed",
    )
    session.add_all([supervisor, member])
    session.flush()
    member.manager_id = supervisor.id
    member.leader_id = supervisor.id
    session.commit()
    return company, admin, manager_user, member_user, supervisor, member


def _current_user(company, user, employee=None, *, clearance: int):
    return AuthenticatedUser(
        user_id=user.id,
        company_id=company.id,
        company_code=company.code,
        company_name=company.name,
        role_id=user.role_id,
        role_name="company_admin" if clearance == 1 else "employee",
        clearance=clearance,
        username=user.username,
        email=user.email,
        employee_id=employee.id if employee else None,
        employee_number=employee.employee_number if employee else None,
        employee_name=employee.full_name if employee else None,
        must_change_password=False,
    )


def test_version_and_ui_markers() -> None:
    assert 'app_version: str = "0.8.8.191"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.191" in _source(".env")
    page = _source("ui/pages/admin/policy_violations.py")
    assert "Bulk Add Violations via Excel" in page
    assert "Download Violation Excel Template" in page
    assert "Validate and Preview Violations" in page
    assert "Import Violations" in page
    assert "policy_violation_bulk_digest" in page
    assert "uploaded_digest != stored_digest" in page
    app = _source("app.py")
    assert "classify_runtime_connection_issue" in app
    assert "Retry Connection" in app


def test_violation_template_and_duplicate_validation_are_company_safe() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, *_ = _seed(session)
        service = PolicyViolationBulkImportService(session)
        payload = service.build_template(company_id=company.id)
        workbook = load_workbook(BytesIO(payload))
        assert "Violation Import Template" in workbook.sheetnames
        assert "Related Policies" in workbook.sheetnames
        assert "Instructions" in workbook.sheetnames
        headers = tuple(cell.value for cell in workbook["Violation Import Template"][1])
        assert headers == VIOLATION_IMPORT_COLUMNS

        sheet = workbook["Violation Import Template"]
        sheet.delete_rows(2, 1)
        row = [
            "ATT-001", "Attendance", "Habitual Tardiness", "Minor",
            "Repeated late arrival.", "Verbal Warning", "Written Warning",
            "Suspension", "Termination", "", "2026-08-24", "Active", "",
        ]
        sheet.append(row)
        sheet.append(row)
        out = BytesIO()
        workbook.save(out)
        preview = service.prepare_preview(
            out.getvalue(), filename="violations.xlsx", company_id=company.id
        )
        assert len(preview) == 2
        assert all("duplicated inside the uploaded file" in str(item["Validation"]) for item in preview)

        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Violation Import Template"]
        sheet.delete_rows(2, 1)
        sheet.append(row)
        out = BytesIO()
        workbook.save(out)
        preview = service.prepare_preview(
            out.getvalue(), filename="violations.xlsx", company_id=company.id
        )
        assert preview[0]["Validation"] == "Ready"
        created = service.import_preview_rows(
            preview, company_id=company.id, actor_user_id=admin.id
        )
        assert [item.violation_code for item in created] == ["ATT-001"]
        assert len(PolicyViolationService(session).list_current(company.id)) == 1

        again = service.prepare_preview(
            out.getvalue(), filename="violations.xlsx", company_id=company.id
        )
        assert "already exists in this company" in str(again[0]["Validation"])
    engine.dispose()


def test_runtime_connection_classifier_distinguishes_ollama_database_and_network() -> None:
    ollama_error = RuntimeServiceUnavailableError(
        "ollama", cause=urllib_error.URLError("connection refused")
    )
    ollama = classify_runtime_connection_issue(ollama_error, service_hint="ollama")
    assert ollama is not None
    assert ollama.code == "ollama_unavailable"
    assert "Ollama" in ollama.message
    assert "connection refused" not in ollama.message

    database = classify_runtime_connection_issue(
        OperationalError("select 1", {}, Exception("db secret")),
        service_hint="database",
    )
    assert database is not None
    assert database.code == "database_unavailable"
    assert "db secret" not in database.message

    network = classify_runtime_connection_issue(urllib_error.URLError("offline"))
    assert network is not None
    assert network.code == "network_unavailable"
    assert "offline" not in network.message

    # An HTTP 4xx/5xx proves the provider was reachable; it must not be
    # mislabeled as an internet/network outage.
    provider_reject = urllib_error.HTTPError(
        "https://provider.invalid", 401, "Unauthorized", hdrs=None, fp=None
    )
    assert classify_runtime_connection_issue(
        provider_reject, service_hint="sms"
    ) is None

    integrity = IntegrityError("insert", {}, Exception("duplicate"))
    assert classify_runtime_connection_issue(integrity) is None


def test_admin_hierarchy_distinguishes_manager_and_leader_from_actual_assignments() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, _manager_user, _member_user, supervisor, member = _seed(session)
        current = _current_user(company, admin, clearance=1)
        answer = AdminHRAssistant(session).answer(
            current_user=current,
            question=f"Who are the members of {supervisor.employee_number}?",
        )
        assert answer.intent == "employee_hierarchy"
        assert "Manager and Leader" in answer.answer
        assert member.employee_number in answer.answer
        assert "As Manager" in answer.answer
        assert "As Leader" in answer.answer

        # The member's job title says Team Lead, but there are no assigned members.
        member_answer = AdminHRAssistant(session).answer(
            current_user=current,
            question=f"Is {member.employee_number} a leader or manager?",
        )
        assert member_answer.intent == "employee_hierarchy"
        assert "No supervisory assignment" in member_answer.answer
    engine.dispose()


def test_employee_hierarchy_is_self_scoped_and_reports_own_assignments() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, _admin, manager_user, member_user, supervisor, member = _seed(session)
        manager_current = _current_user(
            company, manager_user, supervisor, clearance=2
        )
        answer = HRAssistant(session).answer(
            current_user=manager_current,
            question="Who are my members and am I a leader or manager?",
        )
        assert answer.intent == "employee_hierarchy"
        assert "Manager and Leader" in answer.answer
        assert member.employee_number in answer.answer

        member_current = _current_user(company, member_user, member, clearance=2)
        member_answer = HRAssistant(session).answer(
            current_user=member_current,
            question="Do I have members?",
        )
        assert member_answer.intent == "employee_hierarchy"
        assert "No supervisory assignment" in member_answer.answer
        assert supervisor.employee_number not in member_answer.answer

        third_party = HRAssistant(session).answer(
            current_user=member_current,
            question=f"Who are the members of {supervisor.employee_number}?",
        )
        assert third_party.intent == "employee_hierarchy"
        assert "only your own" in third_party.answer
        assert member.employee_number not in third_party.answer
    engine.dispose()


def test_chat_preflight_returns_popup_safe_ollama_issue_without_losing_direct_answer(monkeypatch) -> None:
    engine = _engine()
    with Session(engine) as session:
        assistant = SmartPortalAssistant(session)

        def fail_health_check() -> None:
            raise RuntimeServiceUnavailableError(
                "ollama",
                cause=urllib_error.URLError("refused private detail"),
            )

        monkeypatch.setattr(assistant.ollama, "check_available", fail_health_check)
        issue = assistant.preflight_connection_issue(role_scope="employee")
        assert issue is not None
        assert issue.code == "ollama_unavailable"
        assert "Ollama" in issue.message
        assert "private detail" not in issue.message
    engine.dispose()


def test_bulk_violation_import_is_captured_by_central_audit_listener() -> None:
    import database.session  # noqa: F401 - installs production Session listener

    engine = _engine()
    with Session(engine) as session:
        company, admin, *_ = _seed(session)
        service = PolicyViolationBulkImportService(session)
        payload = service.build_template(company_id=company.id)
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Violation Import Template"]
        sheet.delete_rows(2, 1)
        sheet.append([
            "CON-001", "Conduct", "Insubordination", "Major",
            "Refusal to follow a lawful company instruction.",
            "Written Warning", "Final Written Warning", "Suspension",
            "Termination", "", "2026-08-24", "Active", "",
        ])
        out = BytesIO()
        workbook.save(out)
        preview = service.prepare_preview(
            out.getvalue(), filename="violations.xlsx", company_id=company.id
        )
        set_audit_actor(company_id=company.id, user_id=admin.id, module="Policies")
        try:
            created = service.import_preview_rows(
                preview, company_id=company.id, actor_user_id=admin.id
            )
        finally:
            clear_audit_actor()
        event = session.scalar(
            select(AuditEvent).where(
                AuditEvent.company_id == company.id,
                AuditEvent.entity_type == "PolicyViolation",
                AuditEvent.entity_label == created[0].public_id,
            )
        )
        assert event is not None
        assert event.action == "Created"
        assert event.module == "Policies"
    engine.dispose()
