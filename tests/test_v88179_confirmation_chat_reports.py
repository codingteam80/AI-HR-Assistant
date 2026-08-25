"""v8.8.179 confirmation invalidation and Chat Assistant reporting checks."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from modules.reports.chat_assistant_report import (
    build_chat_report_excel,
    build_chat_report_pdf,
)
from services.chat_report_service import ChatDateRangeParser


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _sample_report(chart_type: str = "bar") -> dict:
    return {
        "title": "Leave Report — August 2026",
        "period": "August 2026",
        "start_date": "2026-08-01",
        "end_date": "2026-08-31",
        "generated_at": "2026-08-22T01:01:00+08:00",
        "timezone": "Asia/Manila",
        "summary": ["Leave requests: 2", "Employees represented: 2"],
        "columns": ["Employee", "Leave Type", "Days"],
        "rows": [
            {"Employee": "Employee A", "Leave Type": "Vacation Leave", "Days": 2.0},
            {"Employee": "Employee B", "Leave Type": "Sick Leave", "Days": 1.0},
        ],
        "chart": {
            "type": chart_type,
            "title": "Leave requests by type",
            "category_label": "Leave Type",
            "value_label": "Requests",
            "data": [
                {"label": "Vacation Leave", "value": 1},
                {"label": "Sick Leave", "value": 1},
            ],
        },
        "domain": "leave",
    }


def test_leave_review_confirmation_tracks_all_five_requested_settings() -> None:
    source = _source("ui/pages/admin/leave_management_page.py")
    assert "invalidate_confirmation_on_change(" in source
    assert '"reset_month": selected_reset_month' in source
    assert '"reset_day": selected_reset_day' in source
    assert '"utilization_enabled": selected_utilization_enabled' in source
    assert '"utilization_percentage": selected_percentage' in source
    assert '"manager_vl_retention": selected_manager_retention' in source
    assert '"I reviewed the reset date and utilization settings."' in source
    assert 'with st.form("company_leave_policy_form")' not in source


def test_project_wide_confirmation_guard_is_used_for_similar_admin_workflows() -> None:
    expected_files = (
        "ui/pages/admin/company_page.py",
        "ui/pages/admin/announcements_page.py",
        "ui/pages/admin/employees_page.py",
        "ui/pages/admin/onboarding_management.py",
        "ui/pages/admin/policies_page.py",
        "ui/pages/admin/hr_contacts_management.py",
        "ui/pages/admin/company_forms_documents_page.py",
    )
    for path in expected_files:
        assert "invalidate_confirmation_on_change" in _source(path), path


def test_date_parser_supports_specific_month_year_range_relative_and_quarters() -> None:
    parser = ChatDateRangeParser(today=date(2026, 8, 22))
    this_august = parser.parse("How many employees are on leave this August?")
    assert this_august is not None
    assert (this_august.start, this_august.end) == (date(2026, 8, 1), date(2026, 8, 31))

    specific = parser.parse("Who is on leave on August 25, 2026?")
    assert specific is not None
    assert specific.start == specific.end == date(2026, 8, 25)

    date_range = parser.parse("Show leave requests from August 1, 2026 to August 15, 2026")
    assert date_range is not None
    assert (date_range.start, date_range.end) == (date(2026, 8, 1), date(2026, 8, 15))

    last_six = parser.parse("Create attendance report for the last 6 months")
    assert last_six is not None
    assert last_six.start == date(2026, 3, 1)
    assert last_six.end == date(2026, 8, 22)

    quarters = parser.parse("Compare attendance from Q1 to Q2 2026")
    assert quarters is not None
    assert (quarters.start, quarters.end) == (date(2026, 1, 1), date(2026, 6, 30))


def test_date_parser_does_not_mistake_modal_may_for_month() -> None:
    parser = ChatDateRangeParser(today=date(2026, 8, 22))
    assert parser.parse("May I carry over unused leave?") is None
    may_period = parser.parse("Show leave in May")
    assert may_period is not None
    assert (may_period.start, may_period.end) == (date(2026, 5, 1), date(2026, 5, 31))


def test_chat_assistant_report_engine_is_live_permission_aware_and_not_qwen_generated() -> None:
    source = _source("services/chat_report_service.py")
    assert "_authorized_employees" in source
    assert "Employee.manager_id == employee_id" in source
    assert "Employee.leader_id == employee_id" in source
    assert 'role_scope == "admin"' in source
    assert "ChatReportService" in _source("modules/hr_assistant/admin_hr_assistant.py")
    assert "ChatReportService" in _source("modules/hr_assistant/hr_assistant.py")
    assert '"live_report"' in _source("modules/smart_ai/portal_ai.py")
    assert '_NEVER_ENHANCE_INTENTS' in _source("modules/smart_ai/portal_ai.py")


def test_chat_pages_store_timestamp_and_report_for_every_new_assistant_answer() -> None:
    for path in ("ui/pages/admin/chat_page.py", "ui/pages/user/chat_page.py"):
        source = _source(path)
        assert '"report": response.report' in source
        assert '"timestamp": current_assistant_timestamp()' in source
        assert "render_chat_report(" in source
        assert "render_assistant_timestamp(" in source


def test_prompt_receives_current_company_datetime_and_keeps_structured_answer_rules() -> None:
    prompt = _source("modules/smart_ai/prompts/hr_assistant_prompt.py")
    portal = _source("modules/smart_ai/portal_ai.py")
    assert "CURRENT COMPANY DATE/TIME" in prompt
    assert "today, yesterday, tomorrow, now, this week, this month, this year" in prompt
    assert "Two or more distinct facts" in prompt
    assert "Markdown bullet list" in prompt
    assert "Markdown numbered list" in prompt
    assert "current_datetime_text=datetime.now(" in portal
    assert "ZoneInfo(get_settings().display_timezone)" in portal


def test_chat_report_excel_uses_same_summary_data_and_chart_artifact() -> None:
    workbook_bytes = build_chat_report_excel(_sample_report("bar"))
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
    assert workbook.sheetnames == ["Summary", "Report Data", "Chart Data"]
    assert workbook["Summary"]["A1"].value == "Leave Report — August 2026"
    assert workbook["Report Data"]["A2"].value == "Employee A"
    assert workbook["Chart Data"]["A2"].value == "Vacation Leave"
    assert workbook["Chart Data"]["B2"].value == 1
    assert len(workbook["Chart Data"]._charts) == 1


def test_chat_report_pdf_supports_bar_line_pie_and_donut() -> None:
    for chart_type in ("bar", "line", "pie", "donut"):
        output = build_chat_report_pdf(_sample_report(chart_type))
        assert output.startswith(b"%PDF")
        assert len(output) > 1500


def test_v88179_version_markers() -> None:
    assert 'app_version: str = "0.8.8.179"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.179" in _source(".env")
    assert "APP_VERSION=0.8.8.179" in _source(".env.example")
