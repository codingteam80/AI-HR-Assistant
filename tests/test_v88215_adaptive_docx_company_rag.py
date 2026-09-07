"""v8.8.215 adaptive DOCX structure and company-document RAG regressions."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document

from config.chat_assistant_settings import get_chat_assistant_settings
from modules.documents.docx_structure_extractor import DOCXStructureExtractor
from modules.documents.policy_file_parser import PolicyFileParser
from modules.smart_ai import portal_ai
from services.company_form_knowledge_service import CompanyFormKnowledgeExtractor


ROOT = Path(__file__).resolve().parents[1]
MAX_BYTES = 2 * 1024 * 1024


def _save(document: Document) -> bytes:
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _bold_paragraph(document: Document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    return paragraph


def test_normal_style_bold_and_semantic_headings_are_detected_without_word_heading_styles() -> None:
    document = Document()
    _bold_paragraph(document, "Policy Statement")
    document.add_paragraph("Employees receive benefits under the company policy.")
    _bold_paragraph(
        document,
        "Section 1. Employees must submit the required form before processing.",
    )

    extracted = DOCXStructureExtractor.extract(_save(document))

    assert extracted.reliable_heading_count >= 2
    assert any(section.heading == "Policy Statement" for section in extracted.sections)
    assert any(section.heading == "Section 1" for section in extracted.sections)
    assert "Employees must submit the required form before processing." in " ".join(
        section.text for section in extracted.sections
    )


def test_docx_tables_remain_in_original_block_order() -> None:
    document = Document()
    document.add_paragraph("Before table explanation")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Level"
    table.cell(0, 1).text = "Award"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "4500"
    document.add_paragraph("After table explanation")

    extracted = DOCXStructureExtractor.extract(_save(document))

    assert extracted.full_text.index("Before table explanation") < extracted.full_text.index("Level | Award")
    assert extracted.full_text.index("Level | Award") < extracted.full_text.index("After table explanation")


def test_no_reliable_headings_falls_back_to_one_file_level_section_and_bounded_chunks() -> None:
    document = Document()
    for index in range(1, 90):
        document.add_paragraph(
            f"ordinary paragraph {index} contains searchable company information without special formatting"
        )
    data = _save(document)

    extracted = DOCXStructureExtractor.extract(data)
    assert extracted.reliable_heading_count == 0
    assert len(extracted.sections) == 1
    assert extracted.sections[0].heading == "Document"

    rendered = CompanyFormKnowledgeExtractor._docx(data)
    assert "## " not in rendered
    sections = portal_ai._document_sections(rendered)
    assert len(sections) == 1
    chunks = portal_ai._prepare_document_chunks(
        [(label, text, None) for label, text in sections],
        chunk_size=300,
        overlap=56,
    )
    assert len(chunks) > 1
    assert all(len(chunk.split()) <= 300 for _label, chunk, _page in chunks)


def test_adaptive_chunking_merges_tiny_adjacent_sections_without_losing_text() -> None:
    chunks = portal_ai._prepare_document_chunks(
        [
            ("Article I", "Article I - Leave Benefits", None),
            ("Section 1", "Section 1. Vacation leave eligibility applies to regular employees.", None),
            ("Eligibility", "Eligibility requires one full year of continuous service.", None),
            ("Long Rule", " ".join(f"word{index}" for index in range(350)), None),
        ],
        chunk_size=260,
        overlap=48,
    )
    joined = " ".join(chunk for _label, chunk, _page in chunks)
    assert "Article I - Leave Benefits" in joined
    assert "one full year of continuous service" in joined
    assert "word349" in joined
    assert all(len(chunk.split()) <= 260 for _label, chunk, _page in chunks)
    assert any("Article I" in label and "Section 1" in label for label, _chunk, _page in chunks)


def test_policy_docx_normal_style_rules_keep_rule_text_searchable() -> None:
    document = Document()
    _bold_paragraph(document, "Service Award")
    _bold_paragraph(document, "Article I - Service Awards")
    _bold_paragraph(
        document,
        "Section 1. Employees with five years of service receive the applicable cash award.",
    )
    parsed = PolicyFileParser.parse(
        filename="service-award.docx",
        file_bytes=_save(document),
        maximum_size_bytes=MAX_BYTES,
    )
    searchable = " ".join(section.text for section in parsed.sections)
    assert "five years of service receive the applicable cash award" in searchable


def test_file_stage_prefers_strong_exact_evidence_over_many_weaker_long_file_hits(monkeypatch) -> None:
    retriever = portal_ai.HybridRetriever()
    monkeypatch.setattr(retriever.vector, "search", lambda *_a, **_k: [])
    monkeypatch.setattr(retriever.file_vector, "search", lambda *_a, **_k: [])
    monkeypatch.setattr(retriever.settings, "file_top_k", 1)

    documents = [
        portal_ai.KnowledgeDocument(
            document_id="exact-1",
            text="Failure to log in and log out automatically disqualifies an employee from the Perfect Attendance Award.",
            title="Perfect Attendance",
            source_type="company_form_content",
            company_id=1,
            role_scope="shared",
            metadata={
                "file_key": "file:perfect",
                "title": "Perfect Attendance",
                "filename": "Perfect Attendance.docx",
                "section": "Section 3",
                "chunk_index": 1,
            },
        ),
    ]
    for index in range(1, 8):
        documents.append(
            portal_ai.KnowledgeDocument(
                document_id=f"long-{index}",
                text=(
                    "Attendance log procedure and employee time record guidance "
                    f"for company operations part {index}."
                ),
                title="Long Work Schedule",
                source_type="company_form_content",
                company_id=1,
                role_scope="shared",
                metadata={
                    "file_key": "file:long",
                    "title": "Long Work Schedule",
                    "filename": "Work Schedule.docx",
                    "section": "Log Procedure",
                    "chunk_index": index,
                },
            )
        )

    selected, _grouped = retriever._select_file_keys(
        "Does failure to log in and log out disqualify an employee from the Perfect Attendance Award?",
        documents,
        access_partition="company:1:role:admin",
    )
    assert selected == {"file:perfect"}


def test_v88215_document_budgets_and_extraction_ceiling() -> None:
    settings = get_chat_assistant_settings()
    assert settings.policy_chunk_size == 260
    assert settings.policy_chunk_overlap == 48
    assert settings.company_document_chunk_size == 300
    assert settings.company_document_chunk_overlap == 56
    assert settings.file_top_k == 5
    assert settings.file_candidate_k == 10
    assert settings.max_chunks_per_file == 3
    assert settings.bm25_top_k == 10
    assert settings.vector_top_k == 10
    assert settings.final_top_k == 6
    assert CompanyFormKnowledgeExtractor.MAX_CHARACTERS == 150_000


def test_v88215_preserves_existing_hybrid_permission_architecture() -> None:
    source = (ROOT / "modules" / "smart_ai" / "portal_ai.py").read_text(encoding="utf-8")
    assert "class BM25Retriever" in source
    assert "class ChromaRetriever" in source
    assert "class HybridRetriever" in source
    assert "_prepare_document_chunks(" in source
    assert "company_id=current_user.company_id" in source
    assert "role_scope=role_scope" in source
    assert "neighbor_chunk_window" in source


def test_v88215_version_markers() -> None:
    assert 'app_version: str = "0.8.8.215"' in (
        ROOT / "config" / "settings.py"
    ).read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.215" in (ROOT / ".env").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.215" in (ROOT / ".env.example").read_text(encoding="utf-8")
