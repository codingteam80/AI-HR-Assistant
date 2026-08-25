"""v8.8.194 hierarchy routing and live Excel-reference synchronization."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import models  # noqa: F401
from authentication.current_user import AuthenticatedUser
from database.base import Base
from models.company import Company
from models.department import Department
from models.employee import Employee
from models.hr_policy import HRPolicy
from models.policy_violation import PolicyViolation
from models.role import Role
from models.user import User
from modules.hr_assistant.admin_hr_assistant import AdminHRAssistant
from modules.hr_assistant.hr_assistant import HRAssistant
from services.disciplinary_record_bulk_import_service import DisciplinaryRecordBulkImportService
from services.employee_bulk_import_service import EmployeeBulkImportService
from services.policy_violation_bulk_import_service import PolicyViolationBulkImportService


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def _seed_company(session: Session, *, code: str):
    company = Company(code=code, name=f"{code} Company", is_active=True)
    session.add(company)
    session.flush()
    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role])
    session.flush()

    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username=f"{code.lower()}_admin",
        email=f"{code.lower()}_admin@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    maria_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username=f"{code.lower()}_maria",
        email=f"{code.lower()}_maria@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    leo_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username=f"{code.lower()}_leo",
        email=f"{code.lower()}_leo@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    ana_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username=f"{code.lower()}_ana",
        email=f"{code.lower()}_ana@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add_all([admin, maria_user, leo_user, ana_user])
    session.flush()

    department = Department(
        company_id=company.id,
        name=f"{code} Operations",
        code=f"{code}-OPS",
        is_active=True,
    )
    session.add(department)
    session.flush()

    maria = Employee(
        company_id=company.id,
        user_id=maria_user.id,
        employee_number=f"{code}-MGR-001",
        first_name="Maria",
        last_name="Santos",
        job_title="Senior Specialist",  # title intentionally does not say Manager
        employment_status="employed",
        department_id=department.id,
    )
    leo = Employee(
        company_id=company.id,
        user_id=leo_user.id,
        employee_number=f"{code}-LDR-001",
        first_name="Leo",
        last_name="Cruz",
        job_title="Engineer",  # title intentionally does not say Leader
        employment_status="employed",
        department_id=department.id,
    )
    ana = Employee(
        company_id=company.id,
        user_id=ana_user.id,
        employee_number=f"{code}-MEM-001",
        first_name="Ana",
        last_name="Reyes",
        job_title="Team Lead",  # title must not create a hierarchy role
        employment_status="employed",
        department_id=department.id,
    )
    session.add_all([maria, leo, ana])
    session.flush()

    # Maria is Manager-only; Leo is Leader-only. Ana's title is misleading but
    # she has no assigned member herself.
    ana.manager_id = maria.id
    ana.leader_id = leo.id

    policy = HRPolicy(
        public_id=f"POL-{code}-001",
        company_id=company.id,
        created_by_user_id=admin.id,
        title=f"{code} Conduct Policy",
        category="Conduct",
        content="Company conduct rules.",
        version="1.0",
        status="published",
    )
    session.add(policy)
    session.flush()

    violation = PolicyViolation(
        public_id=f"VIO-{code}-001",
        company_id=company.id,
        created_by_user_id=admin.id,
        violation_code=f"{code}-ATT-001",
        category="Attendance",
        offense_title="Habitual Tardiness",
        description="Repeated tardiness.",
        severity="minor",
        first_offense_action="Verbal Warning",
        second_offense_action="Written Warning",
        third_offense_action="Suspension",
        final_action="Termination",
        related_policy_id=policy.id,
        status="active",
    )
    session.add(violation)
    session.commit()
    return company, admin, maria_user, leo_user, ana_user, department, maria, leo, ana, policy, violation


def _current_user(company, user, employee=None, *, clearance: int):
    return AuthenticatedUser(
        user_id=user.id,
        company_id=company.id,
        company_code=company.code,
        company_name=company.name,
        role_id=user.role_id,
        role_name="company_admin" if clearance == 1 else "employee",
        clearance=clearance,
        username=user.username,
        email=user.email,
        employee_id=employee.id if employee is not None else None,
        employee_number=employee.employee_number if employee is not None else None,
        employee_name=employee.full_name if employee is not None else None,
        must_change_password=False,
    )


def test_version_markers() -> None:
    assert 'app_version: str = "0.8.8.194"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.194" in _source(".env")


@pytest.mark.parametrize(
    ("question", "expected_role", "expected_member"),
    [
        ("May member ba si Maria?", "Manager", "MAIN-MEM-001"),
        ("Who is under Maria?", "Manager", "MAIN-MEM-001"),
        ("Sino member ni Leo?", "Leader", "MAIN-MEM-001"),
        ("Is Leo a leader or manager?", "Leader", "MAIN-MEM-001"),
        ("What team members does Leo handle?", "Leader", "MAIN-MEM-001"),
    ],
)
def test_admin_hierarchy_understands_varied_member_phrasing_and_unique_first_name(
    question: str,
    expected_role: str,
    expected_member: str,
) -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed_company(session, code="MAIN")
            current = _current_user(company, admin, clearance=1)
            answer = AdminHRAssistant(session).answer(
                current_user=current,
                question=question,
            )
            assert answer.intent == "employee_hierarchy"
            assert f"**Hierarchy Role:** {expected_role}" in answer.answer
            assert expected_member in answer.answer
    finally:
        engine.dispose()


def test_employee_info_and_self_profile_distinguish_role_from_live_assignments() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            (
                company,
                admin,
                maria_user,
                _leo_user,
                _ana_user,
                _department,
                maria,
                _leo,
                ana,
                _policy,
                _violation,
            ) = _seed_company(session, code="MAIN")
            admin_current = _current_user(company, admin, clearance=1)
            admin_answer = AdminHRAssistant(session).answer(
                current_user=admin_current,
                question="Tell me about Maria",
            )
            assert admin_answer.intent == "employee_lookup"
            assert "**Hierarchy Role:** Manager" in admin_answer.answer
            assert "**Manager Direct Reports:** 1" in admin_answer.answer
            assert "**Leader Team Members:** 0" in admin_answer.answer

            maria_current = _current_user(company, maria_user, maria, clearance=2)
            employee_answer = HRAssistant(session).answer(
                current_user=maria_current,
                question="my profile",
            )
            assert employee_answer.intent == "employee_profile"
            assert "**Hierarchy Role:** Manager" in employee_answer.answer
            assert "**Manager Direct Reports:** 1" in employee_answer.answer

            # A leadership-sounding title never creates a role without actual
            # subordinate assignments.
            no_role = AdminHRAssistant(session).answer(
                current_user=admin_current,
                question=f"Is {ana.employee_number} a leader or manager?",
            )
            assert "No supervisory assignment" in no_role.answer
    finally:
        engine.dispose()


def test_employee_hierarchy_remains_self_scoped_for_third_party_questions() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, _admin, maria_user, _leo_user, _ana_user, _department, maria, leo, ana, *_ = _seed_company(session, code="MAIN")
            current = _current_user(company, maria_user, maria, clearance=2)
            assistant = HRAssistant(session)
            own = assistant.answer(current_user=current, question="May member ba ako?")
            assert own.intent == "employee_hierarchy"
            assert "**Hierarchy Role:** Manager" in own.answer
            assert ana.employee_number in own.answer

            third_party = assistant.answer(
                current_user=current,
                question=f"Sino member ni {leo.employee_number}?",
            )
            assert third_party.intent == "employee_hierarchy"
            assert "only your own" in third_party.answer
            assert ana.employee_number not in third_party.answer
    finally:
        engine.dispose()


def test_employee_template_comments_refresh_with_live_company_references() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, _admin, *_rest = _seed_company(session, code="MAIN")
            other, *_ = _seed_company(session, code="OTHER")
            service = EmployeeBulkImportService(session)
            first = load_workbook(BytesIO(service.build_company_template(company_id=company.id)))
            sheet = first["Employee Import Template"]
            assert "MAIN Operations" in sheet["K1"].comment.text
            assert "OTHER Operations" not in sheet["K1"].comment.text
            assert "MAIN-MGR-001" in sheet["L1"].comment.text
            assert "OTHER-MGR-001" not in sheet["L1"].comment.text

            new_department = Department(
                company_id=company.id,
                name="New Live Department",
                code="NEW-LIVE",
                is_active=True,
            )
            session.add(new_department)
            session.commit()

            refreshed = load_workbook(BytesIO(service.build_company_template(company_id=company.id)))
            refreshed_sheet = refreshed["Employee Import Template"]
            assert "New Live Department" in refreshed_sheet["K1"].comment.text
            assert "New Live Department" in {
                str(cell.value or "")
                for cell in refreshed["Department References"]["A"]
            }
    finally:
        engine.dispose()


def test_policy_and_disciplinary_comments_show_current_exact_company_values() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_main = _seed_company(session, code="MAIN")
            other, other_admin, *_other = _seed_company(session, code="OTHER")

            violation_template = load_workbook(
                BytesIO(PolicyViolationBulkImportService(session).build_template(company_id=company.id))
            )
            v_sheet = violation_template["Violation Import Template"]
            assert "POL-MAIN-001" in v_sheet["J1"].comment.text
            assert "POL-OTHER-001" not in v_sheet["J1"].comment.text
            violation_validations = {
                str(item.sqref): str(item.formula1)
                for item in v_sheet.data_validations.dataValidation
            }
            assert violation_validations["J2:J2000"] == "=ViolationRelatedPolicyIds"

            disciplinary_template = load_workbook(
                BytesIO(DisciplinaryRecordBulkImportService(session).build_template(company_id=company.id))
            )
            d_sheet = disciplinary_template["Disciplinary Case Import"]
            assert "MAIN-MGR-001" in d_sheet["A1"].comment.text
            assert "OTHER-MGR-001" not in d_sheet["A1"].comment.text
            assert "MAIN-ATT-001" in d_sheet["B1"].comment.text
            assert "OTHER-ATT-001" not in d_sheet["B1"].comment.text
            assert admin.username in d_sheet["H1"].comment.text
            assert other_admin.username not in d_sheet["H1"].comment.text
    finally:
        engine.dispose()


def test_dynamic_template_references_exclude_inactive_archived_values() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed_company(session, code="LIVE")
            role = next(role for role in session.query(Role).filter(Role.company_id == company.id).all() if role.name == "employee")
            inactive_user = User(
                company_id=company.id,
                role_id=role.id,
                clearance=2,
                username="inactive_user",
                email="inactive@example.com",
                password_hash="hash",
                is_active=False,
                must_change_password=False,
            )
            session.add(inactive_user)
            session.flush()
            inactive_department = Department(
                company_id=company.id,
                name="Inactive Department",
                code="INACTIVE-DEPT",
                is_active=False,
            )
            session.add(inactive_department)
            session.flush()
            resigned = Employee(
                company_id=company.id,
                user_id=inactive_user.id,
                employee_number="LIVE-OLD-001",
                first_name="Old",
                last_name="Employee",
                employment_status="resigned",
                department_id=inactive_department.id,
            )
            session.add(resigned)
            trashed_policy = HRPolicy(
                public_id="POL-LIVE-TRASH",
                company_id=company.id,
                created_by_user_id=admin.id,
                title="Trashed Policy",
                category="Conduct",
                content="Old.",
                version="1.0",
                status="trashed",
            )
            session.add(trashed_policy)
            archived_violation = PolicyViolation(
                public_id="VIO-LIVE-ARCH",
                company_id=company.id,
                created_by_user_id=admin.id,
                violation_code="LIVE-ARCH-001",
                category="Conduct",
                offense_title="Archived Offense",
                description="Old.",
                severity="minor",
                first_offense_action="Warning",
                second_offense_action="Warning",
                third_offense_action="Warning",
                final_action="Warning",
                status="archived",
            )
            session.add(archived_violation)
            session.commit()

            employee_wb = load_workbook(BytesIO(EmployeeBulkImportService(session).build_company_template(company_id=company.id)))
            employee_sheet = employee_wb["Employee Import Template"]
            assert "Inactive Department" not in employee_sheet["K1"].comment.text
            assert "LIVE-OLD-001" not in employee_sheet["L1"].comment.text

            violation_wb = load_workbook(BytesIO(PolicyViolationBulkImportService(session).build_template(company_id=company.id)))
            assert "POL-LIVE-TRASH" not in violation_wb["Violation Import Template"]["J1"].comment.text

            disciplinary_wb = load_workbook(BytesIO(DisciplinaryRecordBulkImportService(session).build_template(company_id=company.id)))
            d_sheet = disciplinary_wb["Disciplinary Case Import"]
            assert "LIVE-OLD-001" not in d_sheet["A1"].comment.text
            assert "LIVE-ARCH-001" not in d_sheet["B1"].comment.text
            assert "inactive_user" not in d_sheet["H1"].comment.text
    finally:
        engine.dispose()


def test_admin_short_name_hierarchy_requires_employee_number_when_ambiguous() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed_company(session, code="AMB")
            role = next(role for role in session.query(Role).filter(Role.company_id == company.id).all() if role.name == "employee")
            second_user = User(
                company_id=company.id,
                role_id=role.id,
                clearance=2,
                username="amb_maria2",
                email="amb_maria2@example.com",
                password_hash="hash",
                is_active=True,
                must_change_password=False,
            )
            session.add(second_user)
            session.flush()
            second_maria = Employee(
                company_id=company.id,
                user_id=second_user.id,
                employee_number="AMB-MGR-002",
                first_name="Maria",
                last_name="Lopez",
                employment_status="employed",
            )
            session.add(second_maria)
            session.commit()

            answer = AdminHRAssistant(session).answer(
                current_user=_current_user(company, admin, clearance=1),
                question="May member ba si Maria?",
            )
            assert answer.intent == "employee_hierarchy"
            assert "Multiple employees matched" in answer.answer
            assert "Employee Number" in answer.answer
    finally:
        engine.dispose()


def test_excel_comment_guidance_uses_reference_sheet_fallback_for_huge_live_lists() -> None:
    from openpyxl import Workbook
    from services.excel_template_guidance import (
        EXCEL_COMMENT_SAFE_TEXT_LIMIT,
        add_header_guidance,
    )

    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Employee Number"
    add_header_guidance(
        sheet,
        "A1",
        requirement="Use a current employee number.",
        choices=(f"EMP-{index:05d}" for index in range(10_000)),
        choices_heading="Current employee numbers:",
        extra="See the Employee References worksheet for the complete current list.",
    )
    text = sheet["A1"].comment.text
    assert len(text) < EXCEL_COMMENT_SAFE_TEXT_LIMIT + 1_000
    assert "too large for one Excel comment" in text
    assert "reference worksheet" in text
