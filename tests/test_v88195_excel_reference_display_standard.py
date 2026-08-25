"""v8.8.195 human-readable exact-reference Excel standard."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import models  # noqa: F401
from database.base import Base
from models.company import Company
from models.department import Department
from models.employee import Employee
from models.hr_policy import HRPolicy
from models.hr_policy_document import HRPolicyDocument
from models.policy_violation import PolicyViolation
from models.role import Role
from models.user import User
from services.disciplinary_record_bulk_import_service import (
    DISCIPLINARY_IMPORT_COLUMNS,
    DisciplinaryRecordBulkImportService,
)
from services.employee_bulk_import_service import EmployeeBulkImportService
from services.policy_violation_bulk_import_service import PolicyViolationBulkImportService

ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def _seed(session: Session):
    company = Company(code="REF", name="Reference Company", is_active=True)
    other = Company(code="OTH", name="Other Company", is_active=True)
    session.add_all([company, other])
    session.flush()

    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    other_role = Role(company_id=other.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role, other_role])
    session.flush()

    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username="ref_admin",
        email="ref_admin@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    maria_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username="maria.santos",
        email="maria@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    reviewer_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username="pedro.reyes",
        email="pedro@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    other_user = User(
        company_id=other.id,
        role_id=other_role.id,
        clearance=2,
        username="other.user",
        email="other@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add_all([admin, maria_user, reviewer_user, other_user])
    session.flush()

    department = Department(company_id=company.id, name="Operations", is_active=True)
    other_department = Department(company_id=other.id, name="Secret", is_active=True)
    session.add_all([department, other_department])
    session.flush()

    maria = Employee(
        company_id=company.id,
        user_id=maria_user.id,
        employee_number="EMP-001",
        first_name="Maria",
        last_name="Santos",
        suffix="Jr.",
        employment_status="employed",
        department_id=department.id,
    )
    pedro = Employee(
        company_id=company.id,
        user_id=reviewer_user.id,
        employee_number="EMP-002",
        first_name="Pedro",
        last_name="Reyes",
        employment_status="employed",
        department_id=department.id,
    )
    other_employee = Employee(
        company_id=other.id,
        user_id=other_user.id,
        employee_number="OTH-001",
        first_name="Other",
        last_name="Person",
        employment_status="employed",
        department_id=other_department.id,
    )
    session.add_all([maria, pedro, other_employee])
    session.flush()

    policy = HRPolicy(
        public_id="POL-REF-001",
        company_id=company.id,
        created_by_user_id=admin.id,
        title="Attendance Policy",
        category="Attendance",
        content="Attendance rules.",
        version="1.0",
        status="published",
    )
    session.add(policy)
    session.flush()
    session.add(
        HRPolicyDocument(
            company_id=company.id,
            policy_id=policy.id,
            uploaded_by_user_id=admin.id,
            original_filename="Attendance_Policy.pdf",
            stored_filename="stored.pdf",
            storage_path="company_1/stored.pdf",
            mime_type="application/pdf",
            file_extension=".pdf",
            sha256="a" * 64,
            size_bytes=123,
            page_count=1,
            extracted_text="Attendance rules.",
        )
    )

    other_policy = HRPolicy(
        public_id="POL-OTH-001",
        company_id=other.id,
        created_by_user_id=other_user.id,
        title="Secret Policy",
        category="Secret",
        content="Secret.",
        version="1.0",
        status="published",
    )
    session.add(other_policy)
    session.flush()
    session.add(
        HRPolicyDocument(
            company_id=other.id,
            policy_id=other_policy.id,
            uploaded_by_user_id=other_user.id,
            original_filename="Secret_Policy.pdf",
            stored_filename="secret.pdf",
            storage_path="company_2/secret.pdf",
            mime_type="application/pdf",
            file_extension=".pdf",
            sha256="b" * 64,
            size_bytes=123,
            page_count=1,
            extracted_text="Secret.",
        )
    )

    violation = PolicyViolation(
        public_id="VIO-REF-001",
        company_id=company.id,
        created_by_user_id=admin.id,
        violation_code="ATT-001",
        category="Attendance",
        offense_title="Habitual Tardiness",
        description="Repeated tardiness.",
        severity="minor",
        first_offense_action="Verbal Warning",
        second_offense_action="Written Warning",
        third_offense_action="Suspension",
        final_action="Termination",
        related_policy_id=policy.id,
        status="active",
    )
    session.add(violation)
    session.commit()
    return company, admin, maria, pedro, policy, violation


def test_version_markers() -> None:
    assert 'app_version: str = "0.8.8.195"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.195" in _source(".env")


def test_disciplinary_template_uses_readable_employee_user_and_violation_selections() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, maria, pedro, _policy, violation = _seed(session)
            service = DisciplinaryRecordBulkImportService(session)
            workbook = load_workbook(BytesIO(service.build_template(company_id=company.id)))
            sheet = workbook["Disciplinary Case Import"]

            assert tuple(cell.value for cell in sheet[1]) == DISCIPLINARY_IMPORT_COLUMNS
            assert sheet["H1"].value == "Issued By"
            assert sheet["I1"].value == "Reviewed / Approved By"

            maria_ref = "EMP-001 - Santos, Maria Jr."
            pedro_ref = "EMP-002 - Reyes, Pedro"
            violation_ref = "ATT-001 - Habitual Tardiness"
            assert maria_ref in sheet["A1"].comment.text
            assert maria_ref in sheet["H1"].comment.text
            assert pedro_ref in sheet["I1"].comment.text
            assert violation_ref in sheet["B1"].comment.text
            assert "OTH-001" not in sheet["A1"].comment.text

            validations = {
                str(item.sqref): str(item.formula1)
                for item in sheet.data_validations.dataValidation
            }
            assert validations["A2:A2000"] == "=DisciplinaryEmployeeNumbers"
            assert validations["B2:B2000"] == "=DisciplinaryViolationCodes"
            assert validations["H2:H2000"] == "=DisciplinaryUsernames"
            assert validations["I2:I2000"] == "=DisciplinaryUsernames"

            # First columns retain raw identifiers for audit/reference-sheet
            # readability; dropdowns point at the human-readable second column.
            emp_ref = workbook["Employee References"]
            employee_rows = {str(row[0].value): str(row[1].value) for row in emp_ref.iter_rows(min_row=2)}
            assert employee_rows["EMP-001"] == maria_ref
            vio_ref = workbook["Violation References"]
            violation_rows = {str(row[0].value): str(row[1].value) for row in vio_ref.iter_rows(min_row=2)}
            assert violation_rows[violation.violation_code] == violation_ref
            user_ref = workbook["User References"]
            visible_user_references = {str(row[1].value or "") for row in user_ref.iter_rows(min_row=2)}
            assert maria_ref in visible_user_references
            assert pedro_ref in visible_user_references
            assert admin.username not in {str(row[0].value or "") for row in user_ref.iter_rows(min_row=2)}

            # Fill using the new display values and verify they resolve to
            # canonical IDs without saving the labels as identifiers.
            values = [
                maria_ref,
                violation_ref,
                "2026-08-25",
                "Reference-label import test.",
                "",
                "Operations",
                "Verbal Warning",
                maria_ref,
                pedro_ref,
                "2026-08-25",
                "Pending",
                "Issued",
                "",
            ]
            for index, value in enumerate(values, start=1):
                sheet.cell(2, index).value = value
            out = BytesIO()
            workbook.save(out)
            preview = service.prepare_preview(
                out.getvalue(), filename="Disciplinary_Case_Import_Template.xlsx", company_id=company.id
            )
            assert preview[0]["Validation"] == "Ready"
            payload = preview[0]["_values"]
            assert payload["employee_id"] == maria.id
            assert payload["issued_by_user_id"] == maria.user_id
            assert payload["reviewed_approved_by_user_id"] == pedro.user_id
    finally:
        engine.dispose()


def test_disciplinary_import_keeps_legacy_raw_headers_and_values_compatible() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, maria, _pedro, _policy, violation = _seed(session)
            service = DisciplinaryRecordBulkImportService(session)
            workbook = load_workbook(BytesIO(service.build_template(company_id=company.id)))
            sheet = workbook["Disciplinary Case Import"]
            sheet["H1"] = "Issued By Username"
            sheet["I1"] = "Reviewed / Approved By Username"
            values = [
                maria.employee_number,
                violation.violation_code,
                "2026-08-25",
                "Legacy raw-reference compatibility.",
                "",
                "",
                "",
                admin.username,
                "",
                "",
                "Pending",
                "Draft",
                "",
            ]
            for index, value in enumerate(values, start=1):
                sheet.cell(2, index).value = value
            out = BytesIO(); workbook.save(out)
            preview = service.prepare_preview(out.getvalue(), filename="legacy.xlsx", company_id=company.id)
            assert preview[0]["Validation"] == "Ready"
            assert preview[0]["_values"]["employee_id"] == maria.id
            assert preview[0]["_values"]["issued_by_user_id"] == admin.id
    finally:
        engine.dispose()


def test_violation_policy_reference_is_policy_id_plus_filename_and_import_resolves_it() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, _admin, _maria, _pedro, policy, _violation = _seed(session)
            service = PolicyViolationBulkImportService(session)
            workbook = load_workbook(BytesIO(service.build_template(company_id=company.id)))
            sheet = workbook["Violation Import Template"]
            policy_ref = "POL-REF-001 - Attendance_Policy.pdf"
            assert policy_ref in sheet["J1"].comment.text
            assert "POL-OTH-001" not in sheet["J1"].comment.text
            assert workbook["Related Policies"]["A2"].value == "POL-REF-001"
            assert workbook["Related Policies"]["B2"].value == policy_ref
            validations = {
                str(item.sqref): str(item.formula1)
                for item in sheet.data_validations.dataValidation
            }
            assert validations["J2:J2000"] == "=ViolationRelatedPolicyIds"

            row = [
                "ATT-NEW-001",
                "Attendance",
                "Late filing",
                "Minor",
                "Late filing of attendance correction.",
                "Verbal Warning",
                "Written Warning",
                "Suspension",
                "Termination",
                policy_ref,
                "2026-08-25",
                "Active",
                "",
            ]
            for index, value in enumerate(row, start=1):
                sheet.cell(2, index).value = value
            out = BytesIO(); workbook.save(out)
            preview = service.prepare_preview(
                out.getvalue(), filename="Violation_Import_Template.xlsx", company_id=company.id
            )
            assert preview[0]["Validation"] == "Ready"
            assert preview[0]["Related Policy ID"] == policy_ref
            assert preview[0]["_values"]["related_policy_id"] == policy.id
    finally:
        engine.dispose()


def test_employee_manager_leader_similar_reference_uses_same_display_standard() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, _admin, maria, _pedro, _policy, _violation = _seed(session)
            service = EmployeeBulkImportService(session)
            workbook = load_workbook(BytesIO(service.build_company_template(company_id=company.id)))
            sheet = workbook["Employee Import Template"]
            maria_ref = "EMP-001 - Santos, Maria Jr."
            assert maria_ref in sheet["L1"].comment.text
            assert maria_ref in sheet["M1"].comment.text
            refs = workbook["Employee References"]
            values = [str(cell.value or "") for cell in refs["B"]]
            assert maria_ref in values
            validations = {
                str(item.sqref): str(item.formula1)
                for item in sheet.data_validations.dataValidation
            }
            assert validations["L2:L2000 M2:M2000"] == "=EmployeeAssignmentReferences"

            row = [
                "EMP-NEW-001", "Dela Cruz", "Ana", "", "", "Female", "Single",
                "1998-01-01", "ana.new@example.com", "", "Operations", maria_ref, "",
                "Analyst", "Employed", "2026-01-01", "", "",
            ]
            for index, value in enumerate(row, start=1):
                sheet.cell(2, index).value = value
            out = BytesIO(); workbook.save(out)
            preview = service.prepare_preview(out.getvalue(), filename="Employee_Import_Template.xlsx", company_id=company.id)
            assert preview[0]["Validation"] == "Ready"
            assert preview[0]["_employee"]["manager_employee_number"] == maria.employee_number
    finally:
        engine.dispose()
