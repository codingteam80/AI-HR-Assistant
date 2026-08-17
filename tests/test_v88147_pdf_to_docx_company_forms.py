"""Regression coverage for v8.8.147 editable company-form downloads."""

from io import BytesIO
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from modules.documents.pdf_to_docx_converter import (
    DOCX_MIME_TYPE,
    convert_pdf_bytes_to_docx,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _sample_pdf() -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=letter)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(72, 720, "Employee Information Form")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(72, 680, "Employee Name:")
    pdf.line(175, 678, 500, 678)
    pdf.drawString(72, 645, "Department:")
    pdf.line(155, 643, 500, 643)
    pdf.save()
    return output.getvalue()


def test_pdf_is_converted_to_a_valid_editable_docx() -> None:
    source = _sample_pdf()
    filename, mime_type, converted = convert_pdf_bytes_to_docx(
        pdf_bytes=source,
        source_filename="Employee Information.pdf",
    )

    assert filename == "Employee Information.docx"
    assert mime_type == DOCX_MIME_TYPE
    assert converted.startswith(b"PK")
    assert source.startswith(b"%PDF")

    document = Document(BytesIO(converted))
    visible_text = " ".join(paragraph.text for paragraph in document.paragraphs)
    visible_text += " " + " ".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Employee Information Form" in visible_text
    assert "Employee Name" in visible_text


def test_empty_pdf_is_rejected_with_a_controlled_message() -> None:
    try:
        convert_pdf_bytes_to_docx(
            pdf_bytes=b"",
            source_filename="empty.pdf",
        )
    except ValueError as error:
        assert "empty" in str(error).lower()
    else:  # pragma: no cover - defensive assertion
        raise AssertionError("An empty PDF must not produce a download.")


def test_malformed_pdf_is_rejected_without_crashing_the_page() -> None:
    try:
        convert_pdf_bytes_to_docx(
            pdf_bytes=b"not a valid pdf",
            source_filename="broken.pdf",
        )
    except ValueError as error:
        assert "could not be converted" in str(error).lower()
    else:  # pragma: no cover - defensive assertion
        raise AssertionError("A malformed PDF must not produce a download.")


def test_admin_and_employee_buttons_use_download_form_label() -> None:
    sources = [
        (PROJECT_ROOT / "ui/pages/admin/company_forms_documents_page.py").read_text(
            encoding="utf-8"
        ),
        (PROJECT_ROOT / "ui/pages/user/company_forms_documents_page.py").read_text(
            encoding="utf-8"
        ),
    ]
    assert all('"Download Form"' in source for source in sources)
    assert all('"Download Original Form"' not in source for source in sources)
    assert all("prepare_editable_company_form_download" in source for source in sources)


def test_submission_download_remains_unconverted() -> None:
    user_source = (
        PROJECT_ROOT / "ui/pages/user/company_forms_documents_page.py"
    ).read_text(encoding="utf-8")
    submission_block = user_source.split("def _render_my_documents", 1)[1]
    assert "get_submission_download" in submission_block
    assert "prepare_editable_company_form_download" not in submission_block
    assert '"Download My Submitted Copy"' in submission_block


def test_release_dependencies_and_warning_guards() -> None:
    requirements = (PROJECT_ROOT / "requirements.txt").read_text(encoding="utf-8")
    settings = (PROJECT_ROOT / "config/settings.py").read_text(encoding="utf-8")
    assert "pdf2docx>=0.5.8,<0.6.0" in requirements
    assert "PyMuPDF>=1.26.7,<1.27.0" in requirements
    assert 'app_version: str = "0.8.8.147"' in settings

    files = [
        path
        for path in PROJECT_ROOT.rglob("*.py")
        if "tests" not in path.parts
        and ".venv" not in path.parts
        and "venv" not in path.parts
    ]
    source = "\n".join(path.read_text(encoding="utf-8") for path in files)
    assert "st.components.v1.html" not in source
    assert "components.html" not in source
    assert "use_container_width" not in source
