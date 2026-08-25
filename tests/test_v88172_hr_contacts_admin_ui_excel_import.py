"""v8.8.172 HR Contacts admin UI, automatic ordering, and Excel import checks."""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from database.base import Base
from models.company import Company
from services.hr_contact_bulk_import_service import HRContactBulkImportService
from services.hr_contact_service import HRContactService


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _company(session, code: str = "CONTACTS") -> Company:
    company = Company(code=code, name=f"{code} Company", is_active=True)
    session.add(company)
    session.commit()
    session.refresh(company)
    return company


def test_display_order_is_automatic_and_edit_preserves_existing_order() -> None:
    factory = _factory()
    with factory() as session:
        company = _company(session)
        service = HRContactService(session)
        first = service.create_contact(
            company_id=company.id,
            values={
                "name": "First Contact",
                "email": "first@example.com",
                "sort_order": 999,
            },
        )
        second = service.create_contact(
            company_id=company.id,
            values={"name": "Second Contact", "phone": "+63 900 111 2222"},
        )

        assert first.sort_order == 10
        assert second.sort_order == 20

        updated = service.update_contact(
            company_id=company.id,
            contact_id=first.id,
            values={
                "name": "First Contact Updated",
                "email": "first@example.com",
                "sort_order": 777,
            },
        )
        assert updated.sort_order == 10


def test_hr_contact_excel_template_preview_and_atomic_import() -> None:
    factory = _factory()
    with factory() as session:
        company = _company(session, "BULK")
        service = HRContactBulkImportService(session)
        template = service.build_template()
        workbook = load_workbook(BytesIO(template))
        sheet = workbook["HR Contact Import Template"]
        sheet["A2"] = "Maria Santos"
        sheet["B2"] = "HR Manager"
        sheet["C2"] = "Human Resources"
        sheet["D2"] = "maria.santos@example.com"
        sheet["E2"] = "+63 912 345 6789"
        sheet.append(
            [
                "John Reyes",
                "Payroll Specialist",
                "Finance / Payroll",
                "john.reyes@example.com",
                "",
                "2nd Floor",
                "Monday-Friday",
                "Payroll support",
            ]
        )
        output = BytesIO()
        workbook.save(output)

        preview = service.prepare_preview(
            output.getvalue(),
            filename="contacts.xlsx",
            company_id=company.id,
        )
        assert len(preview) == 2
        assert all(row["Validation"] == "Ready" for row in preview)

        contacts = service.import_preview_rows(preview, company_id=company.id)
        assert [contact.name for contact in contacts] == ["Maria Santos", "John Reyes"]
        assert [contact.sort_order for contact in contacts] == [10, 20]
        assert len(HRContactService(session).list_contacts(company.id, active_only=True)) == 2


def test_hr_contact_excel_preview_blocks_invalid_contact_rows() -> None:
    factory = _factory()
    with factory() as session:
        company = _company(session, "INVALID")
        service = HRContactBulkImportService(session)
        workbook = load_workbook(BytesIO(service.build_template()))
        sheet = workbook["HR Contact Import Template"]
        sheet["A2"] = "No Contact Method"
        sheet["D2"] = ""
        sheet["E2"] = ""
        output = BytesIO()
        workbook.save(output)

        preview = service.prepare_preview(
            output.getvalue(),
            filename="invalid.xlsx",
            company_id=company.id,
        )
        assert len(preview) == 1
        assert "email address or phone" in preview[0]["Validation"]


def test_hr_contacts_admin_ui_uses_three_row_table_horizontal_forms_and_excel_upload() -> None:
    management = _source("ui/pages/admin/hr_contacts_management.py")

    assert 'max_height=255' in management
    assert 'st.number_input(' not in management
    assert 'st.columns([1.7, 4.3])' in management
    assert 'st.columns(2)' in management
    assert 'st.columns([1.7, 1.7, 2.6])' in management
    assert '"Upload HR Contacts via Excel"' in management
    assert '"Download HR Contact Excel Template"' in management
    assert '"Validate and Preview HR Contacts"' in management
    assert '"Import HR Contacts"' in management


def test_current_version_is_v88172() -> None:
    settings = _source("config/settings.py")
    env = _source(".env.example")

    assert 'app_version: str = "0.8.8.172"' in settings
    assert 'APP_VERSION=0.8.8.172' in env
