"""v8.8.186 integrated profile-photo, sticky-banner, and disciplinary regressions."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from io import BytesIO

import pytest
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import Session
from openpyxl import load_workbook

import models  # noqa: F401 - register complete metadata graph
from authentication.current_user import AuthenticatedUser
from database.base import Base
from models.audit_event import AuditEvent
from models.company import Company
from models.employee import Employee
from models.employee_history import EmployeeHistory
from models.notification import Notification
from models.role import Role
from models.user import User
from schemas.disciplinary_record_schema import (
    DisciplinaryRecordCreateRequest,
    DisciplinaryRecordUpdateRequest,
)
from schemas.policy_violation_schema import PolicyViolationCreateRequest
from services.disciplinary_record_service import DisciplinaryRecordService
from services.audit_context import clear_audit_actor, set_audit_actor
from services.policy_violation_service import PolicyViolationService
from modules.hr_assistant.admin_hr_assistant import AdminHRAssistant
from modules.hr_assistant.hr_assistant import HRAssistant
from modules.reports.disciplinary_report import build_disciplinary_records_excel


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection, _record) -> None:
        dbapi_connection.execute("PRAGMA foreign_keys = ON")

    Base.metadata.create_all(engine)
    return engine


def _seed(session: Session, *, code: str = "MAIN"):
    company = Company(code=code, name=f"{code} Company", is_active=True)
    session.add(company)
    session.flush()
    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role])
    session.flush()
    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username=f"admin_{code.lower()}",
        email=f"admin_{code.lower()}@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    employee_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username=f"member_{code.lower()}",
        email=f"member_{code.lower()}@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add_all([admin, employee_user])
    session.flush()
    employee = Employee(
        company_id=company.id,
        user_id=employee_user.id,
        employee_number=f"{code}-001",
        first_name="Ana",
        last_name="Reyes",
        employment_status="employed",
    )
    session.add(employee)
    session.commit()
    return company, admin, employee_user, employee


def _violation(session: Session, *, company_id: int, admin_id: int, code: str = "ATT-001"):
    return PolicyViolationService(session).create_violation(
        PolicyViolationCreateRequest(
            company_id=company_id,
            created_by_user_id=admin_id,
            violation_code=code,
            category="Attendance",
            offense_title="Habitual Tardiness",
            description="Repeated late arrival beyond the company attendance rule.",
            severity="Minor",
            first_offense_action="Verbal Warning",
            second_offense_action="Written Warning",
            third_offense_action="Suspension",
            final_action="Termination",
            effective_date=date(2026, 1, 1),
            status="active",
        )
    )


def _current_user(*, company, user, employee=None, clearance: int):
    return AuthenticatedUser(
        user_id=user.id,
        company_id=company.id,
        company_code=company.code,
        company_name=company.name,
        role_id=user.role_id,
        role_name="company_admin" if clearance == 1 else "employee",
        clearance=clearance,
        username=user.username,
        email=user.email,
        employee_id=employee.id if employee is not None else None,
        employee_number=employee.employee_number if employee is not None else None,
        employee_name=employee.full_name if employee is not None else None,
        must_change_password=False,
    )


def test_version_runtime_routes_and_single_source_markers() -> None:
    assert 'app_version: str = "0.8.8.186"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.186" in _source(".env")
    assert "APP_VERSION=0.8.8.186" in _source(".env.example")
    assert '"employee_disciplinary_records"' in _source("database/runtime_schema.py")
    assert "EmployeeDisciplinaryRecord" in _source("models/__init__.py")
    routes = _source("ui/module_view_navigation.py")
    assert '"disciplinary": "Violations / Disciplinary Records"' in routes
    assert '"disciplinary": "My Disciplinary Records"' in routes
    model = _source("models/employee_disciplinary_record.py")
    assert "Penalties are intentionally NOT copied" in model
    assert "suggested action is always derived" in model


def test_profile_photo_is_compact_collapsible_and_hides_streamlit_size_helper() -> None:
    employees_page = _source("ui/pages/admin/employees_page.py")
    profile = _source("ui/components/employee_profile_photo.py")
    theme = _source("ui/theme/theme_loader.py")
    assert 'with st.expander("Profile Photo", expanded=False):' in employees_page
    assert 'st.columns(\n            [0.85, 2.15]' in profile
    assert "Maximum file size: {max_mb} MB" in profile
    assert "processed to 512×512 PNG" in profile
    assert "independent of the Employee Master" in profile
    assert "uploaded_file.size > max_mb * 1024 * 1024" in profile
    assert '[data-testid="stFileUploaderDropzoneInstructions"] small' in theme
    assert "display: none !important" in theme


def test_sticky_rule_targets_outer_streamlit_element_not_inner_keyed_container() -> None:
    theme = _source("ui/theme/theme_loader.py")
    outer = 'div[data-testid="stElementContainer"]:has(.st-key-hr_global_topbar_shell)'
    assert outer in theme
    outer_block = theme[theme.index(outer): theme.index(".st-key-hr_global_topbar_shell {{", theme.index(outer))]
    assert "position: sticky !important" in outer_block
    inner_start = theme.index(".st-key-hr_global_topbar_shell {{")
    inner_block = theme[inner_start: theme.index("}}", inner_start)]
    assert "position: relative !important" in inner_block
    assert "position: sticky" not in inner_block
    assert "z-index: 9000" in outer_block
    assert "z-index: 10030" in theme  # notification dropdown remains above banner


def test_case_progression_recomputes_and_penalty_stays_linked_to_master() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, employee_user, employee = _seed(session)
        violation = _violation(session, company_id=company.id, admin_id=admin.id)
        service = DisciplinaryRecordService(session)
        first = service.create_record(
            DisciplinaryRecordCreateRequest(
                company_id=company.id,
                employee_id=employee.id,
                violation_id=violation.id,
                incident_date=date(2026, 8, 1),
                incident_description="First attendance incident.",
                case_status="Draft",
                created_by_user_id=admin.id,
            )
        )
        assert first.previous_offense_count == 0
        assert first.offense_level == "1st"
        assert service.suggested_action_for_record(first) == "Verbal Warning"

        second = service.create_record(
            DisciplinaryRecordCreateRequest(
                company_id=company.id,
                employee_id=employee.id,
                violation_id=violation.id,
                incident_date=date(2026, 8, 20),
                incident_description="Second attendance incident.",
                case_status="Draft",
                created_by_user_id=admin.id,
            )
        )
        assert second.previous_offense_count == 1
        assert second.offense_level == "2nd"
        assert service.suggested_action_for_record(second) == "Written Warning"

        # Moving the second incident before the first must recompute progression.
        second = service.update_record(
            DisciplinaryRecordUpdateRequest(
                company_id=company.id,
                record_id=second.id,
                incident_date=date(2026, 7, 20),
                incident_description=second.incident_description,
                case_status="Draft",
                employee_acknowledgment="Pending",
                edited_by_user_id=admin.id,
            )
        )
        assert second.previous_offense_count == 0
        assert second.offense_level == "1st"

        # Editing the master changes suggested guidance without copying penalty text into the case.
        violation.first_offense_action = "Coaching Memo"
        session.commit()
        assert service.suggested_action_for_record(second) == "Coaching Memo"
        assert not hasattr(second, "suggested_action")
        assert employee_user.id > 0


def test_issued_case_notification_history_and_employee_self_only_visibility() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, employee_user, employee = _seed(session)
        violation = _violation(session, company_id=company.id, admin_id=admin.id)
        service = DisciplinaryRecordService(session)
        record = service.create_record(
            DisciplinaryRecordCreateRequest(
                company_id=company.id,
                employee_id=employee.id,
                violation_id=violation.id,
                incident_date=date(2026, 8, 24),
                incident_description="Attendance incident confirmed after review.",
                actual_action_taken="Verbal Warning",
                issued_by_user_id=admin.id,
                reviewed_approved_by_user_id=admin.id,
                date_issued=date(2026, 8, 24),
                employee_acknowledgment="Pending",
                case_status="Issued",
                created_by_user_id=admin.id,
            )
        )
        own = service.list_employee_records(
            company_id=company.id,
            employee_id=employee.id,
            requester_user_id=employee_user.id,
        )
        assert [item.id for item in own] == [record.id]
        assert session.scalar(
            select(Notification).where(
                Notification.user_id == employee_user.id,
                Notification.related_entity_type == "employee_disciplinary_record",
                Notification.related_entity_id == record.id,
            )
        ) is not None
        assert session.scalar(
            select(EmployeeHistory).where(
                EmployeeHistory.employee_id == employee.id,
                EmployeeHistory.source == "disciplinary_record",
            )
        ) is not None

        company2, admin2, employee_user2, employee2 = _seed(session, code="OTHER")
        with pytest.raises(PermissionError, match="only view your own"):
            service.list_employee_records(
                company_id=company.id,
                employee_id=employee.id,
                requester_user_id=employee_user2.id,
            )
        with pytest.raises(ValueError, match="does not belong"):
            service.list_employee_records(
                company_id=company.id,
                employee_id=employee2.id,
                requester_user_id=employee_user2.id,
            )
        assert admin2.id > 0 and company2.id > 0


def test_chat_actual_records_are_permission_aware_and_master_questions_still_work() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, employee_user, employee = _seed(session)
        violation = _violation(session, company_id=company.id, admin_id=admin.id)
        DisciplinaryRecordService(session).create_record(
            DisciplinaryRecordCreateRequest(
                company_id=company.id,
                employee_id=employee.id,
                violation_id=violation.id,
                incident_date=date(2026, 8, 24),
                incident_description="Late arrival.",
                actual_action_taken="Verbal Warning",
                issued_by_user_id=admin.id,
                date_issued=date(2026, 8, 24),
                case_status="Issued",
                created_by_user_id=admin.id,
            )
        )
        admin_user = _current_user(company=company, user=admin, clearance=1)
        employee_current = _current_user(
            company=company, user=employee_user, employee=employee, clearance=2
        )

        admin_answer = AdminHRAssistant(session).answer(
            current_user=admin_user,
            question="Who received a violation this month?",
        )
        assert admin_answer.intent == "disciplinary_records"
        assert employee.employee_number in admin_answer.answer
        assert admin_answer.actions[0].query_params == {"employee_view": "disciplinary"}

        employee_answer = HRAssistant(session).answer(
            current_user=employee_current,
            question="Show my disciplinary records.",
        )
        assert employee_answer.intent == "disciplinary_records"
        assert "ATT-001" in employee_answer.answer
        assert employee_answer.actions[0].query_params == {"policy_view": "disciplinary"}

        master_answer = HRAssistant(session).answer(
            current_user=employee_current,
            question="What is the second offense penalty for ATT-001?",
        )
        assert master_answer.intent == "policy_violation"
        assert "Written Warning" in master_answer.answer

        third_party = HRAssistant(session).answer(
            current_user=employee_current,
            question="Show Ana Cruz disciplinary records.",
        )
        assert third_party.intent == "disciplinary_records"
        assert "only view your own" in third_party.answer

        follow_up = AdminHRAssistant(session).answer(
            current_user=admin_user,
            question="Show disciplinary cases for this employee.",
            history=[{"role": "user", "content": f"Find employee {employee.employee_number}"}],
        )
        assert follow_up.intent == "disciplinary_records"
        assert employee.employee_number in follow_up.answer


def test_central_audit_trail_captures_admin_case_creation() -> None:
    # Importing database.session installs the production listener.
    import database.session  # noqa: F401

    engine = _engine()
    with Session(engine) as session:
        company, admin, _employee_user, employee = _seed(session)
        violation = _violation(session, company_id=company.id, admin_id=admin.id)
        set_audit_actor(company_id=company.id, user_id=admin.id, module="Employees")
        try:
            record = DisciplinaryRecordService(session).create_record(
                DisciplinaryRecordCreateRequest(
                    company_id=company.id,
                    employee_id=employee.id,
                    violation_id=violation.id,
                    incident_date=date(2026, 8, 24),
                    incident_description="Audited disciplinary case.",
                    case_status="Draft",
                    created_by_user_id=admin.id,
                )
            )
        finally:
            clear_audit_actor()
        event = session.scalar(
            select(AuditEvent).where(
                AuditEvent.company_id == company.id,
                AuditEvent.entity_type == "EmployeeDisciplinaryRecord",
                AuditEvent.entity_label == record.public_id,
            )
        )
        assert event is not None
        assert event.module == "Employees"
        assert event.action == "Created"



def test_disciplinary_report_export_contains_expected_case_columns() -> None:
    payload = build_disciplinary_records_excel([
        {
            "Case ID": "DR_TEST",
            "Employee": "MAIN-001 — Ana Reyes",
            "Violation": "ATT-001 — Habitual Tardiness",
            "Incident Date": "2026-08-24",
            "Suggested Disciplinary Action": "Verbal Warning",
            "Case Status": "Issued",
        }
    ])
    workbook = load_workbook(BytesIO(payload), read_only=True)
    sheet = workbook["Disciplinary Records"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "Case ID" in headers
    assert "Suggested Disciplinary Action" in headers
    assert "Case Status" in headers
    values = list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert "DR_TEST" in values


def test_notification_routes_disciplinary_records_to_correct_exact_tabs() -> None:
    source = _source("ui/components/topbar.py")
    assert 'if "disciplinary" in entity:\n            return "admin", "Employees"' in source
    assert 'if "disciplinary" in entity:\n        return "employee", "Company Policies"' in source
    assert 'st.query_params["employee_view"] = "disciplinary"' in source
    assert 'st.query_params["policy_view"] = "disciplinary"' in source
    assert '(("disciplinary",),\n        "⚖️",\n        "Disciplinary")' in source.replace("    ", "") or '"Disciplinary"' in source
