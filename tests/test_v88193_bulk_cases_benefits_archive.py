"""v8.8.193 bulk disciplinary/benefit imports and recoverable case archive."""

from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, event, inspect
from sqlalchemy.orm import Session

import models  # noqa: F401
from database.base import Base
from database.schema_upgrade import upgrade_existing_schema
from models.company import Company
from models.department import Department
from models.employee import Employee
from models.role import Role
from models.user import User
from schemas.policy_violation_schema import PolicyViolationCreateRequest
from schemas.disciplinary_record_schema import DisciplinaryRecordCreateRequest
from services.benefit_bulk_import_service import BenefitBulkImportService
from services.disciplinary_record_bulk_import_service import DisciplinaryRecordBulkImportService
from services.disciplinary_record_service import ACKNOWLEDGMENT_STATUSES, CASE_STATUSES, DisciplinaryRecordService
from services.onboarding_service import OnboardingService
from services.policy_violation_service import PolicyViolationService
from ui.onboarding_destinations import ONBOARDING_DESTINATIONS


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")

    @event.listens_for(engine, "connect")
    def _enable_fk(dbapi_connection, _record) -> None:
        dbapi_connection.execute("PRAGMA foreign_keys = ON")

    Base.metadata.create_all(engine)
    return engine


def _seed_company(session: Session, *, code: str):
    company = Company(code=code, name=f"{code} Company", is_active=True)
    session.add(company)
    session.flush()
    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role])
    session.flush()
    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username=f"admin_{code.lower()}",
        email=f"admin_{code.lower()}@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    employee_user = User(
        company_id=company.id,
        role_id=employee_role.id,
        clearance=2,
        username=f"member_{code.lower()}",
        email=f"member_{code.lower()}@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add_all([admin, employee_user])
    session.flush()
    department = Department(company_id=company.id, name=f"{code} Department", is_active=True)
    session.add(department)
    session.flush()
    employee = Employee(
        company_id=company.id,
        user_id=employee_user.id,
        employee_number=f"{code}-001",
        first_name="Ana",
        last_name="Reyes",
        employment_status="employed",
        department_id=department.id,
    )
    session.add(employee)
    session.commit()
    violation = PolicyViolationService(session).create_violation(
        PolicyViolationCreateRequest(
            company_id=company.id,
            created_by_user_id=admin.id,
            violation_code=f"{code}-V001",
            category="Attendance",
            offense_title="Habitual Tardiness",
            description="Repeated late arrival.",
            severity="Minor",
            first_offense_action="Verbal Warning",
            second_offense_action="Written Warning",
            third_offense_action="Suspension",
            final_action="Termination",
            status="active",
        )
    )
    return company, admin, employee_user, employee, violation


def test_version_and_admin_ui_contains_requested_bulk_archive_controls() -> None:
    assert 'app_version: str = "0.8.8.193"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.193" in _source(".env")
    disciplinary_ui = _source("ui/components/disciplinary_records.py")
    assert "Bulk Add Disciplinary Cases via Excel" in disciplinary_ui
    assert "Download Disciplinary Case Excel Template" in disciplinary_ui
    assert "Delete / Move Case to Archive" in disciplinary_ui
    assert "Restore Disciplinary Case" in disciplinary_ui
    benefits_ui = _source("ui/pages/admin/onboarding_management.py")
    assert "Bulk Add Benefits via Excel" in benefits_ui
    assert "Download Benefit Excel Template" in benefits_ui


def test_disciplinary_template_guidance_is_complete_and_company_scoped() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, _member, employee, violation = _seed_company(session, code="MAIN")
        other, other_admin, _other_member, other_employee, other_violation = _seed_company(session, code="OTHER")
        payload = DisciplinaryRecordBulkImportService(session).build_template(company_id=company.id)
        employee_number = employee.employee_number
        other_employee_number = other_employee.employee_number
        violation_code = violation.violation_code
        other_violation_code = other_violation.violation_code
        admin_username = admin.username
        other_admin_username = other_admin.username

    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["Disciplinary Case Import"]
    assert all(cell.comment is not None for cell in sheet[1])
    for choice in ACKNOWLEDGMENT_STATUSES:
        assert choice in sheet["K1"].comment.text
    for choice in CASE_STATUSES:
        assert choice in sheet["L1"].comment.text
    assert "Employee References" in sheet["A1"].comment.text
    assert "Violation References" in sheet["B1"].comment.text
    assert "User References" in sheet["H1"].comment.text

    employee_numbers = {str(cell.value or "") for cell in workbook["Employee References"]["A"]}
    violation_codes = {str(cell.value or "") for cell in workbook["Violation References"]["A"]}
    usernames = {str(cell.value or "") for cell in workbook["User References"]["A"]}
    assert employee_number in employee_numbers
    assert other_employee_number not in employee_numbers
    assert violation_code in violation_codes
    assert other_violation_code not in violation_codes
    assert admin_username in usernames
    assert other_admin_username not in usernames

    validations = {str(item.sqref): str(item.formula1) for item in sheet.data_validations.dataValidation}
    assert validations["A2:A2000"] == "=DisciplinaryEmployeeNumbers"
    assert validations["B2:B2000"] == "=DisciplinaryViolationCodes"
    assert validations["H2:H2000"] == "=DisciplinaryUsernames"
    assert "K2:K2000" in validations and "L2:L2000" in validations
    engine.dispose()


def test_disciplinary_bulk_import_is_atomic_and_archive_restore_hides_employee_view() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, employee_user, employee, violation = _seed_company(session, code="CASE")
        service = DisciplinaryRecordBulkImportService(session)
        workbook = load_workbook(BytesIO(service.build_template(company_id=company.id)))
        sheet = workbook["Disciplinary Case Import"]
        values = [
            employee.employee_number,
            violation.violation_code,
            "2026-08-24",
            "First imported case.",
            "",
            "",
            "Verbal Warning",
            admin.username,
            "",
            "2026-08-24",
            "Pending",
            "Issued",
            "",
        ]
        for index, value in enumerate(values, start=1):
            sheet.cell(2, index).value = value
        stream = BytesIO()
        workbook.save(stream)
        preview = service.prepare_preview(stream.getvalue(), filename="cases.xlsx", company_id=company.id)
        assert preview[0]["Validation"] == "Ready"
        created = service.import_preview_rows(preview, company_id=company.id, actor_user_id=admin.id)
        assert len(created) == 1
        record = created[0]
        assert record.previous_offense_count == 0
        assert record.offense_level == "1st"

        case_service = DisciplinaryRecordService(session)
        assert len(case_service.list_employee_records(
            company_id=company.id,
            employee_id=employee.id,
            requester_user_id=employee_user.id,
        )) == 1
        case_service.archive_record(company_id=company.id, record_id=record.id, actor_user_id=admin.id)
        assert case_service.list_admin_records(company_id=company.id, requester_user_id=admin.id) == []
        assert len(case_service.list_admin_archived_records(company_id=company.id, requester_user_id=admin.id)) == 1
        assert case_service.list_employee_records(
            company_id=company.id,
            employee_id=employee.id,
            requester_user_id=employee_user.id,
        ) == []
        case_service.restore_record(company_id=company.id, record_id=record.id, actor_user_id=admin.id)
        assert len(case_service.list_admin_records(company_id=company.id, requester_user_id=admin.id)) == 1

        # Revalidating the same workbook blocks an existing duplicate case.
        duplicate_preview = service.prepare_preview(stream.getvalue(), filename="cases.xlsx", company_id=company.id)
        assert "matching active disciplinary case" in duplicate_preview[0]["Validation"]
    engine.dispose()


def test_benefit_template_guidance_dropdown_and_atomic_duplicate_protection() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, _admin, _member, _employee, _violation = _seed_company(session, code="BEN")
        service = BenefitBulkImportService(session)
        payload = service.build_template()
        workbook = load_workbook(BytesIO(payload))
        sheet = workbook["Benefit Import Template"]
        assert all(cell.comment is not None for cell in sheet[1])
        for choice in ONBOARDING_DESTINATIONS:
            assert choice in sheet["I1"].comment.text
        validations = {str(item.sqref): str(item.formula1) for item in sheet.data_validations.dataValidation}
        assert validations["I2:I2000"] == "=BenefitRelatedWorkspaces"

        values = [
            "Medical Coverage",
            "Health",
            "Company medical coverage.",
            "All employees",
            "2026-08-24",
            "Contact HR.",
            "HR Team",
            10,
            "Onboarding — Benefits",
        ]
        for index, value in enumerate(values, start=1):
            sheet.cell(2, index).value = value
        stream = BytesIO()
        workbook.save(stream)
        preview = service.prepare_preview(stream.getvalue(), filename="benefits.xlsx", company_id=company.id)
        assert preview[0]["Validation"] == "Ready"
        created = service.import_preview_rows(preview, company_id=company.id)
        assert len(created) == 1
        assert created[0].is_active is True
        assert created[0].target_page == "Onboarding"

        duplicate = service.prepare_preview(stream.getvalue(), filename="benefits.xlsx", company_id=company.id)
        assert "already exists" in duplicate[0]["Validation"]
        benefits = OnboardingService(session).list_benefits(company.id)
        assert len([item for item in benefits if item.name == "Medical Coverage"]) == 1
    engine.dispose()




def test_archive_restore_recalculates_later_offense_progression() -> None:
    engine = _engine()
    with Session(engine) as session:
        company, admin, _employee_user, employee, violation = _seed_company(session, code="SEQ")
        service = DisciplinaryRecordService(session)
        records = []
        for day in (1, 10, 20):
            records.append(
                service.create_record(
                    DisciplinaryRecordCreateRequest(
                        company_id=company.id,
                        employee_id=employee.id,
                        violation_id=violation.id,
                        incident_date=date(2026, 8, day),
                        incident_description=f"Incident on day {day}.",
                        case_status="Draft",
                        created_by_user_id=admin.id,
                    )
                )
            )
        assert [item.offense_level for item in records] == ["1st", "2nd", "3rd"]

        service.archive_record(
            company_id=company.id, record_id=records[0].id, actor_user_id=admin.id
        )
        active = service.list_admin_records(
            company_id=company.id, requester_user_id=admin.id
        )
        active_by_id = {item.id: item for item in active}
        assert active_by_id[records[1].id].offense_level == "1st"
        assert active_by_id[records[2].id].offense_level == "2nd"

        service.restore_record(
            company_id=company.id, record_id=records[0].id, actor_user_id=admin.id
        )
        active = service.list_admin_records(
            company_id=company.id, requester_user_id=admin.id
        )
        active_by_id = {item.id: item for item in active}
        assert active_by_id[records[0].id].offense_level == "1st"
        assert active_by_id[records[1].id].offense_level == "2nd"
        assert active_by_id[records[2].id].offense_level == "3rd"
    engine.dispose()


def test_existing_disciplinary_table_upgrades_archive_columns_without_row_loss(tmp_path: Path) -> None:
    copied_db = tmp_path / "legacy_disciplinary.db"
    engine = create_engine(f"sqlite+pysqlite:///{copied_db}")
    with engine.begin() as connection:
        connection.exec_driver_sql(
            """
            CREATE TABLE employee_disciplinary_records (
                id INTEGER PRIMARY KEY,
                company_id INTEGER NOT NULL,
                public_id VARCHAR(30) NOT NULL
            )
            """
        )
        connection.exec_driver_sql(
            "INSERT INTO employee_disciplinary_records (id, company_id, public_id) "
            "VALUES (1, 1, 'DR_LEGACY')"
        )
    before_columns = {
        column["name"]
        for column in inspect(engine).get_columns("employee_disciplinary_records")
    }
    assert "archived_at" not in before_columns

    upgrade_existing_schema(engine)
    after_columns = {
        column["name"]
        for column in inspect(engine).get_columns("employee_disciplinary_records")
    }
    assert {"archived_at", "archived_by_user_id"}.issubset(after_columns)
    with engine.connect() as connection:
        assert connection.exec_driver_sql(
            "SELECT COUNT(*) FROM employee_disciplinary_records"
        ).scalar_one() == 1
    engine.dispose()

