"""v8.8.196 policy Excel references prefer filename and fall back to title."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

import models  # noqa: F401
from database.base import Base
from models.company import Company
from models.hr_policy import HRPolicy
from models.hr_policy_document import HRPolicyDocument
from models.role import Role
from models.user import User
from services.policy_violation_bulk_import_service import PolicyViolationBulkImportService


def _engine():
    engine = create_engine(
        "sqlite+pysqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
        future=True,
    )
    Base.metadata.create_all(engine)
    return engine


def _seed_company_user(session: Session) -> tuple[Company, User]:
    company = Company(name="Template Test Co", code="TPL196")
    session.add(company)
    session.flush()
    role = Role(company_id=company.id, name="company_admin", is_active=True)
    session.add(role)
    session.flush()
    user = User(
        company_id=company.id,
        role_id=role.id,
        username="admin196",
        email="admin196@example.com",
        password_hash="unused",
        clearance=1,
        is_active=True,
        must_change_password=False,
    )
    session.add(user)
    session.flush()
    return company, user




def test_version_markers() -> None:
    root = Path(__file__).resolve().parents[1]
    assert 'app_version: str = "0.8.8.196"' in (root / "config/settings.py").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.196" in (root / ".env").read_text(encoding="utf-8")


def test_related_policy_reference_uses_filename_when_available_and_title_when_not() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, user = _seed_company_user(session)
            with_file = HRPolicy(
                public_id="POL-FILE-001",
                company_id=company.id,
                created_by_user_id=user.id,
                title="Attendance Policy",
                category="Attendance",
                content="Attendance rules.",
                version="1.0",
                status="published",
            )
            title_only = HRPolicy(
                public_id="POL-TITLE-001",
                company_id=company.id,
                created_by_user_id=user.id,
                title="Manual Conduct Policy",
                category="Conduct",
                content="Conduct rules.",
                version="1.0",
                status="published",
            )
            session.add_all([with_file, title_only])
            session.flush()
            session.add(
                HRPolicyDocument(
                    company_id=company.id,
                    policy_id=with_file.id,
                    uploaded_by_user_id=user.id,
                    original_filename="Attendance_Policy.pdf",
                    stored_filename="stored.pdf",
                    storage_path="policies/stored.pdf",
                    mime_type="application/pdf",
                    file_extension=".pdf",
                    sha256="a" * 64,
                    size_bytes=123,
                    extracted_text="Attendance rules.",
                )
            )
            session.commit()

            service = PolicyViolationBulkImportService(session)
            workbook = load_workbook(BytesIO(service.build_template(company_id=company.id)))
            sheet = workbook["Violation Import Template"]
            refs = workbook["Related Policies"]
            rows = {
                str(refs.cell(row, 1).value): str(refs.cell(row, 2).value)
                for row in range(2, refs.max_row + 1)
            }

            assert rows["POL-FILE-001"] == "POL-FILE-001 - Attendance_Policy.pdf"
            assert rows["POL-TITLE-001"] == "POL-TITLE-001 - Manual Conduct Policy"
            assert "POL-FILE-001 - Attendance_Policy.pdf" in sheet["J1"].comment.text
            assert "POL-TITLE-001 - Manual Conduct Policy" in sheet["J1"].comment.text
            assert "Policy ID - Filename / Title (whichever is available)" in sheet["J1"].comment.text

            # The dedicated filename column remains explicit when no source file exists.
            filename_by_id = {
                str(refs.cell(row, 1).value): str(refs.cell(row, 3).value)
                for row in range(2, refs.max_row + 1)
            }
            assert filename_by_id["POL-TITLE-001"] == "[No uploaded file]"

            # The readable title fallback must resolve back to the same internal policy ID.
            sheet["A2"] = "ATT-TITLE-001"
            sheet["B2"] = "Attendance"
            sheet["C2"] = "Title fallback test"
            sheet["D2"] = "Minor"
            sheet["E2"] = "Test description"
            sheet["F2"] = "Verbal Warning"
            sheet["G2"] = "Written Warning"
            sheet["H2"] = "Suspension"
            sheet["I2"] = "Termination"
            sheet["J2"] = "POL-TITLE-001 - Manual Conduct Policy"
            sheet["K2"] = "2026-08-25"
            sheet["L2"] = "Active"
            sheet["M2"] = ""
            output = BytesIO()
            workbook.save(output)
            preview = service.prepare_preview(
                output.getvalue(),
                filename="Violation_Import_Template.xlsx",
                company_id=company.id,
            )
            assert preview[0]["Validation"] == "Ready"
            assert preview[0]["Related Policy ID"] == "POL-TITLE-001 - Manual Conduct Policy"
            assert preview[0]["_values"]["related_policy_id"] == title_only.id
    finally:
        engine.dispose()
