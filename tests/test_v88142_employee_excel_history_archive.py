"""Regression coverage for v8.8.142 Employee Excel, History, and Archive."""

from datetime import date
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import create_engine, inspect, select
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from config.settings import Settings
from database.base import Base
from database.schema_upgrade import upgrade_existing_schema
from models.employee import Employee
from models.employee_history import EmployeeHistory
from models.user import User
from schemas.admin_management_schema import EmployeeAccountCreate, EmployeeMasterUpdate
from scripts.create_initial_data import seed_initial_data
from services.admin_management_service import AdminManagementService
from services.employee_bulk_import_service import (
    EMPLOYEE_IMPORT_COLUMNS,
    SAMPLE_EMPLOYEE_NUMBER,
    SAMPLE_EMPLOYEE_ROW,
    EmployeeBulkImportService,
    calculate_years_of_service,
    generated_temporary_password,
    generated_username,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _settings() -> Settings:
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        initial_company_code="EMPIMPORT",
        initial_company_name="Employee Import Company",
        initial_admin_username="admin",
        initial_admin_email="admin.employee.import@example.com",
        initial_admin_password=SecretStr("Temporary123!"),
        initial_admin_employee_number="ADMIN-001",
        initial_admin_first_name="System",
        initial_admin_last_name="Administrator",
    )


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine, sessionmaker(bind=engine, expire_on_commit=False)


def _filled_template() -> bytes:
    workbook = load_workbook(BytesIO(EmployeeBulkImportService.build_template()))
    sheet = workbook["Employee Import Template"]
    first = {
        "Employee Number": "200001",
        "Last Name": "Rodriguez",
        "First Name": "Robert Jay",
        "Middle Name": "Santos",
        "Suffix": "",
        "Gender": "Male",
        "Civil Status": "Single",
        "Date of Birth": "1990-01-01",
        "Email": "robert.rodriguez@example.com",
        "Telephone / Mobile No.": "09170000001",
        "Department": "Operations",
        "Manager Employee Number": "",
        "Leader Employee Number": "",
        "Job Title / Position": "Manager",
        "Employment Status": "Employed",
        "Hired Date": "2019-08-05",
        "Years of Service": "",
        "Training Checklist": "[x] Orientation; [ ] Safety",
    }
    second = {
        "Employee Number": "200002",
        "Last Name": "Garcia",
        "First Name": "Lander",
        "Middle Name": "",
        "Suffix": "",
        "Gender": "Male",
        "Civil Status": "Single",
        "Date of Birth": "1995-05-10",
        "Email": "lander.garcia@example.com",
        "Telephone / Mobile No.": "09170000002",
        "Department": "Operations",
        "Manager Employee Number": "Robert Jay Rodriguez",
        "Leader Employee Number": "",
        "Job Title / Position": "Developer",
        "Employment Status": "Employed",
        "Hired Date": "2020-01-15",
        "Years of Service": "999",
        "Training Checklist": "[x] Orientation",
    }
    sheet.append([first[column] for column in EMPLOYEE_IMPORT_COLUMNS])
    sheet.append([second[column] for column in EMPLOYEE_IMPORT_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_name_account_defaults_and_service_years() -> None:
    assert generated_username("Lander", "Garcia") == "lgarcia"
    assert generated_username("Robert Jay", "Rodriguez") == "rjrodriguez"
    assert generated_temporary_password("Lander", "Garcia", "191220") == "LG_191220"
    assert calculate_years_of_service(date(2019, 8, 5), as_of=date(2026, 8, 13)) == 7
    assert calculate_years_of_service(date(2019, 8, 14), as_of=date(2026, 8, 13)) == 6


def test_template_has_exact_employee_and_service_columns() -> None:
    workbook = load_workbook(BytesIO(EmployeeBulkImportService.build_template()), data_only=False)
    assert workbook.sheetnames == ["Employee Import Template", "Instructions"]
    sheet = workbook["Employee Import Template"]
    assert tuple(cell.value for cell in sheet[1]) == EMPLOYEE_IMPORT_COLUMNS
    assert "Hired Date" in EMPLOYEE_IMPORT_COLUMNS
    assert "Years of Service" in EMPLOYEE_IMPORT_COLUMNS
    assert sheet.max_row == 2
    assert tuple(cell.value for cell in sheet[2]) == SAMPLE_EMPLOYEE_ROW
    assert sheet["A2"].value == SAMPLE_EMPLOYEE_NUMBER
    assert all(cell.fill.fill_type == "solid" for cell in sheet[2])
    assert all(cell.fill.fgColor.rgb.endswith("D9D9D9") for cell in sheet[2])
    assert all(cell.font.italic for cell in sheet[2])


def test_unchanged_gray_sample_row_is_never_imported() -> None:
    _, factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service = EmployeeBulkImportService(session)
        try:
            service.prepare_preview(
                EmployeeBulkImportService.build_template(),
                filename="employee_template.xlsx",
                company_id=seed["company"].id,
            )
        except ValueError as exc:
            assert str(exc) == "The employee template does not contain any employee rows."
        else:
            raise AssertionError("The gray sample row must not be treated as an employee.")


def test_excel_preview_import_is_atomic_and_assigns_same_file_manager() -> None:
    _, factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service = EmployeeBulkImportService(session)
        preview = service.prepare_preview(
            _filled_template(),
            filename="employees.xlsx",
            company_id=seed["company"].id,
        )
        assert len(preview) == 2
        assert all(row["Validation"] == "Ready" for row in preview)
        assert preview[0]["User Name"] == "rjrodriguez"
        assert preview[1]["User Name"] == "lgarcia"
        assert preview[1]["Years of Service"] != 999

        results = service.import_preview_rows(
            preview,
            company_id=seed["company"].id,
            current_user_id=seed["admin_user"].id,
            filename="employees.xlsx",
        )
        assert len(results) == 2
        manager = session.scalar(
            select(Employee).where(Employee.employee_number == "200001")
        )
        member = session.scalar(
            select(Employee).where(Employee.employee_number == "200002")
        )
        assert member.manager_id == manager.id
        assert session.scalar(
            select(User).where(User.username == "lgarcia")
        ).must_change_password is True
        history = list(
            session.scalars(
                select(EmployeeHistory).where(
                    EmployeeHistory.upload_filename == "employees.xlsx"
                )
            ).all()
        )
        assert len(history) == 2
        assert all(item.old_values_json is None for item in history)
        assert all("password" not in (item.new_values_json or "").casefold() for item in history)


def test_duplicate_default_username_receives_two_random_digits() -> None:
    _, factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        AdminManagementService(session).create_employee_with_optional_account(
            EmployeeAccountCreate(
                company_id=seed["company"].id,
                employee_number="EXIST-1",
                first_name="Robert Jay",
                last_name="Rodriguez",
                work_email="existing.robert@example.com",
                create_login_account=True,
                username="rjrodriguez",
                login_email="existing.robert@example.com",
                temporary_password="RR_EXIST-1",
                clearance=2,
            ),
            current_user_id=seed["admin_user"].id,
        )
        preview = EmployeeBulkImportService(session).prepare_preview(
            _filled_template(),
            filename="employees.xlsx",
            company_id=seed["company"].id,
        )
        assert re.fullmatch(r"rjrodriguez\d{2}", str(preview[0]["User Name"]))


def test_unexpected_batch_failure_rolls_back_every_new_employee() -> None:
    _, factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service = EmployeeBulkImportService(session)
        preview = service.prepare_preview(
            _filled_template(),
            filename="employees.xlsx",
            company_id=seed["company"].id,
        )
        preview[1]["_employee"]["manager_employee_number"] = "MISSING"
        try:
            service.import_preview_rows(
                preview,
                company_id=seed["company"].id,
                current_user_id=seed["admin_user"].id,
                filename="employees.xlsx",
            )
        except ValueError as error:
            assert "MISSING" in str(error).upper()
        else:
            raise AssertionError("The invalid manager should block the whole batch.")
        imported = list(
            session.scalars(
                select(Employee).where(Employee.employee_number.in_(["200001", "200002"]))
            ).all()
        )
        assert imported == []


def test_resigned_employee_moves_to_archive_and_restores_same_records() -> None:
    _, factory = _factory()
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        service = AdminManagementService(session)
        employee = service.create_employee_with_optional_account(
            EmployeeAccountCreate(
                company_id=seed["company"].id,
                employee_number="300001",
                first_name="Archive",
                last_name="Employee",
                work_email="archive.employee@example.com",
                hire_date=date(2020, 1, 1),
                create_login_account=True,
                username="aemployee",
                login_email="archive.employee@example.com",
                temporary_password="AE_300001",
                clearance=2,
            ),
            current_user_id=seed["admin_user"].id,
        )
        employee_id = employee.id
        user_id = employee.user.id
        archived = service.update_employee_master_record(
            EmployeeMasterUpdate(
                company_id=seed["company"].id,
                employee_id=employee.id,
                employee_number=employee.employee_number,
                first_name=employee.first_name,
                last_name=employee.last_name,
                work_email=employee.work_email,
                hire_date=employee.hire_date,
                employment_status="resigned",
                username=employee.user.username,
                clearance=2,
            ),
            current_user_id=seed["admin_user"].id,
        )
        assert archived.archived_at is not None
        assert archived.user.is_active is False
        assert all(item.id != employee_id for item in service.list_employees(seed["company"].id))
        assert any(item.id == employee_id for item in service.list_archived_employees(seed["company"].id))

        restored = service.restore_archived_employee(
            company_id=seed["company"].id,
            employee_id=employee_id,
            current_user_id=seed["admin_user"].id,
        )
        assert restored.id == employee_id
        assert restored.user.id == user_id
        assert restored.employment_status == "employed"
        assert restored.archived_at is None
        assert restored.user.is_active is True
        actions = [row.action_type for row in service.list_employee_history(seed["company"].id)]
        assert "archived" in actions
        assert "restored" in actions


def test_runtime_upgrade_adds_archive_columns_without_deleting_records(tmp_path) -> None:
    database_path = tmp_path / "legacy.db"
    engine = create_engine(f"sqlite+pysqlite:///{database_path}")
    with engine.begin() as connection:
        connection.exec_driver_sql(
            "CREATE TABLE employees (id INTEGER PRIMARY KEY, company_id INTEGER NOT NULL, "
            "employee_number VARCHAR(80) NOT NULL, first_name VARCHAR(100) NOT NULL, "
            "last_name VARCHAR(100) NOT NULL, employment_status VARCHAR(50) NOT NULL, "
            "created_at DATETIME DEFAULT CURRENT_TIMESTAMP NOT NULL, "
            "updated_at DATETIME DEFAULT CURRENT_TIMESTAMP NOT NULL)"
        )
        connection.exec_driver_sql(
            "INSERT INTO employees (id, company_id, employee_number, first_name, last_name, employment_status) "
            "VALUES (1, 1, 'OLD-1', 'Old', 'Employee', 'resigned')"
        )
    upgrade_existing_schema(engine)
    columns = {column["name"] for column in inspect(engine).get_columns("employees")}
    assert {"archived_at", "archived_by_user_id"}.issubset(columns)
    with engine.connect() as connection:
        row = connection.exec_driver_sql(
            "SELECT employee_number, archived_at FROM employees WHERE id = 1"
        ).one()
    assert row.employee_number == "OLD-1"
    assert row.archived_at is not None


def test_employee_ui_has_new_tabs_and_no_deprecated_calls() -> None:
    source = (PROJECT_ROOT / "ui/pages/admin/employees_page.py").read_text(encoding="utf-8")
    assert '"Hired Date"' in source
    assert '"Years of Service"' in source
    assert '"History"' in source
    assert 'f"Archive ({len(archived_employees)})"' in source
    assert "Upload Employees via Excel" in source
    assert "st.components.v1.html" not in source
    assert "components.html" not in source
    assert "use_container_width" not in source
