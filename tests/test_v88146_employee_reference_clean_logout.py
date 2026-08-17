"""Regression coverage for v8.8.146 assignment references and logout."""

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from database.base import Base
from models.company import Company
from models.employee import Employee
from services.employee_bulk_import_service import (
    EMPLOYEE_IMPORT_COLUMNS,
    EmployeeBulkImportService,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _employee_row(**overrides) -> dict[str, object]:
    row = {
        "Employee Number": "E-002",
        "Last Name": "Garcia",
        "First Name": "Lander",
        "Middle Name": "",
        "Suffix": "",
        "Gender": "Male",
        "Civil Status": "Single",
        "Date of Birth": "1995-01-15",
        "Email": "lander@example.com",
        "Telephone / Mobile No.": "",
        "Department": "IT",
        "Manager Employee Number": "",
        "Leader Employee Number": "",
        "Job Title / Position": "Developer",
        "Employment Status": "Employed",
        "Hired Date": "2022-01-15",
        "Years of Service": "",
        "Training Checklist": "",
    }
    row.update(overrides)
    return row


def _workbook_bytes(*rows: dict[str, object]) -> bytes:
    workbook = load_workbook(BytesIO(EmployeeBulkImportService.build_template()))
    sheet = workbook["Employee Import Template"]
    for row in rows:
        sheet.append([row.get(column, "") for column in EMPLOYEE_IMPORT_COLUMNS])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def test_template_keeps_columns_and_documents_accepted_name_references() -> None:
    workbook = load_workbook(BytesIO(EmployeeBulkImportService.build_template()))
    sheet = workbook["Employee Import Template"]
    assert "Manager Employee Number" in [cell.value for cell in sheet[1]]
    assert "Leader Employee Number" in [cell.value for cell in sheet[1]]
    instructions = " ".join(
        str(cell.value or "")
        for row in workbook["Instructions"].iter_rows()
        for cell in row
    )
    assert "exact full name" in instructions
    assert "unique first/last name" in instructions
    assert sheet["L2"].value == "100001"
    assert sheet["M2"].value == "Maria Santos"


def test_existing_employee_can_be_resolved_by_unique_first_or_last_name() -> None:
    factory = _factory()
    with factory() as session:
        company = Company(code="REF", name="Reference Company")
        session.add(company)
        session.flush()
        session.add(
            Employee(
                company_id=company.id,
                employee_number="M-001",
                first_name="Robert Jay",
                last_name="Rodriguez",
                employment_status="employed",
            )
        )
        session.commit()
        preview = EmployeeBulkImportService(session).prepare_preview(
            _workbook_bytes(
                _employee_row(
                    **{
                        "Manager Employee Number": "  Robert Jay  ",
                        "Leader Employee Number": "rodriguez",
                    }
                )
            ),
            filename="employees.xlsx",
            company_id=company.id,
        )
        assert preview[0]["Validation"] == "Ready"
        assert preview[0]["_employee"]["manager_employee_number"] == "M-001"
        assert preview[0]["_employee"]["leader_employee_number"] == "M-001"


def test_same_upload_employee_can_be_resolved_by_first_and_last_name() -> None:
    factory = _factory()
    with factory() as session:
        company = Company(code="BATCH", name="Batch Company")
        session.add(company)
        session.commit()
        manager = _employee_row(
            **{
                "Employee Number": "M-100",
                "First Name": "Maria Elena",
                "Last Name": "Santos",
                "Email": "maria@example.com",
                "Job Title / Position": "Manager",
            }
        )
        member = _employee_row(
            **{
                "Manager Employee Number": "Maria Elena Santos",
            }
        )
        preview = EmployeeBulkImportService(session).prepare_preview(
            _workbook_bytes(manager, member),
            filename="batch.xlsx",
            company_id=company.id,
        )
        assert all(row["Validation"] == "Ready" for row in preview)
        assert preview[1]["_employee"]["manager_employee_number"] == "M-100"


def test_ambiguous_name_is_rejected_without_guessing() -> None:
    factory = _factory()
    with factory() as session:
        company = Company(code="AMB", name="Ambiguous Company")
        session.add(company)
        session.flush()
        session.add_all(
            [
                Employee(
                    company_id=company.id,
                    employee_number=f"M-00{index}",
                    first_name="John",
                    last_name=last_name,
                    employment_status="employed",
                )
                for index, last_name in enumerate(("Reyes", "Santos"), start=1)
            ]
        )
        session.commit()
        preview = EmployeeBulkImportService(session).prepare_preview(
            _workbook_bytes(
                _employee_row(**{"Manager Employee Number": "John"})
            ),
            filename="ambiguous.xlsx",
            company_id=company.id,
        )
        assert "matches multiple employees" in preview[0]["Validation"]
        assert "Employee Number" in preview[0]["Validation"]


def test_logout_uses_clean_supported_browser_reload() -> None:
    session_source = (PROJECT_ROOT / "authentication/session_manager.py").read_text(
        encoding="utf-8"
    )
    storage_source = (
        PROJECT_ROOT / "authentication/browser_auth_storage.py"
    ).read_text(encoding="utf-8")
    app_source = (PROJECT_ROOT / "app.py").read_text(encoding="utf-8")
    assert "window.parent.location.reload()" in session_source
    assert "def complete_logout_transition" in session_source
    assert "AuthSessionManager.complete_logout_transition()" in app_source
    logout_block = session_source.split("def logout(cls)", 1)[1].split(
        "def complete_logout_transition", 1
    )[0]
    assert "st.rerun()" in logout_block
    assert "Signing out…" in storage_source
    assert "st.rerun()" not in storage_source.split(
        "def remove_browser_auth_token", 1
    )[1]


def test_release_keeps_deprecated_streamlit_calls_out() -> None:
    files = [
        path
        for path in PROJECT_ROOT.rglob("*.py")
        if "tests" not in path.parts
        and ".venv" not in path.parts
        and "venv" not in path.parts
    ]
    source = "\n".join(path.read_text(encoding="utf-8") for path in files)
    assert "st.components.v1.html" not in source
    assert "components.html" not in source
    assert "use_container_width" not in source
