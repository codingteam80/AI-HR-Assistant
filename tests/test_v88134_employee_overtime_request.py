"""Employee OT form, DTR grounding, approval, and export regression tests."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pydantic import SecretStr
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import models  # noqa: F401
from config.settings import Settings
from database.base import Base
from models.attendance_record import AttendanceRecord
from modules.reports.combined_hr_report import build_combined_hr_excel
from schemas.overtime_schema import OvertimeRequestInput, OvertimeReviewInput
from scripts.create_initial_data import seed_initial_data
from services.overtime_service import OvertimeService


ROOT = Path(__file__).resolve().parents[1]


def _settings() -> Settings:
    return Settings(
        _env_file=None,
        database_url="sqlite+pysqlite:///:memory:",
        initial_company_code="V88134",
        initial_company_name="Overtime Company",
        initial_admin_username="admin",
        initial_admin_email="admin@example.com",
        initial_admin_password=SecretStr("Temporary123!"),
        initial_admin_employee_number="ADMIN-134",
    )


def _factory():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine, expire_on_commit=False)


def _attendance(seed, rendered: date) -> AttendanceRecord:
    return AttendanceRecord(
        company_id=seed["company"].id,
        employee_id=seed["admin_employee"].id,
        attendance_date=rendered,
        time_in=datetime(2026, 8, 14, 0, 0, tzinfo=timezone.utc),
        time_out=datetime(2026, 8, 14, 12, 0, tzinfo=timezone.utc),
        work_status="WFO",
        scheduled_workday=True,
        regular_hours_target=Decimal("8.00"),
        lunch_break_minutes=60,
        total_hours=Decimal("11.00"),
        ot_hours=Decimal("3.00"),
    )


def _input(seed, rendered: date, *, edited: bool) -> OvertimeRequestInput:
    return OvertimeRequestInput(
        company_id=seed["company"].id,
        employee_id=seed["admin_employee"].id,
        requested_by_user_id=seed["admin_user"].id,
        date_rendered=rendered,
        ot_time_start=datetime(2026, 8, 14, 17, 15 if edited else 0),
        ot_time_end=datetime(2026, 8, 14, 20, 0),
        estimated_hours=Decimal("2.75" if edited else "3.00"),
        ot_type="Regular Overtime",
        ot_purpose="System deployment",
        travel_fare=Decimal("100"),
        travel_route="Office to Bacoor, Cavite",
        dinner_break_flag=True,
    )


def test_employee_ot_request_keeps_dtr_snapshot_and_edit_mismatch() -> None:
    factory = _factory()
    rendered = date(2026, 8, 14)
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        record = _attendance(seed, rendered)
        session.add(record)
        session.commit()

        request = OvertimeService(session).submit(_input(seed, rendered, edited=True))
        assert request.status == "pending_approval"
        assert request.attendance_record_id == record.id
        assert Decimal(request.dtr_estimated_hours) == Decimal("3.00")
        assert request.has_dtr_mismatch is True
        assert record.time_out == datetime(2026, 8, 14, 12, 0, tzinfo=timezone.utc)


def test_approved_request_supplies_full_ot_export_fields() -> None:
    factory = _factory()
    rendered = date(2026, 8, 14)
    with factory() as session:
        seed = seed_initial_data(session, _settings())
        employee = seed["admin_employee"]
        record = _attendance(seed, rendered)
        session.add(record)
        session.commit()
        service = OvertimeService(session)
        request = service.submit(_input(seed, rendered, edited=False))
        service.review(
            OvertimeReviewInput(
                company_id=seed["company"].id,
                overtime_request_id=request.id,
                reviewed_by_user_id=seed["admin_user"].id,
                reviewer_employee_id=employee.id,
                clearance=1,
                decision="approved",
                comment="Validated against DTR",
            )
        )
        workbook_data = build_combined_hr_excel(
            employees=[employee],
            attendance_records=[record],
            leave_requests=[],
            overtime_requests=[request],
            start_date=rendered,
            end_date=rendered,
            scheduled_workdays={rendered: True},
            timezone_name="Asia/Manila",
        )
        sheet = load_workbook(BytesIO(workbook_data))["Overtime File"]
        assert sheet.max_row == 2
        assert sheet["D2"].value == "05:00"
        assert sheet["E2"].value == "PM"
        assert sheet["F2"].value == "08:00"
        assert sheet["G2"].value == "PM"
        assert sheet["H2"].value == 3
        assert sheet["I2"].value == "Regular Overtime"
        assert sheet["J2"].value == "System deployment"
        assert sheet["K2"].value == 100
        assert sheet["L2"].value == "Office to Bacoor, Cavite"
        assert sheet["M2"].value == 1
        assert sheet["N2"].value == "CC"


def test_employee_ui_combines_filing_and_history_in_one_expander() -> None:
    source = (ROOT / "ui/pages/user/attendance_workspace.py").read_text()
    assert 'with st.expander("Overtime Request")' in source
    assert "**File Overtime Request**" in source
    assert "**My Overtime Requests**" in source
