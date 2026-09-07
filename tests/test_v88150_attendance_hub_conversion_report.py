"""Regression checks for Attendance Hub and Leave Conversion reporting."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from modules.reports.leave_conversion_report import (
    HEADERS,
    build_leave_conversion_excel,
    build_leave_conversion_rows,
)


ROOT = Path(__file__).resolve().parents[1]


def _source(relative_path: str) -> str:
    return (ROOT / relative_path).read_text(encoding="utf-8")


def test_attendance_hub_is_available_in_both_portals() -> None:
    constants = _source("core/constants.py")
    admin_sidebar = _source("ui/components/admin_sidebar.py")
    employee_layout = _source("ui/layouts/user_layout.py")
    admin_layout = _source("ui/layouts/admin_layout.py")
    assert '"Attendance Hub"' in constants
    assert '"Attendance Hub"' in admin_sidebar
    assert 'current_page == "Attendance Hub"' in employee_layout
    assert 'page == "Attendance Hub"' in admin_layout
    assert "render_employee_attendance_hub_page" in employee_layout
    assert "render_admin_attendance_hub_page" in admin_layout


def test_dashboard_uses_one_persisted_dynamic_punch_button() -> None:
    attendance = _source("ui/pages/user/attendance_workspace.py")
    punch_area = attendance[
        attendance.index("daily_sessions ="):
        attendance.index("if punch_clicked:")
    ]
    assert 'punch_label = "Time Out" if has_any_punch else "Time In"' in punch_area
    assert "disabled=has_completed_session" in punch_area
    assert "button_columns = st.columns(3)" not in attendance
    assert '"Save Status"' not in attendance
    assert '"Login"' not in punch_area
    assert '"Logout"' not in punch_area
    dashboard = _source("ui/pages/user/dashboard_page.py")
    assert "show_punch_controls=True" in dashboard
    assert "show_record_management=False" in dashboard
    employee_hub = _source("ui/pages/user/attendance_hub_page.py")
    assert "show_punch_controls=False" in employee_hub
    assert "show_record_management=True" in employee_hub


def test_admin_management_tools_are_moved_to_attendance_hub() -> None:
    dashboard = _source("ui/pages/admin/admin_dashboard_page.py")
    hub = _source("ui/pages/admin/attendance_hub_page.py")
    attendance = _source("ui/pages/admin/attendance_dashboard.py")
    assert "show_management_tools=False" in dashboard
    assert "show_management_tools=True" in hub
    for label in (
        "Attendance Schedule & OT Rules",
        "Overtime Requests & Validation",
        "Admin Attendance Correction",
        "Attendance Edit History",
    ):
        assert label in attendance


def test_leave_conversion_rows_use_ledger_semantics_and_blank_amount() -> None:
    employee = SimpleNamespace(
        id=1,
        employee_number="EMP-001",
        first_name="Lander",
        middle_name="Sotto",
        last_name="Garcia",
        suffix=None,
    )
    vacation = SimpleNamespace(
        employee_id=1,
        leave_type=SimpleNamespace(code="VACATION"),
        beginning_credit_days=Decimal("30"),
        credit_days=Decimal("17"),
        converted_to_cash_days=Decimal("2"),
    )
    sick = SimpleNamespace(
        employee_id=1,
        leave_type=SimpleNamespace(code="SICK"),
        beginning_credit_days=Decimal("8"),
        credit_days=Decimal("17"),
        converted_to_cash_days=Decimal("10"),
    )
    rows = build_leave_conversion_rows(
        employees=[employee],
        balances=[vacation, sick],
        leave_year=2027,
    )
    row = rows[0]
    assert list(row) == list(HEADERS)
    assert row["Employee Name"] == "Garcia, Lander, S."
    assert row["Available Vacation Leave (Last Year)"] == Decimal("30.00")
    assert row["Credit Vacation Leave (Current Year)"] == Decimal("17.00")
    assert row["Remaining Sick Leave (Last Year)"] == Decimal("8.00")
    assert row["Total Conversion"] == Decimal("12.00")
    assert row["Total Conversion Amount"] == ""

    workbook = load_workbook(BytesIO(build_leave_conversion_excel(rows)))
    sheet = workbook["Leave Conversion to Cash"]
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.cell(2, 9).value == 12
    assert sheet.cell(2, 10).value is None


def test_report_page_has_separate_conversion_tab_and_all_column_search() -> None:
    report = _source("ui/pages/admin/reports_page.py")
    assert '"Combined HR Report", "Leave Conversion to Cash"' in report
    assert "build_leave_conversion_rows" in report
    assert "build_leave_conversion_excel" in report
    assert "multi_search_input(" in report
    assert "matches_search_terms(search_terms, row.values())" in report
    for excluded in ("Conversion Status", "Date Processed", "Remarks", "Daily Rate"):
        assert excluded not in _source(
            "modules/reports/leave_conversion_report.py"
        )
