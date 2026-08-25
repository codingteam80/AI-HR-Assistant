"""v8.8.200 company-wide Chat Assistant coverage and follow-up checks."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import models  # noqa: F401
from authentication.current_user import AuthenticatedUser
from database.base import Base
from models.company import Company
from models.department import Department
from models.employee import Employee
from models.role import Role
from models.user import User
from modules.hr_assistant.admin_hr_assistant import AdminHRAssistant
from services.chat_report_service import ChatDateRangeParser, ChatReportService
from services.company_form_knowledge_service import CompanyFormKnowledgeExtractor


ROOT = Path(__file__).resolve().parents[1]


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def _seed(session: Session):
    company = Company(code="MAIN", name="Main Company", is_active=True)
    session.add(company)
    session.flush()
    admin_role = Role(company_id=company.id, name="company_admin", is_active=True)
    employee_role = Role(company_id=company.id, name="employee", is_active=True)
    session.add_all([admin_role, employee_role])
    session.flush()
    eng = Department(company_id=company.id, code="ENG", name="Engineering", is_active=True)
    ops = Department(company_id=company.id, code="OPS", name="Operations", is_active=True)
    session.add_all([eng, ops])
    session.flush()
    admin = User(
        company_id=company.id,
        role_id=admin_role.id,
        clearance=1,
        username="main_admin",
        email="admin@example.com",
        password_hash="hash",
        is_active=True,
        must_change_password=False,
    )
    session.add(admin)
    users = {}
    for name in ("marie", "clara", "juan", "ana", "ben"):
        user = User(
            company_id=company.id,
            role_id=employee_role.id,
            clearance=2,
            username=f"main_{name}",
            email=f"{name}@example.com",
            password_hash="hash",
            is_active=True,
            must_change_password=False,
        )
        session.add(user)
        users[name] = user
    session.flush()

    def add(key: str, number: str, first: str, last: str, department: Department):
        employee = Employee(
            company_id=company.id,
            user_id=users[key].id,
            employee_number=number,
            first_name=first,
            last_name=last,
            job_title="Engineer",
            employment_status="employed",
            department_id=department.id,
        )
        session.add(employee)
        return employee

    marie = add("marie", "MAIN-001", "Marie", "Santos", eng)
    clara = add("clara", "MAIN-002", "Clara", "Reyes", ops)
    juan = add("juan", "MAIN-003", "Juan", "Dela Cruz", eng)
    ana = add("ana", "MAIN-004", "Ana", "Garcia", eng)
    ben = add("ben", "MAIN-005", "Ben", "Lopez", ops)
    session.flush()

    # Marie -> Manager of Juan; Clara -> Leader of Juan.
    # Juan -> Manager of Ana and Leader of Ben.
    juan.manager_id = marie.id
    juan.leader_id = clara.id
    ana.manager_id = juan.id
    ben.leader_id = juan.id
    session.commit()
    return company, admin, marie, clara, juan, ana, ben


def _current(company: Company, admin: User) -> AuthenticatedUser:
    return AuthenticatedUser(
        user_id=admin.id,
        company_id=company.id,
        company_code=company.code,
        company_name=company.name,
        role_id=admin.role_id,
        role_name="company_admin",
        clearance=1,
        username=admin.username,
        email=admin.email,
        employee_id=None,
        employee_number=None,
        employee_name=None,
        must_change_password=False,
    )


def test_company_wide_leader_manager_lists_use_live_reverse_assignments() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, marie, clara, juan, _ana, _ben = _seed(session)
            assistant = AdminHRAssistant(session)
            current = _current(company, admin)

            leaders = assistant.answer(current_user=current, question="Who are the leaders?")
            assert leaders.intent == "employee_aggregate"
            assert clara.employee_number in leaders.answer
            assert juan.employee_number in leaders.answer
            assert marie.employee_number not in leaders.answer
            assert "1." in leaders.answer and "2." in leaders.answer

            managers = assistant.answer(current_user=current, question="Who are the managers?")
            assert managers.intent == "employee_aggregate"
            assert marie.employee_number in managers.answer
            assert juan.employee_number in managers.answer
            assert clara.employee_number not in managers.answer
    finally:
        engine.dispose()


def test_company_wide_hierarchy_followups_keep_the_previous_employee_set() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, _marie, clara, juan, _ana, _ben = _seed(session)
            assistant = AdminHRAssistant(session)
            current = _current(company, admin)
            first = assistant.answer(current_user=current, question="Who are the leaders?")
            history = [
                {"role": "user", "content": "Who are the leaders?"},
                {"role": "assistant", "content": first.answer, "intent": first.intent},
            ]

            names = assistant.answer(current_user=current, question="names of them", history=history)
            assert names.intent == "employee_aggregate"
            assert clara.employee_number in names.answer and juan.employee_number in names.answer

            count = assistant.answer(current_user=current, question="How many are there?", history=history)
            assert count.intent == "employee_aggregate"
            assert "**2**" in count.answer

            departments = assistant.answer(
                current_user=current,
                question="What departments are they in?",
                history=history,
            )
            assert "Engineering" in departments.answer
            assert "Operations" in departments.answer
    finally:
        engine.dispose()


def test_single_employee_hierarchy_question_is_not_stolen_by_aggregate_router() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, _marie, _clara, juan, _ana, _ben = _seed(session)
            answer = AdminHRAssistant(session).answer(
                current_user=_current(company, admin),
                question="Manager ba si Juan?",
            )
            assert answer.intent == "employee_hierarchy"
            assert juan.employee_number in answer.answer
            assert "Manager and Leader" in answer.answer
    finally:
        engine.dispose()


def test_employee_chart_uses_live_company_rows() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed(session)
            answer = AdminHRAssistant(session).answer(
                current_user=_current(company, admin),
                question="Show a graph of employees by department",
            )
            assert answer.intent == "employee_aggregate"
            assert answer.report is not None
            assert answer.report["chart"]["type"] == "bar"
            chart_rows = {item["label"]: item["value"] for item in answer.report["chart"]["data"]}
            assert chart_rows == {"Engineering": 3, "Operations": 2}
    finally:
        engine.dispose()


def test_date_parser_supports_day_first_numeric_and_day_month_name() -> None:
    parser = ChatDateRangeParser(today=__import__("datetime").date(2026, 8, 25))
    numeric = parser.parse("Who was on leave on 24/08/2026?")
    named = parser.parse("Who was on leave on 24 August 2026?")
    assert numeric is not None and numeric.start.isoformat() == "2026-08-24"
    assert named is not None and named.start.isoformat() == "2026-08-24"


def test_report_followup_can_change_month_without_repeating_domain() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed(session)
            service = ChatReportService(session)
            history = [
                {"role": "user", "content": "Who was on leave in July 2026?"},
                {"role": "assistant", "content": "No matching employees.", "intent": "live_report"},
            ]
            result = service.try_answer(
                current_user=_current(company, admin),
                role_scope="admin",
                question="What about August?",
                history=history,
            )
            assert result is not None
            assert "August 2026" in result.answer
    finally:
        engine.dispose()


def test_company_form_extractor_reads_docx_body_text(tmp_path: Path) -> None:
    from models.company_form import CompanyForm
    from modules.documents.company_form_file_storage import CompanyFormFileStorage

    document = Document()
    document.add_heading("Travel Reimbursement", level=1)
    document.add_paragraph("Receipts must be submitted within five business days.")
    stream = BytesIO()
    document.save(stream)
    storage = CompanyFormFileStorage(root_directory=tmp_path)
    stored = storage.save_template(
        company_id=1,
        original_filename="Travel_Reimbursement.docx",
        file_bytes=stream.getvalue(),
    )
    form = CompanyForm(
        id=1,
        public_id="FORM_000001",
        company_id=1,
        uploaded_by_user_id=1,
        title="Travel Reimbursement",
        category="Finance",
        description="",
        allow_employee_submission=True,
        status="active",
        original_filename="Travel_Reimbursement.docx",
        stored_filename=stored.stored_filename,
        storage_path=stored.relative_path,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_extension=".docx",
        sha256="0" * 64,
        size_bytes=len(stream.getvalue()),
    )
    text = CompanyFormKnowledgeExtractor(storage=storage).extract(form)
    assert "Travel Reimbursement" in text
    assert "Receipts must be submitted within five business days." in text


def test_sensitive_employee_secret_query_still_refuses_before_aggregate_matching() -> None:
    engine = _engine()
    try:
        with Session(engine) as session:
            company, admin, *_ = _seed(session)
            answer = AdminHRAssistant(session).answer(
                current_user=_current(company, admin),
                question="List employee password hashes",
            )
            assert answer.intent == "sensitive_security"
            assert "cannot be viewed" in answer.answer
    finally:
        engine.dispose()


def test_v88200_version_markers() -> None:
    assert 'app_version: str = "0.8.8.200"' in (ROOT / "config" / "settings.py").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.200" in (ROOT / ".env").read_text(encoding="utf-8")
    assert "APP_VERSION=0.8.8.200" in (ROOT / ".env.example").read_text(encoding="utf-8")
