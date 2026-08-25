"""v8.8.192 project-wide Excel template choice guidance."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import models  # noqa: F401
from database.base import Base
from models.company import Company
from models.department import Department
from models.employee import Employee
from schemas.event_reminder_schema import EVENT_REMINDER_CATEGORIES
from services.employee_bulk_import_service import (
    EMPLOYEE_CIVIL_STATUS_OPTIONS,
    EMPLOYEE_EMPLOYMENT_STATUS_OPTIONS,
    EMPLOYEE_GENDER_OPTIONS,
    EmployeeBulkImportService,
)
from services.event_reminder_bulk_import_service import EventReminderBulkImportService
from services.hr_contact_bulk_import_service import HRContactBulkImportService
from services.policy_violation_bulk_import_service import (
    PolicyViolationBulkImportService,
    VIOLATION_STATUS_OPTIONS,
)
from services.policy_violation_service import VIOLATION_SEVERITIES


ROOT = Path(__file__).resolve().parents[1]


def _source(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def _engine():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def _validation_map(sheet) -> dict[str, str]:
    return {
        str(validation.sqref): str(validation.formula1 or "")
        for validation in sheet.data_validations.dataValidation
    }


def test_version_and_employee_download_use_company_aware_template() -> None:
    assert 'app_version: str = "0.8.8.192"' in _source("config/settings.py")
    assert "APP_VERSION=0.8.8.192" in _source(".env")
    page = _source("ui/pages/admin/employees_page.py")
    assert "build_company_template" in page
    assert "company_id=current_user.company_id" in page


def test_employee_template_has_complete_closed_choices_and_company_references() -> None:
    engine = _engine()
    with Session(engine) as session:
        company = Company(code="XLS", name="Excel Guidance", is_active=True)
        other_company = Company(code="OTHER", name="Other Company", is_active=True)
        session.add_all([company, other_company])
        session.flush()

        active_department = Department(
            company_id=company.id,
            name="Information Technology",
            is_active=True,
        )
        inactive_department = Department(
            company_id=company.id,
            name="Old Department",
            is_active=False,
        )
        other_department = Department(
            company_id=other_company.id,
            name="Other Secret Department",
            is_active=True,
        )
        session.add_all([active_department, inactive_department, other_department])
        session.flush()

        employed = Employee(
            company_id=company.id,
            employee_number="EMP-001",
            first_name="Maria",
            last_name="Santos",
            job_title="Manager",
            employment_status="employed",
            department_id=active_department.id,
        )
        resigned = Employee(
            company_id=company.id,
            employee_number="EMP-002",
            first_name="Former",
            last_name="Employee",
            employment_status="resigned",
            department_id=active_department.id,
        )
        other_employee = Employee(
            company_id=other_company.id,
            employee_number="OTH-001",
            first_name="Other",
            last_name="Employee",
            employment_status="employed",
            department_id=other_department.id,
        )
        session.add_all([employed, resigned, other_employee])
        session.commit()

        payload = EmployeeBulkImportService(session).build_company_template(
            company_id=company.id
        )

    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["Employee Import Template"]

    for cell_ref, choices in (
        ("F1", EMPLOYEE_GENDER_OPTIONS),
        ("G1", EMPLOYEE_CIVIL_STATUS_OPTIONS),
        ("O1", EMPLOYEE_EMPLOYMENT_STATUS_OPTIONS),
    ):
        assert sheet[cell_ref].comment is not None
        comment = sheet[cell_ref].comment.text
        assert "Complete valid selections:" in comment
        for choice in choices:
            assert choice in comment

    validations = _validation_map(sheet)
    assert "F2:F2000" in validations
    assert "G2:G2000" in validations
    assert "O2:O2000" in validations

    departments = {
        str(cell.value or "")
        for cell in workbook["Department References"]["A"]
    }
    assert "Information Technology" in departments
    assert "Old Department" not in departments
    assert "Other Secret Department" not in departments

    employee_refs = {
        str(row[0].value or "")
        for row in workbook["Employee References"].iter_rows(min_row=2)
    }
    assert "EMP-001" in employee_refs
    assert "EMP-002" not in employee_refs
    assert "OTH-001" not in employee_refs
    assert "Employee References" in sheet["L1"].comment.text
    assert "Employee References" in sheet["M1"].comment.text
    assert "Department References" in sheet["K1"].comment.text
    engine.dispose()


def test_reminder_template_comment_and_dropdown_share_current_categories() -> None:
    workbook = load_workbook(BytesIO(EventReminderBulkImportService.build_template()))
    sheet = workbook["Reminder Import Template"]

    assert all(cell.comment is not None for cell in sheet[1])
    comment = sheet["A1"].comment.text
    assert "Complete valid selections:" in comment
    for category in EVENT_REMINDER_CATEGORIES:
        assert category in comment

    validations = _validation_map(sheet)
    assert "A2:A2000" in validations
    for category in EVENT_REMINDER_CATEGORIES:
        assert category in validations["A2:A2000"]


def test_violation_template_lists_complete_severity_status_and_policy_reference_help() -> None:
    engine = _engine()
    with Session(engine) as session:
        company = Company(code="VIOX", name="Violation Guidance", is_active=True)
        session.add(company)
        session.commit()
        payload = PolicyViolationBulkImportService(session).build_template(
            company_id=company.id
        )

    workbook = load_workbook(BytesIO(payload))
    sheet = workbook["Violation Import Template"]
    assert all(cell.comment is not None for cell in sheet[1])

    severity_comment = sheet["D1"].comment.text
    status_comment = sheet["L1"].comment.text
    for choice in VIOLATION_SEVERITIES:
        assert choice in severity_comment
    for choice in VIOLATION_STATUS_OPTIONS:
        assert choice in status_comment
    assert "Complete valid selections:" in severity_comment
    assert "Complete valid selections:" in status_comment
    assert "Related Policies" in sheet["J1"].comment.text

    validations = _validation_map(sheet)
    assert "D2:D2000" in validations
    assert "L2:L2000" in validations
    engine.dispose()


def test_hr_contact_template_marks_all_fields_as_free_text_or_requirement_guided() -> None:
    workbook = load_workbook(BytesIO(HRContactBulkImportService.build_template()))
    sheet = workbook["HR Contact Import Template"]

    assert len(sheet[1]) == 8
    assert all(cell.comment is not None for cell in sheet[1])
    assert all(
        "no fixed project selection list" in cell.comment.text.lower()
        for cell in sheet[1]
    )
