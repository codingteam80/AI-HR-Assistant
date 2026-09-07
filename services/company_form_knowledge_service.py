"""Safe searchable-text extraction for authorized Company Form/Documents files."""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
import re

from openpyxl import load_workbook
from pypdf import PdfReader

from models.company_form import CompanyForm
from modules.documents.company_form_file_storage import CompanyFormFileStorage
from modules.documents.docx_structure_extractor import DOCXStructureExtractor
from modules.documents.pdf_layout_extractor import (
    PDFLayoutExtractionError,
    PDFLayoutExtractor,
)


class CompanyFormKnowledgeExtractor:
    """Extract bounded text from modern Company Form/Documents formats.

    The extractor is retrieval-only: it never modifies the source file and it
    fails closed (empty text) for unreadable/unsupported files. Old binary
    ``.doc``/``.xls`` files remain downloadable but are not falsely treated as
    text-searchable when no safe parser is available.
    """

    MAX_CHARACTERS = 150_000
    MAX_PDF_PAGES = 120
    MAX_SHEETS = 20
    MAX_ROWS_PER_SHEET = 1_000
    MAX_COLUMNS_PER_SHEET = 60

    def __init__(self, *, storage: CompanyFormFileStorage | None = None) -> None:
        self.storage = storage or CompanyFormFileStorage()

    @staticmethod
    def _clean(value: str) -> str:
        value = (value or "").replace("\r\n", "\n").replace("\r", "\n")
        value = re.sub(r"[ \t]+", " ", value)
        value = re.sub(r"\n{3,}", "\n\n", value)
        return value.strip()

    @classmethod
    def _bounded(cls, value: str) -> str:
        clean = cls._clean(value)
        if len(clean) <= cls.MAX_CHARACTERS:
            return clean
        return clean[: cls.MAX_CHARACTERS].rstrip() + "\n[Content truncated for search safety]"

    @staticmethod
    def _decode_text(data: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="replace")

    @classmethod
    def _pdf(cls, data: bytes) -> str:
        try:
            extracted = PDFLayoutExtractor.extract(
                data,
                max_pages=cls.MAX_PDF_PAGES,
            )
        except PDFLayoutExtractionError:
            extracted = None

        if extracted is not None and extracted.full_text:
            sections: list[str] = []
            for section in extracted.sections:
                page_heading = f"Page {section.page_number}"
                section_heading = section.heading.strip()
                if section_heading.casefold() == page_heading.casefold():
                    display_heading = page_heading
                else:
                    display_heading = f"{page_heading} — {section_heading}"
                sections.append(
                    f"## {display_heading}\n{section.text}"
                )
            if sections:
                return "\n\n".join(sections)

        # Compatibility fallback for unusual PDFs that PyPDF can extract even
        # when layout parsing cannot. This keeps one damaged/unusual document
        # from breaking or reducing the existing company-wide search surface.
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                if reader.decrypt("") == 0:
                    return ""
            except Exception:
                return ""
        pages = []
        for page_number, page in enumerate(reader.pages[: cls.MAX_PDF_PAGES], start=1):
            text = cls._clean(page.extract_text() or "")
            if text:
                pages.append(f"Page {page_number}: {text}")
        return "\n\n".join(pages)

    @classmethod
    def _docx(cls, data: bytes) -> str:
        """Emit ordered semantic markers while retaining a file-level fallback."""

        extracted = DOCXStructureExtractor.extract(
            data,
            default_heading="Document",
        )
        if not extracted.full_text:
            return ""
        if extracted.reliable_heading_count == 0:
            # No dependable section structure: keep the file intact here and
            # let the existing bounded chunker split it per file with overlap.
            return extracted.full_text
        return "\n\n".join(
            f"## {section.heading}\n{section.text}"
            for section in extracted.sections
            if section.text.strip()
        )

    @classmethod
    def _csv(cls, data: bytes) -> str:
        reader = csv.reader(StringIO(cls._decode_text(data)))
        lines: list[str] = []
        for index, row in enumerate(reader):
            if index >= cls.MAX_ROWS_PER_SHEET:
                break
            values = [str(value).strip() for value in row[: cls.MAX_COLUMNS_PER_SHEET]]
            if any(values):
                lines.append(" | ".join(values))
        return "\n".join(lines)

    @classmethod
    def _xlsx(cls, data: bytes) -> str:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        try:
            for sheet in workbook.worksheets[: cls.MAX_SHEETS]:
                lines.append(f"Sheet: {sheet.title}")
                for row_index, row in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    if row_index > cls.MAX_ROWS_PER_SHEET:
                        break
                    values = [
                        "" if value is None else str(value).strip()
                        for value in row[: cls.MAX_COLUMNS_PER_SHEET]
                    ]
                    if any(values):
                        lines.append(" | ".join(values))
        finally:
            workbook.close()
        return "\n".join(lines)

    def extract(self, form: CompanyForm) -> str:
        """Return searchable source-file text or an empty string safely."""

        extension = (form.file_extension or Path(form.original_filename).suffix).casefold()
        if extension not in {".pdf", ".docx", ".txt", ".csv", ".xlsx"}:
            return ""
        try:
            data = self.storage.read(form.storage_path)
            if extension == ".pdf":
                text = self._pdf(data)
            elif extension == ".docx":
                text = self._docx(data)
            elif extension == ".csv":
                text = self._csv(data)
            elif extension == ".xlsx":
                text = self._xlsx(data)
            else:
                text = self._decode_text(data)
            return self._bounded(text)
        except Exception:
            # Search indexing must never break the Chat page because one local
            # document is damaged, missing, encrypted, or temporarily unreadable.
            return ""
