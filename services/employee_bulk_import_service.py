"""Validated, atomic Excel onboarding for the Employee workspace."""

from __future__ import annotations

from datetime import date, datetime, timezone
from io import BytesIO
import json
import re
import secrets
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from authentication.password_manager import PasswordManager
from models.department import Department
from models.employee import Employee
from models.employee_history import EmployeeHistory
from models.employee_training import EmployeeTraining
from models.role import Role
from models.user import User
from schemas.admin_management_schema import EmployeeAccountCreate
from services.excel_template_guidance import (
    add_header_guidance,
    add_inline_list_validation,
    employee_reference_label,
)


EMPLOYEE_IMPORT_COLUMNS = (
    "Employee Number",
    "Last Name",
    "First Name",
    "Middle Name",
    "Suffix",
    "Gender",
    "Civil Status",
    "Date of Birth",
    "Email",
    "Telephone / Mobile No.",
    "Department",
    "Manager Employee Number",
    "Leader Employee Number",
    "Job Title / Position",
    "Employment Status",
    "Hired Date",
    "Years of Service",
    "Training Checklist",
)

ACCOUNT_PREVIEW_COLUMNS = (
    "Row",
    "Employee Number",
    "Employee Name",
    "Hired Date",
    "Years of Service",
    "User ID",
    "Clearance",
    "User Name",
    "Temporary Password",
    "Validation",
)

SAMPLE_EMPLOYEE_NUMBER = "SAMPLE ONLY - 191220"
SAMPLE_EMPLOYEE_ROW = (
    SAMPLE_EMPLOYEE_NUMBER,
    "Garcia",
    "Lander",
    "Santos",
    "Jr.",
    "Male",
    "Single",
    "1995-01-15",
    "lander.garcia@example.com",
    "09171234567",
    "Information Technology",
    "100001 - Santos, Maria",
    "100002 - Reyes, Pedro Jr.",
    "Software Developer",
    "Employed",
    "2022-01-15",
    "4 (automatic)",
    "[x] Orientation; [ ] Data Privacy",
)

EMPLOYEE_GENDER_OPTIONS = ("Male", "Female")
EMPLOYEE_CIVIL_STATUS_OPTIONS = (
    "N/A",
    "Single",
    "Married",
    "Widowed",
    "Separated",
    "Divorced",
)
EMPLOYEE_EMPLOYMENT_STATUS_OPTIONS = ("Employed", "Resigned")


def calculate_years_of_service(hire_date: date | None, *, as_of: date | None = None) -> int | None:
    """Return completed service years and never return a negative value."""

    if hire_date is None:
        return None
    reference = as_of or date.today()
    if hire_date > reference:
        return 0
    return (
        reference.year
        - hire_date.year
        - ((reference.month, reference.day) < (hire_date.month, hire_date.day))
    )


def generated_username(first_name: str, last_name: str) -> str:
    """Create all given-name initials followed by the normalized surname."""

    initials = "".join(
        token[0]
        for token in re.findall(r"[A-Za-z0-9]+", first_name)
        if token
    )
    surname = "".join(re.findall(r"[A-Za-z0-9]+", last_name))
    return f"{initials}{surname}".lower()


def generated_temporary_password(
    first_name: str,
    last_name: str,
    employee_number: str,
) -> str:
    """Create XX_EMPLOYEE-NUMBER using first-name and surname initials."""

    first_initial = next((c for c in first_name if c.isalnum()), "X")
    last_initial = next((c for c in last_name if c.isalnum()), "X")
    return f"{first_initial}{last_initial}_{employee_number}".upper()


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalized_employee_name(
    first_name: object,
    middle_name: object,
    last_name: object,
    suffix: object,
) -> str:
    """Normalize an employee name for duplicate prevention."""

    return " ".join(
        " ".join(_text(value).split()).casefold()
        for value in (first_name, middle_name, last_name, suffix)
        if _text(value)
    )


def _optional_text(value: object) -> str | None:
    normalized = _text(value)
    return normalized or None


def _normalized_name_reference(value: object) -> str:
    """Normalize an employee-name reference without changing employee IDs."""

    return " ".join(
        re.findall(r"[\w]+", _text(value).casefold(), flags=re.UNICODE)
    )


def _employee_reference_aliases(
    *,
    first_name: object,
    middle_name: object,
    last_name: object,
    suffix: object,
) -> set[str]:
    """Return supported full, first/last, first-only, and last-only aliases."""

    first = _text(first_name)
    middle = _text(middle_name)
    last = _text(last_name)
    suffix_value = _text(suffix)
    variants = (
        " ".join(part for part in (first, middle, last, suffix_value) if part),
        " ".join(part for part in (first, last) if part),
        first,
        last,
    )
    return {
        normalized
        for value in variants
        if (normalized := _normalized_name_reference(value))
    }


def _resolve_assignment_reference(
    value: object,
    candidates: list[dict[str, object]],
    *,
    field_name: str,
) -> tuple[str | None, str | None]:
    """Resolve one Manager/Leader cell to a unique canonical employee number."""

    raw = _text(value)
    if not raw:
        return None, None

    # Employee Number is always authoritative and is matched before names.
    # Company-aware templates display references as
    # "Employee Number - Last Name, First Name Suffix"; keep raw employee
    # numbers and older name-based entries backward compatible.
    number_candidate = raw.split(" - ", 1)[0].strip() if " - " in raw else raw
    number_key = number_candidate.casefold()
    number_matches = {
        str(candidate["employee_number"]): candidate
        for candidate in candidates
        if str(candidate["employee_number"]).casefold() == number_key
    }
    matches = list(number_matches.values())
    if not matches:
        name_key = _normalized_name_reference(raw)
        matches_by_number = {
            str(candidate["employee_number"]): candidate
            for candidate in candidates
            if name_key and name_key in set(candidate["name_aliases"])
        }
        matches = list(matches_by_number.values())

    if not matches:
        return None, (
            f"{field_name} '{raw}' does not match an employee number or name."
        )
    if len(matches) > 1:
        return None, (
            f"{field_name} '{raw}' matches multiple employees. Use the "
            "Employee Number or a more complete unique name."
        )

    selected = matches[0]
    if str(selected.get("employment_status") or "").casefold() != "employed":
        return None, (
            f"{field_name} '{raw}' matches an employee who is not Employed."
        )
    return str(selected["employee_number"]), None


def _parse_date(value: object, *, required: bool, field_name: str) -> date | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"{field_name} is required.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text_value = _text(value)
    try:
        return date.fromisoformat(text_value)
    except ValueError as exc:
        raise ValueError(f"{field_name} must use YYYY-MM-DD.") from exc


def _status(value: object) -> str:
    normalized = _text(value).casefold() or "employed"
    aliases = {
        "employed": "employed",
        "active": "employed",
        "resigned": "resigned",
        "inactive": "resigned",
    }
    if normalized not in aliases:
        raise ValueError("Employment Status must be Employed or Resigned.")
    return aliases[normalized]


def _candidate_employment_status(value: object) -> str:
    """Return a non-raising status used only while resolving references."""

    try:
        return _status(value)
    except ValueError:
        return "invalid"


def _training_items(value: object) -> list[dict[str, object]]:
    raw = _text(value)
    if not raw:
        return []
    items: list[dict[str, object]] = []
    for index, raw_item in enumerate(re.split(r"[;\n]+", raw)):
        item = raw_item.strip()
        if not item:
            continue
        completed = bool(re.match(r"^\[(x|X|✓)\]", item))
        title = re.sub(r"^\[(x|X|✓|\s)\]\s*", "", item).strip()
        if title:
            items.append(
                {
                    "title": title[:255],
                    "is_completed": completed,
                    "display_order": index,
                }
            )
    return items


class EmployeeBulkImportService:
    """Prepare Excel previews and commit a whole valid batch atomically."""

    def __init__(self, session: Session) -> None:
        self.session = session
        self.password_manager = PasswordManager()

    @staticmethod
    def build_template(
        *,
        department_names: tuple[str, ...] = (),
        employee_references: tuple[tuple[str, str, str, str], ...] = (),
        include_reference_sheets: bool = False,
    ) -> bytes:
        """Return a styled template with complete current selection guidance."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Employee Import Template"
        sheet.append(list(EMPLOYEE_IMPORT_COLUMNS))
        sheet.append(list(SAMPLE_EMPLOYEE_ROW))

        header_fill = PatternFill("solid", fgColor="2F740B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_fill = PatternFill("solid", fgColor="D9D9D9")
        sample_font = Font(color="666666", italic=True)
        for cell in sheet[2]:
            cell.fill = sample_fill
            cell.font = sample_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 34

        widths = {
            "A": 18, "B": 20, "C": 22, "D": 20, "E": 12,
            "F": 14, "G": 16, "H": 16, "I": 30, "J": 24,
            "K": 24, "L": 46, "M": 46, "N": 24, "O": 20,
            "P": 16, "Q": 18, "R": 36,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:R2000"

        add_header_guidance(
            sheet, "A1",
            requirement="Required and unique inside the company. Free-text employee number; no fixed selection.",
        )
        add_header_guidance(sheet, "B1", requirement="Required. Employee last name; free text.")
        add_header_guidance(sheet, "C1", requirement="Required. Employee first name; free text.")
        add_header_guidance(sheet, "D1", requirement="Optional middle name; free text.")
        add_header_guidance(sheet, "E1", requirement="Optional suffix such as Jr., Sr., II, or III; free text.")
        add_header_guidance(
            sheet, "F1",
            requirement="Optional. Use one of the same Gender choices available in Employee Master.",
            choices=EMPLOYEE_GENDER_OPTIONS,
        )
        add_header_guidance(
            sheet, "G1",
            requirement="Optional. Use one of the same Civil Status choices available in Employee Master.",
            choices=EMPLOYEE_CIVIL_STATUS_OPTIONS,
        )
        add_header_guidance(sheet, "H1", requirement="Optional date of birth. Use YYYY-MM-DD or an Excel date cell.")
        add_header_guidance(sheet, "I1", requirement="Required unique work/login email address.")
        add_header_guidance(sheet, "J1", requirement="Optional telephone or mobile number; free text.")
        add_header_guidance(
            sheet, "K1",
            requirement="Enter an existing department name or a new department name.",
            choices=department_names,
            choices_heading="Current company department names available at download time:",
            extra=(
                (
                    "The same current list is in the 'Department References' worksheet. "
                    "This field intentionally remains open because Employee Master permits creating a new department."
                    if department_names
                    else "No current company departments are available yet. You may enter a new department name."
                )
                if include_reference_sheets
                else (
                    "The company-aware template downloaded from the app includes the current Department References worksheet. "
                    "This field remains open because Employee Master permits creating a new department."
                )
            ),
        )
        assignment_choices = tuple(
            f"{number} - {name}"
            for number, name, _job_title, _department in employee_references
        )
        assignment_extra = (
            (
                "The complete current employed-employee reference list is in the 'Employee References' worksheet. "
                "You may also reference another employed row in the same upload by Employee Number."
            )
            if include_reference_sheets
            else (
                "The company-aware template downloaded from the app includes an Employee References worksheet. "
                "You may also reference another employed row in the same upload by Employee Number."
            )
        )
        add_header_guidance(
            sheet, "L1",
            requirement="Optional Manager reference. For current employees, choose or enter: Employee Number - Last Name, First Name Suffix. Another employed row in the same upload may still be referenced by Employee Number.",
            choices=assignment_choices,
            choices_heading="Current employed employee selections available at download time:",
            extra=(assignment_extra if assignment_choices else assignment_extra + " No current employed employee references are available yet."),
        )
        add_header_guidance(
            sheet, "M1",
            requirement="Optional Leader reference. For current employees, choose or enter: Employee Number - Last Name, First Name Suffix. Another employed row in the same upload may still be referenced by Employee Number.",
            choices=assignment_choices,
            choices_heading="Current employed employee selections available at download time:",
            extra=(assignment_extra if assignment_choices else assignment_extra + " No current employed employee references are available yet."),
        )
        add_header_guidance(sheet, "N1", requirement="Optional job title or position; free text.")
        add_header_guidance(
            sheet, "O1",
            requirement="Optional. Blank defaults to Employed.",
            choices=EMPLOYEE_EMPLOYMENT_STATUS_OPTIONS,
        )
        add_header_guidance(sheet, "P1", requirement="Required hired date. Use YYYY-MM-DD or an Excel date cell.")
        add_header_guidance(sheet, "Q1", requirement="Automatic from Hired Date during preview/import. Leave this column blank.")
        add_header_guidance(
            sheet, "R1",
            requirement="Optional checklist. Separate items with semicolons; prefix completed items with [x].",
            extra="Example: [x] Orientation; [ ] Data Privacy",
        )

        add_inline_list_validation(sheet, "F2:F2000", EMPLOYEE_GENDER_OPTIONS, allow_blank=True)
        add_inline_list_validation(sheet, "G2:G2000", EMPLOYEE_CIVIL_STATUS_OPTIONS, allow_blank=True)
        add_inline_list_validation(
            sheet, "O2:O2000", EMPLOYEE_EMPLOYMENT_STATUS_OPTIONS, allow_blank=True
        )

        if include_reference_sheets:
            department_sheet = workbook.create_sheet("Department References")
            department_sheet.append(["Current Department Name"])
            for cell in department_sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            if department_names:
                for name in department_names:
                    department_sheet.append([name])
            else:
                department_sheet.append(["No current departments. You may enter a new department name."])
            department_sheet.column_dimensions["A"].width = 52
            department_sheet.freeze_panes = "A2"

            employee_sheet = workbook.create_sheet("Employee References")
            employee_sheet.append(["Employee Number", "Employee Reference", "Employee Name", "Job Title / Position", "Department"])
            for cell in employee_sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            if employee_references:
                for number, name, job_title, department in employee_references:
                    employee_sheet.append([number, f"{number} - {name}", name, job_title, department])
                workbook.defined_names.add(
                    DefinedName(
                        "EmployeeAssignmentReferences",
                        attr_text=f"'Employee References'!$B$2:$B${len(employee_references) + 1}",
                    )
                )
                # Suggest the live reference list but allow a raw Employee Number
                # for another employed employee created in the same workbook.
                assignment_validation = DataValidation(
                    type="list", formula1="=EmployeeAssignmentReferences", allow_blank=True
                )
                assignment_validation.showErrorMessage = False
                assignment_validation.promptTitle = "Current employee references"
                assignment_validation.prompt = (
                    "Choose Employee Number - Last Name, First Name Suffix, or enter a same-upload Employee Number."
                )
                assignment_validation.showInputMessage = True
                sheet.add_data_validation(assignment_validation)
                assignment_validation.add("L2:L2000")
                assignment_validation.add("M2:M2000")
            else:
                employee_sheet.append(["", "", "No current employed employees available.", "", ""])
            employee_sheet.column_dimensions["A"].width = 24
            employee_sheet.column_dimensions["B"].width = 52
            employee_sheet.column_dimensions["C"].width = 38
            employee_sheet.column_dimensions["D"].width = 34
            employee_sheet.column_dimensions["E"].width = 30
            employee_sheet.freeze_panes = "A2"

        instructions = workbook.create_sheet("Instructions")
        instructions_rows = [
            ("Column", "Requirement"),
            ("Employee Number", "Required; unique inside the company."),
            ("First Name / Last Name", "Required."),
            ("Email", "Required; unique and used as the login email."),
            ("Hired Date", "Required; YYYY-MM-DD."),
            ("Years of Service", "Automatic from Hired Date during preview/import; leave this column blank."),
            ("Gender / Civil Status / Employment Status", "Use the dropdowns or hover the related column headers to see the complete current selections."),
            ("Department", "Current department names are listed in Department References. New department names are also allowed, matching Employee Master behavior."),
            ("Gray Sample Row", "Row 2 is a reference example only. Its unchanged SAMPLE ONLY employee number is automatically ignored during upload. Delete it or replace it with actual employee information."),
            (
                "Manager/Leader Employee Number",
                "Optional. Current employees are shown as Employee Number - Last Name, First Name Suffix in Employee References and the dropdown. Another employed row in the same file may be referenced by Employee Number. Legacy exact/unique employee-name references remain accepted for compatibility.",
            ),
            ("Employment Status", "Employed or Resigned; blank defaults to Employed."),
            ("Training Checklist", "Optional; separate items with semicolons. Use [x] for completed."),
            ("Account Information", "User ID, default clearance, username, and temporary password are generated after upload preview."),
        ]
        for row in instructions_rows:
            instructions.append(row)
        for cell in instructions[1]:
            cell.fill = header_fill
            cell.font = header_font
        instructions.column_dimensions["A"].width = 34
        instructions.column_dimensions["B"].width = 110
        instructions.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def build_company_template(self, *, company_id: int) -> bytes:
        """Return the employee template with current company reference sheets."""

        departments = list(
            self.session.scalars(
                select(Department)
                .where(
                    Department.company_id == company_id,
                    Department.is_active.is_(True),
                )
                .order_by(Department.name.asc())
            )
        )
        employees = list(
            self.session.scalars(
                select(Employee)
                .where(
                    Employee.company_id == company_id,
                    Employee.employment_status == "employed",
                )
                .order_by(Employee.last_name.asc(), Employee.first_name.asc())
            )
        )
        department_lookup = {department.id: department.name for department in departments}
        references = tuple(
            (
                employee.employee_number,
                employee_reference_label(
                    employee.employee_number, employee.last_name, employee.first_name, employee.suffix
                ).split(" - ", 1)[1],
                employee.job_title or "",
                department_lookup.get(employee.department_id, ""),
            )
            for employee in employees
        )
        return self.build_template(
            department_names=tuple(department.name for department in departments),
            employee_references=references,
            include_reference_sheets=True,
        )

    def _unique_username(self, base: str, reserved: set[str]) -> str:
        normalized = base.casefold()
        if normalized not in reserved:
            reserved.add(normalized)
            return base
        for _ in range(200):
            candidate = f"{base}{secrets.randbelow(100):02d}"
            if candidate.casefold() not in reserved:
                reserved.add(candidate.casefold())
                return candidate
        raise ValueError(f"Could not generate a unique username for '{base}'.")

    def prepare_preview(self, file_bytes: bytes, *, filename: str, company_id: int) -> list[dict[str, object]]:
        """Parse, validate, and return editable account-preview rows."""

        try:
            workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        except Exception as exc:
            raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc

        if "Employee Import Template" not in workbook.sheetnames:
            raise ValueError("Use the downloadable template and keep the 'Employee Import Template' worksheet name.")
        sheet = workbook["Employee Import Template"]
        headers = tuple(_text(cell.value) for cell in sheet[1])
        if headers != EMPLOYEE_IMPORT_COLUMNS:
            raise ValueError("The Excel columns or their order do not match the downloadable employee template.")

        existing_employees = list(
            self.session.scalars(
                select(Employee).where(Employee.company_id == company_id)
            ).all()
        )
        existing_users = list(
            self.session.scalars(select(User).where(User.company_id == company_id)).all()
        )
        existing_numbers = {employee.employee_number.strip().casefold() for employee in existing_employees}
        existing_emails = {user.email.strip().casefold() for user in existing_users}
        reserved_usernames = {user.username.strip().casefold() for user in existing_users}
        existing_names = {
            _normalized_employee_name(
                employee.first_name,
                employee.middle_name,
                employee.last_name,
                employee.suffix,
            )
            for employee in existing_employees
        }

        raw_rows: list[dict[str, object]] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if _text(values[0]).casefold() == SAMPLE_EMPLOYEE_NUMBER.casefold():
                continue
            raw_rows.append({"Row": row_number, **dict(zip(EMPLOYEE_IMPORT_COLUMNS, values))})
        if not raw_rows:
            raise ValueError("The employee template does not contain any employee rows.")

        batch_numbers = [_text(row["Employee Number"]).casefold() for row in raw_rows]
        batch_emails = [_text(row["Email"]).casefold() for row in raw_rows]
        batch_names = [
            _normalized_employee_name(
                row["First Name"],
                row["Middle Name"],
                row["Last Name"],
                row["Suffix"],
            )
            for row in raw_rows
        ]
        assignment_candidates = [
            {
                "employee_number": employee.employee_number,
                "employment_status": employee.employment_status,
                "name_aliases": _employee_reference_aliases(
                    first_name=employee.first_name,
                    middle_name=employee.middle_name,
                    last_name=employee.last_name,
                    suffix=employee.suffix,
                ),
            }
            for employee in existing_employees
        ]
        assignment_candidates.extend(
            {
                "employee_number": _text(row["Employee Number"]),
                "employment_status": _candidate_employment_status(
                    row["Employment Status"]
                ),
                "name_aliases": _employee_reference_aliases(
                    first_name=row["First Name"],
                    middle_name=row["Middle Name"],
                    last_name=row["Last Name"],
                    suffix=row["Suffix"],
                ),
            }
            for row in raw_rows
            if _text(row["Employee Number"])
        )
        preview: list[dict[str, object]] = []

        for row in raw_rows:
            errors: list[str] = []
            employee_number = _text(row["Employee Number"])
            first_name = _text(row["First Name"])
            last_name = _text(row["Last Name"])
            email = _text(row["Email"])
            normalized_name = _normalized_employee_name(
                row["First Name"],
                row["Middle Name"],
                row["Last Name"],
                row["Suffix"],
            )
            try:
                hired_date = _parse_date(row["Hired Date"], required=True, field_name="Hired Date")
                date_of_birth = _parse_date(row["Date of Birth"], required=False, field_name="Date of Birth")
                employment_status = _status(row["Employment Status"])
            except ValueError as exc:
                errors.append(str(exc))
                hired_date = None
                date_of_birth = None
                employment_status = "employed"

            if not employee_number:
                errors.append("Employee Number is required.")
            elif employee_number.casefold() in existing_numbers:
                errors.append("Employee Number already exists in this company.")
            elif batch_numbers.count(employee_number.casefold()) > 1:
                errors.append("Employee Number is duplicated in this file.")
            if not first_name:
                errors.append("First Name is required.")
            if not last_name:
                errors.append("Last Name is required.")
            if normalized_name in existing_names:
                errors.append(
                    "Possible duplicate employee name already exists in this company."
                )
            elif normalized_name and batch_names.count(normalized_name) > 1:
                errors.append(
                    "Employee name is duplicated in this file after ignoring case and spacing."
                )
            if not email:
                errors.append("Email is required.")
            elif email.casefold() in existing_emails:
                errors.append("Email already exists in this company.")
            elif batch_emails.count(email.casefold()) > 1:
                errors.append("Email is duplicated in this file.")

            resolved_assignments: dict[str, str | None] = {}
            for field_name in ("Manager Employee Number", "Leader Employee Number"):
                resolved_number, assignment_error = _resolve_assignment_reference(
                    row[field_name],
                    assignment_candidates,
                    field_name=field_name,
                )
                resolved_assignments[field_name] = resolved_number
                if assignment_error:
                    errors.append(assignment_error)
                elif resolved_number and resolved_number.casefold() == employee_number.casefold():
                    errors.append(
                        f"{field_name} cannot refer to the employee in the same row."
                    )

            base_username = generated_username(first_name, last_name)
            if len(base_username) < 3:
                errors.append("Generated username is shorter than three characters; edit the name data.")
                username = base_username
            else:
                username = self._unique_username(base_username, reserved_usernames)
            temporary_password = generated_temporary_password(first_name, last_name, employee_number)

            try:
                EmployeeAccountCreate(
                    company_id=company_id,
                    employee_number=employee_number,
                    first_name=first_name,
                    middle_name=_optional_text(row["Middle Name"]),
                    last_name=last_name,
                    suffix=_optional_text(row["Suffix"]),
                    work_email=email,
                    telephone_mobile_no=_optional_text(row["Telephone / Mobile No."]),
                    job_title=_optional_text(row["Job Title / Position"]),
                    hire_date=hired_date,
                    date_of_birth=date_of_birth,
                    gender=_optional_text(row["Gender"]),
                    civil_status=_optional_text(row["Civil Status"]),
                    employment_status=employment_status,
                    department_name=_optional_text(row["Department"]),
                    create_login_account=True,
                    username=username,
                    login_email=email,
                    temporary_password=temporary_password,
                    clearance=2,
                )
            except ValidationError as exc:
                errors.extend(str(error["msg"]).removeprefix("Value error, ") for error in exc.errors())

            normalized = {
                "Row": int(row["Row"]),
                "Employee Number": employee_number,
                "Employee Name": " ".join(part for part in (first_name, _text(row["Middle Name"]), last_name, _text(row["Suffix"])) if part),
                "Hired Date": hired_date.isoformat() if hired_date else "",
                "Years of Service": calculate_years_of_service(hired_date),
                "User ID": "Generated after saving",
                "Clearance": "2 - User",
                "User Name": username,
                "Temporary Password": temporary_password,
                "Validation": "Ready" if not errors else " | ".join(dict.fromkeys(errors)),
                "_filename": filename,
                "_employee": {
                    "employee_number": employee_number,
                    "first_name": first_name,
                    "middle_name": _optional_text(row["Middle Name"]),
                    "last_name": last_name,
                    "suffix": _optional_text(row["Suffix"]),
                    "work_email": email,
                    "telephone_mobile_no": _optional_text(row["Telephone / Mobile No."]),
                    "department_name": _optional_text(row["Department"]),
                    "manager_employee_number": resolved_assignments["Manager Employee Number"],
                    "leader_employee_number": resolved_assignments["Leader Employee Number"],
                    "job_title": _optional_text(row["Job Title / Position"]),
                    "hire_date": hired_date.isoformat() if hired_date else None,
                    "date_of_birth": date_of_birth.isoformat() if date_of_birth else None,
                    "gender": _optional_text(row["Gender"]),
                    "civil_status": _optional_text(row["Civil Status"]),
                    "employment_status": employment_status,
                    "trainings": _training_items(row["Training Checklist"]),
                },
            }
            preview.append(normalized)
        return preview

    def import_preview_rows(
        self,
        preview_rows: list[dict[str, object]],
        *,
        company_id: int,
        current_user_id: int,
        filename: str,
    ) -> list[dict[str, object]]:
        """Revalidate and atomically create employees, accounts, and history."""

        if not preview_rows:
            raise ValueError("There are no employee rows ready to import.")
        if any(str(row.get("Validation", "")) != "Ready" for row in preview_rows):
            raise ValueError("Resolve every validation error before importing employees.")

        usernames = [str(row.get("User Name", "")).strip() for row in preview_rows]
        if len({value.casefold() for value in usernames}) != len(usernames):
            raise ValueError("User Name values must be unique inside the upload preview.")

        clearance_map = {"1 - Admin": 1, "2 - User": 2, "1": 1, "2": 2}
        existing_numbers = {
            value.casefold()
            for value in self.session.scalars(
                select(Employee.employee_number).where(Employee.company_id == company_id)
            ).all()
        }
        existing_emails = {
            value.casefold()
            for value in self.session.scalars(
                select(User.email).where(User.company_id == company_id)
            ).all()
        }
        existing_usernames = {
            value.casefold()
            for value in self.session.scalars(
                select(User.username).where(User.company_id == company_id)
            ).all()
        }
        existing_names = {
            _normalized_employee_name(
                employee.first_name,
                employee.middle_name,
                employee.last_name,
                employee.suffix,
            )
            for employee in self.session.scalars(
                select(Employee).where(Employee.company_id == company_id)
            ).all()
        }

        validated: list[dict[str, object]] = []
        for row in preview_rows:
            employee = dict(row.get("_employee") or {})
            employee_number = str(employee.get("employee_number") or "").strip()
            email = str(employee.get("work_email") or "").strip()
            username = str(row.get("User Name") or "").strip()
            password = str(row.get("Temporary Password") or "")
            clearance_value = str(row.get("Clearance") or "2 - User").strip()
            clearance = clearance_map.get(clearance_value)
            if clearance is None:
                raise ValueError(f"Row {row.get('Row')}: Clearance must be 1 - Admin or 2 - User.")
            if employee_number.casefold() in existing_numbers:
                raise ValueError(f"Row {row.get('Row')}: Employee Number already exists.")
            if email.casefold() in existing_emails:
                raise ValueError(f"Row {row.get('Row')}: Email already exists.")
            if username.casefold() in existing_usernames:
                raise ValueError(f"Row {row.get('Row')}: User Name already exists.")
            normalized_name = _normalized_employee_name(
                employee.get("first_name"),
                employee.get("middle_name"),
                employee.get("last_name"),
                employee.get("suffix"),
            )
            if normalized_name in existing_names:
                raise ValueError(
                    f"Row {row.get('Row')}: Possible duplicate employee name already exists."
                )
            request = EmployeeAccountCreate(
                company_id=company_id,
                employee_number=employee_number,
                first_name=str(employee.get("first_name") or ""),
                middle_name=employee.get("middle_name"),
                last_name=str(employee.get("last_name") or ""),
                suffix=employee.get("suffix"),
                work_email=email,
                telephone_mobile_no=employee.get("telephone_mobile_no"),
                job_title=employee.get("job_title"),
                hire_date=date.fromisoformat(str(employee["hire_date"])),
                date_of_birth=(date.fromisoformat(str(employee["date_of_birth"])) if employee.get("date_of_birth") else None),
                gender=employee.get("gender"),
                civil_status=employee.get("civil_status"),
                employment_status=str(employee.get("employment_status") or "employed"),
                department_name=employee.get("department_name"),
                create_login_account=True,
                username=username,
                login_email=email,
                temporary_password=password,
                clearance=clearance,
            )
            validated.append({"row": row, "employee": employee, "request": request})
            existing_numbers.add(employee_number.casefold())
            existing_emails.add(email.casefold())
            existing_usernames.add(username.casefold())
            existing_names.add(normalized_name)

        roles = {
            role.name: role
            for role in self.session.scalars(
                select(Role).where(Role.company_id == company_id)
            ).all()
        }
        if "company_admin" not in roles or "employee" not in roles:
            raise ValueError("Internal account access mapping is unavailable. Run the initial-data script.")

        departments = {
            department.name.casefold(): department
            for department in self.session.scalars(
                select(Department).where(Department.company_id == company_id)
            ).all()
        }
        employee_map = {
            employee.employee_number.casefold(): employee
            for employee in self.session.scalars(
                select(Employee).where(Employee.company_id == company_id)
            ).all()
        }
        batch_id = f"EMP-{uuid4().hex[:12].upper()}"
        results: list[dict[str, object]] = []
        created_pairs: list[tuple[Employee, dict[str, object], EmployeeAccountCreate]] = []

        try:
            for item in validated:
                request = item["request"]
                employee_values = item["employee"]
                department = None
                if request.department_name:
                    key = request.department_name.strip().casefold()
                    department = departments.get(key)
                    if department is None:
                        department = Department(
                            company_id=company_id,
                            name=request.department_name.strip(),
                            code=None,
                            is_active=True,
                        )
                        self.session.add(department)
                        self.session.flush()
                        departments[key] = department

                role = roles["company_admin" if request.clearance == 1 else "employee"]
                user = User(
                    company_id=company_id,
                    role_id=role.id,
                    clearance=request.clearance,
                    username=str(request.username).strip(),
                    email=str(request.login_email),
                    password_hash=self.password_manager.hash_password(str(request.temporary_password)),
                    is_active=request.employment_status == "employed",
                    must_change_password=True,
                )
                self.session.add(user)
                self.session.flush()

                employee = Employee(
                    company_id=company_id,
                    user_id=user.id,
                    department_id=department.id if department else None,
                    employee_number=request.employee_number.strip(),
                    first_name=request.first_name.strip(),
                    middle_name=request.middle_name.strip() if request.middle_name else None,
                    last_name=request.last_name.strip(),
                    suffix=request.suffix.strip() if request.suffix else None,
                    work_email=str(request.work_email),
                    telephone_mobile_no=request.telephone_mobile_no,
                    job_title=request.job_title,
                    employment_status=request.employment_status,
                    hire_date=request.hire_date,
                    date_of_birth=request.date_of_birth,
                    gender=request.gender,
                    civil_status=request.civil_status,
                    archived_at=(datetime.now(timezone.utc) if request.employment_status == "resigned" else None),
                    archived_by_user_id=(current_user_id if request.employment_status == "resigned" else None),
                )
                self.session.add(employee)
                self.session.flush()
                employee_map[employee.employee_number.casefold()] = employee

                for training in employee_values.get("trainings", []):
                    self.session.add(
                        EmployeeTraining(
                            company_id=company_id,
                            employee_id=employee.id,
                            title=str(training["title"]),
                            is_completed=bool(training["is_completed"]),
                            display_order=int(training["display_order"]),
                        )
                    )
                created_pairs.append((employee, employee_values, request))
                results.append(
                    {
                        "Row": item["row"]["Row"],
                        "Employee Number": employee.employee_number,
                        "Employee Name": employee.full_name,
                        "User ID": user.id,
                        "Clearance": f"{user.clearance} - {'Admin' if user.clearance == 1 else 'User'}",
                        "User Name": user.username,
                        "Temporary Password": str(request.temporary_password),
                        "Status": "Imported",
                    }
                )

            for employee, employee_values, request in created_pairs:
                manager_number = str(employee_values.get("manager_employee_number") or "").casefold()
                leader_number = str(employee_values.get("leader_employee_number") or "").casefold()
                manager = employee_map.get(manager_number) if manager_number else None
                leader = employee_map.get(leader_number) if leader_number else None
                if manager_number and manager is None:
                    raise ValueError(f"Manager Employee Number '{manager_number}' is unavailable.")
                if leader_number and leader is None:
                    raise ValueError(f"Leader Employee Number '{leader_number}' is unavailable.")
                if manager is employee or leader is employee:
                    raise ValueError(f"Employee {employee.employee_number} cannot be their own manager or leader.")
                if manager is not None and manager.employment_status != "employed":
                    raise ValueError(f"Manager '{manager.employee_number}' must be an employed record.")
                if leader is not None and leader.employment_status != "employed":
                    raise ValueError(f"Leader '{leader.employee_number}' must be an employed record.")
                employee.manager_id = manager.id if manager else None
                employee.leader_id = leader.id if leader else None
                self.session.add(
                    EmployeeHistory(
                        company_id=company_id,
                        employee_id=employee.id,
                        employee_number=employee.employee_number,
                        employee_name=employee.full_name,
                        performed_by_user_id=current_user_id,
                        action_type=("archived" if employee.employment_status == "resigned" else "bulk_imported"),
                        source="excel_upload",
                        summary=("Imported through Excel and archived as Resigned." if employee.employment_status == "resigned" else "Employee and login account imported through Excel."),
                        new_values_json=json.dumps({"employment_status": employee.employment_status, "clearance": request.clearance, "username": str(request.username)}, sort_keys=True),
                        upload_batch_id=batch_id,
                        upload_filename=filename[:255],
                    )
                )
            self.session.commit()
            return results
        except Exception:
            self.session.rollback()
            raise

    @staticmethod
    def build_import_result(results: list[dict[str, object]]) -> bytes:
        """Return controlled initial credentials after a successful import."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Import Result"
        headers = list(results[0].keys()) if results else ["Status"]
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="2F740B")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in results:
            sheet.append([row.get(header, "") for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()
