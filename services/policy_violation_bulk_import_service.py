"""Validated, atomic Excel import for violation master-list definitions."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy.orm import Session

from models.policy_violation import PolicyViolation
from schemas.policy_violation_schema import PolicyViolationCreateRequest
from services.policy_service import PolicyService
from services.excel_template_guidance import (
    add_header_guidance,
    add_inline_list_validation,
    identifier_name_reference,
)
from services.policy_violation_service import (
    PolicyViolationService,
    VIOLATION_SEVERITIES,
    VIOLATION_STATUSES,
)


VIOLATION_IMPORT_COLUMNS = (
    "Violation Code",
    "Category",
    "Violation / Offense",
    "Severity",
    "Description",
    "1st Offense",
    "2nd Offense",
    "3rd Offense",
    "Final / Maximum Action",
    "Related Policy ID",
    "Effective Date",
    "Status",
    "Notes",
)

VIOLATION_PREVIEW_COLUMNS = (
    "Row",
    *VIOLATION_IMPORT_COLUMNS,
    "Validation",
)

SAMPLE_VIOLATION_CODE = "SAMPLE ONLY - ATT-001"
SAMPLE_VIOLATION_ROW = (
    SAMPLE_VIOLATION_CODE,
    "Attendance",
    "Habitual Tardiness",
    "Minor",
    "Repeated late arrival beyond the company attendance rule.",
    "Verbal Warning",
    "Written Warning",
    "Suspension",
    "Termination",
    "",
    "2026-08-24",
    "Active",
    "Replace or delete this sample row before upload.",
)

VIOLATION_STATUS_OPTIONS = tuple(value.title() for value in VIOLATION_STATUSES)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional_text(value: object) -> str | None:
    value_text = _text(value)
    return value_text or None


def _parse_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(_text(value))
    except ValueError as exc:
        raise ValueError("Effective Date must use YYYY-MM-DD or be blank.") from exc


def _validation_message(error: ValidationError) -> str:
    messages: list[str] = []
    for item in error.errors():
        location = " / ".join(str(value) for value in item.get("loc", ()))
        message = str(item.get("msg", "Invalid value"))
        messages.append(f"{location}: {message}" if location else message)
    return " | ".join(messages[:4]) or "Invalid row values."


class PolicyViolationBulkImportService:
    """Prepare a safe Excel preview and atomically create valid violations."""

    def __init__(self, session: Session) -> None:
        self.session = session
        self.violation_service = PolicyViolationService(session)
        self.policy_service = PolicyService(session)

    def _policy_reference_data(
        self, company_id: int
    ) -> tuple[list[object], dict[int, object], dict[str, int], dict[int, str]]:
        policies = self.policy_service.list_for_admin(company_id)
        document_map = self.policy_service.get_document_map(
            company_id=company_id, policies=policies
        )
        lookup: dict[str, int] = {}
        display_by_id: dict[int, str] = {}
        for policy in policies:
            public_id = PolicyService.public_id_for(policy).strip()
            document = document_map.get(policy.id)
            filename = (
                document.original_filename.strip()
                if document is not None and str(document.original_filename or "").strip()
                else ""
            )
            title = str(policy.title or "").strip()
            filename_or_title = filename or title or "[Untitled policy]"
            display = identifier_name_reference(public_id, filename_or_title)
            if public_id:
                lookup[public_id.casefold()] = policy.id  # legacy/raw ID remains accepted
            if display:
                lookup[display.casefold()] = policy.id
                display_by_id[policy.id] = display
        return policies, document_map, lookup, display_by_id

    def build_template(self, *, company_id: int) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Violation Import Template"
        sheet.append(list(VIOLATION_IMPORT_COLUMNS))
        sheet.append(list(SAMPLE_VIOLATION_ROW))

        header_fill = PatternFill("solid", fgColor="2F740B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        sample_fill = PatternFill("solid", fgColor="D9D9D9")
        sample_font = Font(color="666666", italic=True)
        for cell in sheet[2]:
            cell.fill = sample_fill
            cell.font = sample_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[2].height = 48

        widths = {
            "A": 22, "B": 22, "C": 32, "D": 16, "E": 48,
            "F": 34, "G": 34, "H": 34, "I": 36, "J": 56,
            "K": 18, "L": 14, "M": 44,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:M2000"

        policies, policy_documents, _policy_lookup, policy_display_by_id = (
            self._policy_reference_data(company_id)
        )
        policy_choices = tuple(
            policy_display_by_id[policy.id]
            for policy in policies
            if policy.id in policy_display_by_id
        )

        add_header_guidance(
            sheet, "A1",
            requirement="Required and unique inside the company. Letters/numbers plus -, _, ., and / are allowed; free text.",
        )
        add_header_guidance(sheet, "B1", requirement="Required violation category; free text. There is no fixed category list.")
        add_header_guidance(sheet, "C1", requirement="Required violation/offense name; free text.")
        add_header_guidance(
            sheet, "D1",
            requirement="Required. Use one of the same Severity choices available in the violation master form.",
            choices=VIOLATION_SEVERITIES,
        )
        add_header_guidance(sheet, "E1", requirement="Required violation description; free text.")
        add_header_guidance(sheet, "F1", requirement="Required disciplinary action for the 1st offense; free text.")
        add_header_guidance(sheet, "G1", requirement="Required disciplinary action for the 2nd offense; free text.")
        add_header_guidance(sheet, "H1", requirement="Required disciplinary action for the 3rd offense; free text.")
        add_header_guidance(sheet, "I1", requirement="Required final/maximum disciplinary action; free text.")
        add_header_guidance(
            sheet, "J1",
            requirement="Optional related policy reference. Choose or enter: Policy ID - Filename / Title (whichever is available). Leave blank for no policy link.",
            choices=policy_choices,
            choices_heading="Current company policy selections available at download time:",
            extra=(
                "The same current company references are listed in the 'Related Policies' worksheet. Leave blank for no policy link."
                if policy_choices
                else "No active company policy references are currently available. The 'Related Policies' worksheet is currently empty; leave this field blank unless a policy is added before downloading a fresh template."
            ),
        )
        add_header_guidance(sheet, "K1", requirement="Optional effective date. Use YYYY-MM-DD or an Excel date cell; blank means immediate.")
        add_header_guidance(
            sheet, "L1",
            requirement="Required. Use one of the current violation Status selections.",
            choices=VIOLATION_STATUS_OPTIONS,
        )
        add_header_guidance(sheet, "M1", requirement="Optional notes; free text.")

        add_inline_list_validation(
            sheet, "D2:D2000", VIOLATION_SEVERITIES, allow_blank=False
        )
        add_inline_list_validation(
            sheet, "L2:L2000", VIOLATION_STATUS_OPTIONS, allow_blank=False
        )

        policies_sheet = workbook.create_sheet("Related Policies")
        policies_sheet.append(["Policy ID", "Policy Reference", "Filename", "Policy / Version"])
        for cell in policies_sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for policy in policies:
            document = policy_documents.get(policy.id)
            filename = (
                document.original_filename.strip()
                if document is not None and str(document.original_filename or "").strip()
                else ""
            )
            title = str(policy.title or "").strip()
            filename_or_title = filename or title or "[Untitled policy]"
            policies_sheet.append([
                PolicyService.public_id_for(policy),
                policy_display_by_id.get(
                    policy.id,
                    identifier_name_reference(
                        PolicyService.public_id_for(policy),
                        filename_or_title,
                    ),
                ),
                filename or "[No uploaded file]",
                f"{policy.title} · v{policy.version}",
            ])
        policies_sheet.column_dimensions["A"].width = 24
        policies_sheet.column_dimensions["B"].width = 64
        policies_sheet.column_dimensions["C"].width = 56
        policies_sheet.column_dimensions["D"].width = 68
        policies_sheet.freeze_panes = "A2"
        if policies:
            workbook.defined_names.add(
                DefinedName(
                    "ViolationRelatedPolicyIds",
                    attr_text=f"'Related Policies'!$B$2:$B${len(policies) + 1}",
                )
            )
            policy_validation = DataValidation(
                type="list",
                formula1="=ViolationRelatedPolicyIds",
                allow_blank=True,
            )
            policy_validation.errorTitle = "Invalid Policy reference"
            policy_validation.error = (
                "Choose a current Policy ID - Filename / Title reference from the dropdown/Related Policies worksheet, "
                "or leave the cell blank."
            )
            policy_validation.promptTitle = "Current policy references"
            policy_validation.prompt = (
                "Choose the current Policy ID - Filename / Title reference available in this company."
            )
            policy_validation.showErrorMessage = True
            policy_validation.showInputMessage = True
            sheet.add_data_validation(policy_validation)
            policy_validation.add("J2:J2000")

        instructions = workbook.create_sheet("Instructions")
        instruction_rows = (
            ("Column", "Requirement"),
            ("Violation Code", "Required and unique inside the company. Letters/numbers plus -, _, ., / are allowed."),
            ("Category", "Required."),
            ("Violation / Offense", "Required."),
            ("Severity", "Required. Use the dropdown or hover the Severity header to view the complete current selections."),
            ("Description", "Required."),
            ("1st/2nd/3rd/Final", "All four disciplinary-action fields are required."),
            ("Related Policy ID", "Optional. Choose Policy ID - Filename / Title (whichever is available) from the Related Policies worksheet; leave blank for no link. Legacy raw Policy ID remains accepted during import."),
            ("Effective Date", "Optional. Use YYYY-MM-DD; blank means immediate."),
            ("Status", "Required. Use the dropdown or hover the Status header to view the complete current selections."),
            ("Notes", "Optional."),
            ("Gray Sample Row", "Row 2 is an example only. Delete it or replace it. An unchanged SAMPLE ONLY row is ignored."),
            ("Import Safety", "Every row must validate before import. The batch is saved atomically; no partial batch is created."),
        )
        for row in instruction_rows:
            instructions.append(row)
        for cell in instructions[1]:
            cell.fill = header_fill
            cell.font = header_font
        instructions.column_dimensions["A"].width = 30
        instructions.column_dimensions["B"].width = 110
        instructions.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _raw_values(
        self,
        row: dict[str, object],
        *,
        company_id: int,
        policy_lookup: dict[str, int],
    ) -> dict[str, object]:
        policy_reference = _text(row.get("Related Policy ID"))
        related_policy_id: int | None = None
        if policy_reference:
            related_policy_id = policy_lookup.get(policy_reference.casefold())
            if related_policy_id is None:
                raise ValueError(
                    f"Related Policy reference {policy_reference} is not an active policy in this company."
                )
        return {
            "company_id": company_id,
            "created_by_user_id": 1,  # Preview-only Pydantic validation; actor is replaced on import.
            "violation_code": _text(row.get("Violation Code")),
            "category": _text(row.get("Category")),
            "offense_title": _text(row.get("Violation / Offense")),
            "severity": _text(row.get("Severity")),
            "description": _text(row.get("Description")),
            "first_offense_action": _text(row.get("1st Offense")),
            "second_offense_action": _text(row.get("2nd Offense")),
            "third_offense_action": _text(row.get("3rd Offense")),
            "final_action": _text(row.get("Final / Maximum Action")),
            "related_policy_id": related_policy_id,
            "effective_date": _parse_date(row.get("Effective Date")),
            "status": _text(row.get("Status")) or "Active",
            "notes": _optional_text(row.get("Notes")),
        }

    def prepare_preview(
        self,
        file_bytes: bytes,
        *,
        filename: str,
        company_id: int,
    ) -> list[dict[str, object]]:
        try:
            workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        except Exception as exc:
            raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc

        sheet_name = "Violation Import Template"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                "Use the downloadable template and keep the 'Violation Import Template' worksheet name."
            )
        sheet = workbook[sheet_name]
        headers = tuple(_text(cell.value) for cell in sheet[1])
        if headers != VIOLATION_IMPORT_COLUMNS:
            raise ValueError(
                "The Excel columns or their order do not match the downloadable Violation template."
            )

        raw_rows: list[dict[str, object]] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            if _text(values[0]).casefold() == SAMPLE_VIOLATION_CODE.casefold():
                continue
            raw_rows.append({"Row": row_number, **dict(zip(VIOLATION_IMPORT_COLUMNS, values))})
        if not raw_rows:
            raise ValueError("The Violation template does not contain any violation rows.")
        if len(raw_rows) > 1000:
            raise ValueError("A single Violation import is limited to 1,000 rows.")

        _policies, _documents, policy_lookup, policy_display_by_id = self._policy_reference_data(company_id)
        normalized_codes: list[str | None] = []
        row_values: list[dict[str, object] | None] = []
        row_errors: list[list[str]] = []
        for row in raw_rows:
            errors: list[str] = []
            normalized: dict[str, object] | None = None
            code: str | None = None
            try:
                values = self._raw_values(row, company_id=company_id, policy_lookup=policy_lookup)
                request_values = PolicyViolationCreateRequest(**values)
                normalized = self.violation_service._normalized_values(
                    company_id=company_id,
                    violation_code=request_values.violation_code,
                    category=request_values.category,
                    offense_title=request_values.offense_title,
                    description=request_values.description,
                    severity=request_values.severity,
                    first_offense_action=request_values.first_offense_action,
                    second_offense_action=request_values.second_offense_action,
                    third_offense_action=request_values.third_offense_action,
                    final_action=request_values.final_action,
                    related_policy_id=request_values.related_policy_id,
                    effective_date=request_values.effective_date,
                    status=request_values.status,
                    notes=request_values.notes,
                )
                code = str(normalized["violation_code"])
            except ValidationError as exc:
                errors.append(_validation_message(exc))
            except ValueError as exc:
                errors.append(str(exc))
            normalized_codes.append(code)
            row_values.append(normalized)
            row_errors.append(errors)

        counts: dict[str, int] = {}
        for code in normalized_codes:
            if code:
                counts[code] = counts.get(code, 0) + 1

        preview: list[dict[str, object]] = []
        for row, code, normalized, errors in zip(raw_rows, normalized_codes, row_values, row_errors):
            if code and counts.get(code, 0) > 1:
                errors.append(f"Violation Code {code} is duplicated inside the uploaded file.")
            if code and self.violation_service.repository.get_by_code(
                company_id=company_id,
                violation_code=code,
            ) is not None:
                errors.append(f"Violation Code {code} already exists in this company.")

            safe = normalized or {}
            policy_id = safe.get("related_policy_id")
            policy_reference = policy_display_by_id.get(int(policy_id), "") if policy_id else ""
            effective = safe.get("effective_date")
            preview.append({
                "Row": int(row["Row"]),
                "Violation Code": safe.get("violation_code", _text(row.get("Violation Code"))),
                "Category": safe.get("category", _text(row.get("Category"))),
                "Violation / Offense": safe.get("offense_title", _text(row.get("Violation / Offense"))),
                "Severity": safe.get("severity", _text(row.get("Severity"))),
                "Description": safe.get("description", _text(row.get("Description"))),
                "1st Offense": safe.get("first_offense_action", _text(row.get("1st Offense"))),
                "2nd Offense": safe.get("second_offense_action", _text(row.get("2nd Offense"))),
                "3rd Offense": safe.get("third_offense_action", _text(row.get("3rd Offense"))),
                "Final / Maximum Action": safe.get("final_action", _text(row.get("Final / Maximum Action"))),
                "Related Policy ID": policy_reference or _text(row.get("Related Policy ID")),
                "Effective Date": effective.isoformat() if isinstance(effective, date) else _text(row.get("Effective Date")),
                "Status": str(safe.get("status", _text(row.get("Status")) or "active")).title(),
                "Notes": safe.get("notes") or "",
                "Validation": "Ready" if not errors else " | ".join(dict.fromkeys(errors)),
                "_filename": filename,
                "_values": safe,
            })
        return preview

    def import_preview_rows(
        self,
        preview_rows: list[dict[str, object]],
        *,
        company_id: int,
        actor_user_id: int,
    ) -> list[PolicyViolation]:
        if not preview_rows:
            raise ValueError("There are no violation rows ready to import.")
        if any(str(row.get("Validation", "")) != "Ready" for row in preview_rows):
            raise ValueError("Resolve every validation error before importing violations.")

        requests: list[PolicyViolationCreateRequest] = []
        for row in preview_rows:
            values = dict(row.get("_values") or {})
            if not values:
                raise ValueError(f"Row {row.get('Row')} must be validated again before import.")
            requests.append(PolicyViolationCreateRequest(
                company_id=company_id,
                created_by_user_id=actor_user_id,
                violation_code=str(values["violation_code"]),
                category=str(values["category"]),
                offense_title=str(values["offense_title"]),
                description=str(values["description"]),
                severity=str(values["severity"]),
                first_offense_action=str(values["first_offense_action"]),
                second_offense_action=str(values["second_offense_action"]),
                third_offense_action=str(values["third_offense_action"]),
                final_action=str(values["final_action"]),
                related_policy_id=values.get("related_policy_id"),
                effective_date=values.get("effective_date"),
                status=str(values["status"]),
                notes=values.get("notes"),
            ))
        return self.violation_service.create_many(requests)
