"""Validated, atomic Excel import for employee disciplinary cases."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy.orm import Session

from repositories.employee_repository import EmployeeRepository
from repositories.policy_violation_repository import PolicyViolationRepository
from repositories.user_repository import UserRepository
from schemas.disciplinary_record_schema import DisciplinaryRecordCreateRequest
from services.disciplinary_record_service import (
    ACKNOWLEDGMENT_STATUSES,
    CASE_STATUSES,
    DisciplinaryRecordService,
)
from services.excel_template_guidance import (
    add_header_guidance,
    add_inline_list_validation,
    employee_reference_label,
    identifier_name_reference,
)


DISCIPLINARY_IMPORT_COLUMNS = (
    "Employee Number",
    "Violation Code",
    "Incident Date",
    "Incident / Case Description",
    "Evidence / Remarks",
    "Project / Team / Department",
    "Actual Action Taken",
    "Issued By",
    "Reviewed / Approved By",
    "Date Issued",
    "Employee Acknowledgment",
    "Case Status",
    "Notes",
)

LEGACY_DISCIPLINARY_IMPORT_COLUMNS = tuple(
    "Issued By Username" if column == "Issued By" else
    "Reviewed / Approved By Username" if column == "Reviewed / Approved By" else
    column
    for column in DISCIPLINARY_IMPORT_COLUMNS
)

DISCIPLINARY_PREVIEW_COLUMNS = (
    "Row",
    *DISCIPLINARY_IMPORT_COLUMNS,
    "Validation",
)

SAMPLE_DISCIPLINARY_ROW = (
    "SAMPLE-EMP-001 - Dela Cruz, Juan",
    "ATT-001 - Habitual Tardiness",
    "2026-08-24",
    "Sample incident description. Replace or delete this row.",
    "Optional evidence or remarks.",
    "Sample Team / Department",
    "",
    "",
    "",
    "",
    "Pending",
    "Draft",
    "Replace or delete this sample row before upload.",
)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional_text(value: object) -> str | None:
    text = _text(value)
    return text or None


def _parse_date(value: object, *, field_name: str, required: bool) -> date | None:
    if value in (None, ""):
        if required:
            raise ValueError(f"{field_name} is required.")
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(_text(value))
    except ValueError as exc:
        raise ValueError(f"{field_name} must use YYYY-MM-DD or an Excel date cell.") from exc


def _validation_message(error: ValidationError) -> str:
    messages: list[str] = []
    for item in error.errors():
        location = " / ".join(str(value) for value in item.get("loc", ()))
        message = str(item.get("msg", "Invalid value"))
        messages.append(f"{location}: {message}" if location else message)
    return " | ".join(messages[:4]) or "Invalid row values."


def _reference_validation(sheet, cell_range: str, formula: str, *, allow_blank: bool) -> None:
    validation = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
    validation.errorTitle = "Invalid selection"
    validation.error = "Choose a value from the corresponding reference worksheet."
    validation.promptTitle = "Project reference"
    validation.prompt = "Choose a valid project value from the dropdown/reference worksheet."
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add(cell_range)



def _employee_reference(employee) -> str:
    return employee_reference_label(
        employee.employee_number,
        employee.last_name,
        employee.first_name,
        employee.suffix,
    )


def _user_reference(user) -> str:
    employee = user.employee
    if employee is not None and employee.employee_number:
        return _employee_reference(employee)
    return identifier_name_reference(user.username, "Account without employee record")


def _lookup_with_display(items, *, raw_value, display_value) -> dict[str, object]:
    """Map both legacy/raw and new human-readable Excel references."""

    lookup: dict[str, object] = {}
    for item in items:
        raw = str(raw_value(item) or "").strip()
        display = str(display_value(item) or "").strip()
        if raw:
            lookup[raw.casefold()] = item
        if display:
            lookup[display.casefold()] = item
    return lookup


class DisciplinaryRecordBulkImportService:
    """Build a company-aware template and atomically import validated cases."""

    def __init__(self, session: Session) -> None:
        self.session = session
        self.employee_repository = EmployeeRepository(session)
        self.violation_repository = PolicyViolationRepository(session)
        self.user_repository = UserRepository(session)
        self.case_service = DisciplinaryRecordService(session)

    def _employees(self, company_id: int):
        return self.employee_repository.list_with_details(company_id, archived=False)

    def _violations(self, company_id: int):
        return [
            item
            for item in self.violation_repository.list_current(company_id)
            if item.status == "active" and item.archived_at is None
        ]

    def _users(self, company_id: int):
        """Return active company users accepted by the legacy/raw importer."""

        return [
            item
            for item in self.user_repository.list_with_details(company_id)
            if item.is_active
        ]

    def _reference_users(self, company_id: int):
        """Return users eligible for the new human-readable Excel dropdown.

        The requested project standard for Issued By / Reviewed By is strictly
        ``Employee Number - Last Name, First Name Suffix``.  Therefore the
        current dropdown/reference sheet only exposes active users that are
        linked to an employed employee record.  Raw usernames remain accepted
        by the importer for backward compatibility (including system/admin
        accounts without an employee record).
        """

        return [
            item
            for item in self._users(company_id)
            if item.employee is not None
            and item.employee.employee_number
            and item.employee.employment_status == "employed"
        ]

    def build_template(self, *, company_id: int) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Disciplinary Case Import"
        sheet.append(list(DISCIPLINARY_IMPORT_COLUMNS))
        sheet.append(list(SAMPLE_DISCIPLINARY_ROW))

        header_fill = PatternFill("solid", fgColor="2F740B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sample_fill = PatternFill("solid", fgColor="D9D9D9")
        sample_font = Font(color="666666", italic=True)
        for cell in sheet[2]:
            cell.fill = sample_fill
            cell.font = sample_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[2].height = 54
        widths = {
            "A": 46, "B": 44, "C": 16, "D": 48, "E": 40, "F": 34,
            "G": 38, "H": 46, "I": 46, "J": 16, "K": 26, "L": 18, "M": 40,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:M2000"

        employees = self._employees(company_id)
        violations = self._violations(company_id)
        users = self._reference_users(company_id)
        employee_choices = tuple(_employee_reference(employee) for employee in employees)
        violation_choices = tuple(
            identifier_name_reference(item.violation_code, item.offense_title)
            for item in violations
        )
        user_choices = tuple(_user_reference(user) for user in users)

        add_header_guidance(
            sheet, "A1",
            requirement="Required. Choose or enter the current employee reference in the format: Employee Number - Last Name, First Name Suffix.",
            choices=employee_choices,
            choices_heading="Current active employee selections available at download time:",
            extra=(
                "The same current company selections are listed in the 'Employee References' worksheet."
                if employee_choices
                else "No active employees are currently available for a disciplinary case. Download a fresh template after an employee becomes available."
            ),
        )
        add_header_guidance(
            sheet, "B1",
            requirement="Required. Choose or enter the current violation reference in the format: Violation Code - Violation / Offense.",
            choices=violation_choices,
            choices_heading="Current active violation selections available at download time:",
            extra=(
                "The same current company selections are listed in the 'Violation References' worksheet. Suggested penalties are not entered here; they stay linked to the master violation."
                if violation_choices
                else "No active master violations are currently available. Add/activate a violation and download a fresh template before importing cases."
            ),
        )
        add_header_guidance(sheet, "C1", requirement="Required incident date. Use YYYY-MM-DD or an Excel date cell.")
        add_header_guidance(sheet, "D1", requirement="Required incident/case description; free text.")
        add_header_guidance(sheet, "E1", requirement="Optional evidence/remarks; free text.")
        add_header_guidance(sheet, "F1", requirement="Optional project/team/department snapshot; free text. Blank defaults to the employee's current department at import time.")
        add_header_guidance(sheet, "G1", requirement="Optional while Draft/For Review; required when Case Status is Issued or Closed. Free text.")
        add_header_guidance(
            sheet, "H1",
            requirement="Optional while Draft/For Review; required when Case Status is Issued or Closed. Choose or enter: Employee Number - Last Name, First Name Suffix.",
            choices=user_choices,
            choices_heading="Current active company user selections available at download time:",
            extra=(
                "The same current company selections are listed in the 'User References' worksheet."
                if user_choices
                else "No active employee-linked user references are currently available. Download a fresh template after an authorized employee account becomes available."
            ),
        )
        add_header_guidance(
            sheet, "I1",
            requirement="Optional reviewer/approver. Choose or enter: Employee Number - Last Name, First Name Suffix, or leave blank.",
            choices=user_choices,
            choices_heading="Current active company user selections available at download time:",
            extra=(
                "The same current company selections are listed in the 'User References' worksheet."
                if user_choices
                else "No active employee-linked user references are currently available. Leave blank where allowed and download a fresh template after an authorized employee account becomes available."
            ),
        )
        add_header_guidance(sheet, "J1", requirement="Optional while Draft/For Review; required when Case Status is Issued or Closed. Use YYYY-MM-DD or an Excel date cell.")
        add_header_guidance(
            sheet, "K1",
            requirement="Required. Use one of the current Employee Acknowledgment selections.",
            choices=ACKNOWLEDGMENT_STATUSES,
        )
        add_header_guidance(
            sheet, "L1",
            requirement="Required. Use one of the current disciplinary Case Status selections.",
            choices=CASE_STATUSES,
        )
        add_header_guidance(sheet, "M1", requirement="Optional notes; free text.")
        add_inline_list_validation(sheet, "K2:K2000", ACKNOWLEDGMENT_STATUSES, allow_blank=False)
        add_inline_list_validation(sheet, "L2:L2000", CASE_STATUSES, allow_blank=False)

        emp_sheet = workbook.create_sheet("Employee References")
        emp_sheet.append(["Employee Number", "Employee Reference", "Employee Name", "Department", "Position"])
        for employee in employees:
            reference = _employee_reference(employee)
            display_name = reference.split(" - ", 1)[1] if " - " in reference else reference
            emp_sheet.append([
                employee.employee_number,
                reference,
                display_name,
                employee.department.name if employee.department else "",
                employee.job_title or "",
            ])
        for cell in emp_sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        emp_sheet.freeze_panes = "A2"
        if employees:
            workbook.defined_names.add(
                DefinedName(
                    "DisciplinaryEmployeeNumbers",
                    attr_text=f"'Employee References'!$B$2:$B${len(employees) + 1}",
                )
            )
            _reference_validation(
                sheet, "A2:A2000", "=DisciplinaryEmployeeNumbers", allow_blank=False
            )

        vio_sheet = workbook.create_sheet("Violation References")
        vio_sheet.append(["Violation Code", "Violation Reference", "Violation / Offense", "Category", "Severity"])
        for item in violations:
            vio_sheet.append([
                item.violation_code,
                identifier_name_reference(item.violation_code, item.offense_title),
                item.offense_title,
                item.category,
                item.severity,
            ])
        for cell in vio_sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        vio_sheet.freeze_panes = "A2"
        if violations:
            workbook.defined_names.add(
                DefinedName(
                    "DisciplinaryViolationCodes",
                    attr_text=f"'Violation References'!$B$2:$B${len(violations) + 1}",
                )
            )
            _reference_validation(
                sheet, "B2:B2000", "=DisciplinaryViolationCodes", allow_blank=False
            )

        user_sheet = workbook.create_sheet("User References")
        user_sheet.append(["Username", "User Reference", "Employee Number", "Employee Name", "Clearance"])
        for user in users:
            employee = user.employee
            reference = _user_reference(user)
            employee_name = (
                _employee_reference(employee).split(" - ", 1)[1]
                if employee is not None and " - " in _employee_reference(employee)
                else (_employee_reference(employee) if employee is not None else "")
            )
            user_sheet.append([
                user.username,
                reference,
                employee.employee_number if employee is not None else "",
                employee_name,
                user.clearance,
            ])
        for cell in user_sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        user_sheet.freeze_panes = "A2"
        if users:
            workbook.defined_names.add(
                DefinedName(
                    "DisciplinaryUsernames",
                    attr_text=f"'User References'!$B$2:$B${len(users) + 1}",
                )
            )
            _reference_validation(sheet, "H2:H2000", "=DisciplinaryUsernames", allow_blank=True)
            _reference_validation(sheet, "I2:I2000", "=DisciplinaryUsernames", allow_blank=True)

        for ref_sheet in (emp_sheet, vio_sheet, user_sheet):
            for column in ref_sheet.columns:
                letter = column[0].column_letter
                ref_sheet.column_dimensions[letter].width = min(
                    48,
                    max(16, max(len(str(cell.value or "")) for cell in column) + 2),
                )

        stream = BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def prepare_preview(self, file_bytes: bytes, *, filename: str, company_id: int) -> list[dict[str, object]]:
        if not filename.lower().endswith(".xlsx"):
            raise ValueError("Upload the downloadable .xlsx disciplinary case template.")
        if not file_bytes:
            raise ValueError("The uploaded disciplinary case file is empty.")
        if len(file_bytes) > 10 * 1024 * 1024:
            raise ValueError("The Disciplinary Case Excel file must not exceed 10 MB.")
        try:
            workbook = load_workbook(BytesIO(file_bytes), data_only=True)
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise ValueError(
                "The uploaded disciplinary case file is not a readable .xlsx workbook."
            ) from exc
        if "Disciplinary Case Import" not in workbook.sheetnames:
            raise ValueError("The workbook does not contain the required 'Disciplinary Case Import' worksheet.")
        sheet = workbook["Disciplinary Case Import"]
        headers = tuple(_text(cell.value) for cell in sheet[1])
        current_headers = headers[: len(DISCIPLINARY_IMPORT_COLUMNS)]
        if current_headers not in (DISCIPLINARY_IMPORT_COLUMNS, LEGACY_DISCIPLINARY_IMPORT_COLUMNS):
            raise ValueError("The disciplinary case template columns were changed. Download a fresh template and try again.")

        active_employees = self._employees(company_id)
        active_violations = self._violations(company_id)
        active_users = self._users(company_id)
        reference_users = self._reference_users(company_id)
        employee_lookup = _lookup_with_display(
            active_employees, raw_value=lambda item: item.employee_number, display_value=_employee_reference
        )
        violation_lookup = _lookup_with_display(
            active_violations,
            raw_value=lambda item: item.violation_code,
            display_value=lambda item: identifier_name_reference(item.violation_code, item.offense_title),
        )
        # Raw usernames remain accepted for backward compatibility, while the
        # new human-readable display labels are only registered for active
        # employee-linked users shown by the current template.
        user_lookup = {
            str(item.username).strip().casefold(): item
            for item in active_users
            if str(item.username or "").strip()
        }
        user_lookup.update(
            _lookup_with_display(
                reference_users,
                raw_value=lambda item: item.username,
                display_value=_user_reference,
            )
        )
        for user in active_users:
            if user.employee is not None and user.employee.employee_number:
                user_lookup[user.employee.employee_number.strip().casefold()] = user
        existing_signatures = {
            (
                item.employee_number.strip().casefold(),
                item.violation_code.strip().casefold(),
                item.incident_date.isoformat(),
                " ".join(item.incident_description.split()).casefold(),
            )
            for item in self.case_service.repository.list_company(company_id)
        }

        rows: list[dict[str, object]] = []
        seen_signatures: set[tuple[str, str, str, str]] = set()
        nonblank_count = 0
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, max_col=len(DISCIPLINARY_IMPORT_COLUMNS), values_only=True), start=2):
            if not any(value not in (None, "") for value in values):
                continue
            nonblank_count += 1
            if nonblank_count > 1000:
                raise ValueError("A maximum of 1,000 disciplinary cases can be imported at once.")
            row = {canonical: values[index] for index, canonical in enumerate(DISCIPLINARY_IMPORT_COLUMNS)}
            display = {"Row": row_number, **{column: _text(row[column]) for column in DISCIPLINARY_IMPORT_COLUMNS}}
            errors: list[str] = []
            request: DisciplinaryRecordCreateRequest | None = None
            try:
                employee_reference = _text(row["Employee Number"])
                violation_reference = _text(row["Violation Code"])
                employee = employee_lookup.get(employee_reference.casefold())
                if employee is None:
                    raise ValueError("Employee Number does not match an active employee reference in this company.")
                violation = violation_lookup.get(violation_reference.casefold())
                if violation is None:
                    raise ValueError("Violation Code does not match an active master violation reference in this company.")
                incident_date = _parse_date(row["Incident Date"], field_name="Incident Date", required=True)
                description = _text(row["Incident / Case Description"])
                if not description:
                    raise ValueError("Incident / Case Description is required.")
                issuer_reference = _text(row["Issued By"])
                reviewer_reference = _text(row["Reviewed / Approved By"])
                issuer = user_lookup.get(issuer_reference.casefold()) if issuer_reference else None
                reviewer = user_lookup.get(reviewer_reference.casefold()) if reviewer_reference else None
                if issuer_reference and issuer is None:
                    raise ValueError("Issued By does not match an active user reference in this company.")
                if reviewer_reference and reviewer is None:
                    raise ValueError("Reviewed / Approved By does not match an active user reference in this company.")
                status = _text(row["Case Status"]) or "Draft"
                acknowledgment = _text(row["Employee Acknowledgment"]) or "Pending"
                date_issued = _parse_date(row["Date Issued"], field_name="Date Issued", required=False)
                request = DisciplinaryRecordCreateRequest(
                    company_id=company_id,
                    employee_id=employee.id,
                    violation_id=violation.id,
                    incident_date=incident_date,
                    incident_description=description,
                    evidence_remarks=_optional_text(row["Evidence / Remarks"]),
                    context_snapshot=_optional_text(row["Project / Team / Department"]),
                    actual_action_taken=_optional_text(row["Actual Action Taken"]),
                    issued_by_user_id=issuer.id if issuer else None,
                    reviewed_approved_by_user_id=reviewer.id if reviewer else None,
                    date_issued=date_issued,
                    employee_acknowledgment=acknowledgment,
                    case_status=status,
                    notes=_optional_text(row["Notes"]),
                    created_by_user_id=0,  # replaced at import after preview validation
                )
                # Reuse production rule validation without writing a record.
                self.case_service._status(request.case_status)
                self.case_service._acknowledgment(request.employee_acknowledgment)
                self.case_service._validate_issue_fields(
                    status=request.case_status,
                    actual_action_taken=request.actual_action_taken,
                    issued_by_user_id=request.issued_by_user_id,
                    date_issued=request.date_issued,
                )
                signature = (
                    employee.employee_number.casefold(),
                    violation.violation_code.casefold(),
                    incident_date.isoformat(),
                    " ".join(description.split()).casefold(),
                )
                if signature in existing_signatures:
                    errors.append("A matching active disciplinary case already exists in this company.")
                if signature in seen_signatures:
                    errors.append("Duplicate disciplinary case row inside this uploaded workbook.")
                seen_signatures.add(signature)
            except (ValueError, ValidationError) as error:
                errors.append(_validation_message(error) if isinstance(error, ValidationError) else str(error))

            display["Validation"] = " | ".join(errors) if errors else "Ready"
            display["_values"] = request.model_dump() if request is not None and not errors else None
            rows.append(display)

        if not rows:
            raise ValueError("The disciplinary case workbook does not contain any data rows.")
        return rows

    def import_preview_rows(self, rows: list[dict[str, object]], *, company_id: int, actor_user_id: int):
        requests: list[DisciplinaryRecordCreateRequest] = []
        for row in rows:
            if row.get("Validation") != "Ready":
                raise ValueError("Every disciplinary case row must be valid before import.")
            values = row.get("_values")
            if not isinstance(values, dict):
                raise ValueError(f"Row {row.get('Row')} must be validated again before import.")
            payload = dict(values)
            payload["company_id"] = company_id
            payload["created_by_user_id"] = actor_user_id
            requests.append(DisciplinaryRecordCreateRequest(**payload))
        return self.case_service.create_many(requests)
