"""Validated, atomic Excel import for admin planning reminders."""

from __future__ import annotations

from datetime import date, datetime, time, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError
from sqlalchemy.orm import Session

from config.settings import Settings, get_settings
from models.event_reminder import EventReminder
from schemas.event_reminder_schema import (
    EVENT_REMINDER_CATEGORIES,
    EventReminderInput,
)
from services.event_reminder_service import EventReminderService
from services.excel_template_guidance import (
    add_header_guidance,
    add_inline_list_validation,
)


EVENT_REMINDER_IMPORT_COLUMNS = (
    "Category",
    "Event Date",
    "Event / Activity Title",
    "Preparation Notes",
)

EVENT_REMINDER_PREVIEW_COLUMNS = (
    "Row",
    *EVENT_REMINDER_IMPORT_COLUMNS,
    "Validation",
)

SAMPLE_REMINDER_TITLE = "SAMPLE ONLY - Company Anniversary"
SAMPLE_REMINDER_ROW = (
    "Company Event",
    "2026/12/12",
    SAMPLE_REMINDER_TITLE,
    "Prepare the employee announcement and confirm the activity details with HR.",
)


def _text(value: object) -> str:
    """Return one normalized spreadsheet text value."""

    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _validation_message(error: ValidationError) -> str:
    """Return concise field-validation messages for one preview row."""

    messages: list[str] = []
    for item in error.errors():
        message = str(item.get("msg") or "Invalid reminder value.")
        if message.lower().startswith("value error, "):
            message = message[13:]
        messages.append(message)
    return " | ".join(dict.fromkeys(messages))


class EventReminderBulkImportService:
    """Build templates, preview rows, and atomically import reminders."""

    def __init__(
        self,
        session: Session,
        *,
        settings: Settings | None = None,
    ) -> None:
        self.session = session
        self.settings = settings or get_settings()
        self.reminder_service = EventReminderService(
            session,
            settings=self.settings,
        )

    @staticmethod
    def build_template() -> bytes:
        """Return the exact Reminder Excel template with hidden category guidance."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reminder Import Template"
        sheet.append(list(EVENT_REMINDER_IMPORT_COLUMNS))
        sheet.append(list(SAMPLE_REMINDER_ROW))

        header_fill = PatternFill("solid", fgColor="2F740B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        # Keep the complete choice list tied to the same category constant
        # used by the live Create Reminder form and validator.
        category_guidance = (
            "Required. Use one of the current Reminder Category selections.\n\n"
            "Complete valid selections:\n"
            + "\n".join(f"- {value}" for value in EVENT_REMINDER_CATEGORIES)
            + "\n\nChoose from the dropdown or use one of these values exactly."
        )
        sheet["A1"].comment = Comment(category_guidance, "AI HR Assistant")
        sheet["A1"].comment.width = 520
        sheet["A1"].comment.height = min(
            520, max(120, 18 * (category_guidance.count("\n") + 3))
        )
        add_header_guidance(
            sheet,
            "B1",
            requirement="Required event date. Use YYYY/MM/DD or an Excel date cell.",
        )
        add_header_guidance(
            sheet,
            "C1",
            requirement="Required event/activity title; free text.",
        )
        add_header_guidance(
            sheet,
            "D1",
            requirement="Optional preparation notes; free text.",
        )

        sample_fill = PatternFill("solid", fgColor="D9D9D9")
        sample_font = Font(color="666666", italic=True)
        for cell in sheet[2]:
            cell.fill = sample_fill
            cell.font = sample_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 44

        widths = {
            "A": 28,
            "B": 18,
            "C": 42,
            "D": 72,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:D2000"
        add_inline_list_validation(
            sheet,
            "A2:A2000",
            EVENT_REMINDER_CATEGORIES,
            allow_blank=False,
        )

        instructions = workbook.create_sheet("Instructions")
        instruction_rows = (
            ("Column", "Requirement"),
            (
                "Category",
                "Required. Use the Category dropdown or hover/open the header comment to view the complete current selections.",
            ),
            (
                "Event Date",
                "Required. Use YYYY/MM/DD. Excel date cells are also accepted.",
            ),
            (
                "Event / Activity Title",
                "Required. This is the same title entered after YYYY/MM/DD - in the manual Entry Box.",
            ),
            (
                "Preparation Notes",
                "Optional. This is the same preparation-notes content entered below a dated title in the manual Entry Box.",
            ),
            (
                "Schedule",
                "Imported reminders use the same existing behavior as manual reminders: 9:00 AM event time and automatic admin notifications at 1 month, 2 weeks, and 1 week before the event.",
            ),
            (
                "Gray Sample Row",
                "Row 2 is only an example. Delete it or replace it with actual reminder data. An unchanged SAMPLE ONLY row is ignored during upload.",
            ),
        )
        for row in instruction_rows:
            instructions.append(row)
        for cell in instructions[1]:
            cell.fill = header_fill
            cell.font = header_font
        instructions.column_dimensions["A"].width = 34
        instructions.column_dimensions["B"].width = 115
        instructions.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _normalize_category(value: object) -> str:
        """Return one canonical current reminder category."""

        cleaned = _text(value)
        if not cleaned:
            raise ValueError("Category is required.")

        category_map = {
            item.casefold(): item
            for item in EVENT_REMINDER_CATEGORIES
        }
        category = category_map.get(cleaned.casefold())
        if category is None:
            raise ValueError(
                "Category must match one of the current allowed reminder categories."
            )
        return category

    @staticmethod
    def _parse_event_date(value: object) -> date:
        """Accept the template YYYY/MM/DD value or an Excel date cell."""

        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value

        cleaned = _text(value)
        if not cleaned:
            raise ValueError("Event Date is required.")

        for pattern in ("%Y/%m/%d", "%Y-%m-%d"):
            try:
                return datetime.strptime(cleaned, pattern).date()
            except ValueError:
                continue
        raise ValueError("Event Date must use YYYY/MM/DD.")

    def _event_start_at(self, event_date: date) -> datetime:
        """Apply the same 9:00 AM local event time as manual reminder entry."""

        local_value = datetime.combine(
            event_date,
            time(9, 0),
            tzinfo=ZoneInfo(self.settings.display_timezone),
        )
        return local_value.astimezone(timezone.utc)

    def _values_from_row(
        self,
        row: dict[str, object],
        *,
        company_id: int,
    ) -> tuple[EventReminderInput, date]:
        """Convert one spreadsheet row into validated reminder input."""

        category = self._normalize_category(row.get("Category"))
        event_date = self._parse_event_date(row.get("Event Date"))
        title = _text(row.get("Event / Activity Title"))
        notes = _text(row.get("Preparation Notes"))

        values = EventReminderInput(
            company_id=company_id,
            title=title,
            category=category,
            notes=notes,
            event_start_at=self._event_start_at(event_date),
            event_end_at=None,
            status="planned",
            announcement_id=None,
        )
        return values, event_date

    def prepare_preview(
        self,
        file_bytes: bytes,
        *,
        filename: str,
        company_id: int,
    ) -> list[dict[str, object]]:
        """Parse and validate the completed Reminder workbook."""

        try:
            workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        except Exception as exc:
            raise ValueError(
                "The uploaded file is not a readable .xlsx workbook."
            ) from exc

        sheet_name = "Reminder Import Template"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                "Use the downloadable template and keep the "
                "'Reminder Import Template' worksheet name."
            )

        sheet = workbook[sheet_name]
        headers = tuple(_text(cell.value) for cell in sheet[1])
        if headers != EVENT_REMINDER_IMPORT_COLUMNS:
            raise ValueError(
                "The Excel columns or their order do not match the downloadable "
                "Reminder template."
            )

        raw_rows: list[dict[str, object]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value not in (None, "") for value in values):
                continue
            if _text(values[2]).casefold() == SAMPLE_REMINDER_TITLE.casefold():
                continue
            raw_rows.append(
                {
                    "Row": row_number,
                    **dict(zip(EVENT_REMINDER_IMPORT_COLUMNS, values)),
                }
            )

        if not raw_rows:
            raise ValueError(
                "The Reminder template does not contain any reminder rows."
            )

        preview: list[dict[str, object]] = []
        for row in raw_rows:
            errors: list[str] = []
            reminder_values: EventReminderInput | None = None
            event_date: date | None = None
            try:
                reminder_values, event_date = self._values_from_row(
                    row,
                    company_id=company_id,
                )
            except ValidationError as error:
                errors.append(_validation_message(error))
            except ValueError as error:
                errors.append(str(error))

            category = (
                reminder_values.category
                if reminder_values is not None
                else _text(row.get("Category"))
            )
            title = (
                reminder_values.title
                if reminder_values is not None
                else _text(row.get("Event / Activity Title"))
            )
            notes = (
                reminder_values.notes
                if reminder_values is not None
                else _text(row.get("Preparation Notes"))
            )
            event_date_text = (
                event_date.strftime("%Y/%m/%d")
                if event_date is not None
                else _text(row.get("Event Date"))
            )

            preview.append(
                {
                    "Row": int(row["Row"]),
                    "Category": category,
                    "Event Date": event_date_text,
                    "Event / Activity Title": title,
                    "Preparation Notes": notes,
                    "Validation": (
                        "Ready"
                        if not errors
                        else " | ".join(dict.fromkeys(errors))
                    ),
                    "_filename": filename,
                    "_values": reminder_values.model_dump()
                    if reminder_values is not None
                    else None,
                }
            )
        return preview

    def import_preview_rows(
        self,
        preview_rows: list[dict[str, object]],
        *,
        company_id: int,
        actor_user_id: int,
    ) -> list[EventReminder]:
        """Revalidate and create the complete valid batch atomically."""

        if not preview_rows:
            raise ValueError("There are no reminder rows ready to import.")
        if any(str(row.get("Validation", "")) != "Ready" for row in preview_rows):
            raise ValueError(
                "Resolve every validation error before importing reminders."
            )

        values_list: list[EventReminderInput] = []
        for row in preview_rows:
            try:
                values, _ = self._values_from_row(
                    row,
                    company_id=company_id,
                )
            except (ValidationError, ValueError) as error:
                if isinstance(error, ValidationError):
                    message = _validation_message(error)
                else:
                    message = str(error)
                raise ValueError(f"Row {row.get('Row')}: {message}") from error
            values_list.append(values)

        return self.reminder_service.create_many(
            values_list,
            actor_user_id=actor_user_id,
        )
