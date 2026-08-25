"""Validated, atomic Excel import for the company HR contact directory."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from models.hr_contact import HRContact
from services.hr_contact_service import HRContactService
from services.excel_template_guidance import add_header_guidance


HR_CONTACT_IMPORT_COLUMNS = (
    "Contact Name",
    "Position / Role",
    "Department / Team",
    "Email",
    "Phone / Mobile",
    "Office / Location",
    "Availability / Office Hours",
    "Notes / Support Coverage",
)

HR_CONTACT_PREVIEW_COLUMNS = (
    "Row",
    *HR_CONTACT_IMPORT_COLUMNS,
    "Validation",
)

SAMPLE_CONTACT_NAME = "SAMPLE ONLY - Maria Santos"
SAMPLE_CONTACT_ROW = (
    SAMPLE_CONTACT_NAME,
    "HR Manager",
    "Human Resources",
    "maria.santos@company.com",
    "+63 912 345 6789",
    "3rd Floor, Main Office, Manila",
    "Monday-Friday, 9:00 AM - 5:00 PM",
    "Handles employee relations, recruitment, and benefits inquiries",
)


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


class HRContactBulkImportService:
    """Prepare HR-contact Excel previews and atomically import valid rows."""

    def __init__(self, session: Session) -> None:
        self.session = session
        self.contact_service = HRContactService(session)

    @staticmethod
    def build_template() -> bytes:
        """Return a styled HR Contact Excel template with guidance."""

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "HR Contact Import Template"
        sheet.append(list(HR_CONTACT_IMPORT_COLUMNS))
        sheet.append(list(SAMPLE_CONTACT_ROW))

        header_fill = PatternFill("solid", fgColor="2F740B")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sample_fill = PatternFill("solid", fgColor="D9D9D9")
        sample_font = Font(color="666666", italic=True)
        for cell in sheet[2]:
            cell.fill = sample_fill
            cell.font = sample_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 44

        widths = {
            "A": 28,
            "B": 24,
            "C": 26,
            "D": 34,
            "E": 24,
            "F": 34,
            "G": 38,
            "H": 52,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:H2000"

        # HR Contact fields are intentionally free text rather than closed
        # project selections.  Header comments make that explicit so users do
        # not waste time looking for a hidden list of allowed values.
        header_guidance = {
            "A1": "Required contact name; free text.",
            "B1": "Optional position, function, or HR responsibility; free text.",
            "C1": "Optional department or team name; free text.",
            "D1": "Optional email address. At least one Email or Phone / Mobile value is required.",
            "E1": "Optional phone/mobile value. At least one Email or Phone / Mobile value is required.",
            "F1": "Optional office/location information; free text.",
            "G1": "Optional availability or office-hours information; free text.",
            "H1": "Optional notes/support coverage such as payroll, benefits, leave, or employee relations; free text.",
        }
        for cell_reference, requirement in header_guidance.items():
            add_header_guidance(
                sheet,
                cell_reference,
                requirement=requirement,
                extra="This field has no fixed project selection list.",
            )

        instructions = workbook.create_sheet("Instructions")
        instruction_rows = (
            ("Column", "Requirement"),
            ("Contact Name", "Required."),
            (
                "Email / Phone",
                "At least one Email or Phone / Mobile value is required for every contact.",
            ),
            (
                "Position / Role",
                "Optional position, function, or HR responsibility.",
            ),
            ("Department / Team", "Optional team or department name."),
            ("Office / Location", "Optional office/location information."),
            (
                "Availability / Office Hours",
                "Optional availability or office-hours text.",
            ),
            (
                "Notes / Support Coverage",
                "Optional support details such as payroll, benefits, leave, or employee relations.",
            ),
            (
                "Display Order",
                "Automatic. Do not add a Display Order column to the template.",
            ),
            (
                "Gray Sample Row",
                "Row 2 is only an example. Delete it or replace it with actual HR contact data. An unchanged SAMPLE ONLY row is ignored during upload.",
            ),
        )
        for row in instruction_rows:
            instructions.append(row)
        for cell in instructions[1]:
            cell.fill = header_fill
            cell.font = header_font
        instructions.column_dimensions["A"].width = 34
        instructions.column_dimensions["B"].width = 110
        instructions.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _values_from_row(row: dict[str, object]) -> dict[str, object]:
        return {
            "name": _text(row.get("Contact Name")),
            "job_title": _text(row.get("Position / Role")),
            "team": _text(row.get("Department / Team")),
            "email": _text(row.get("Email")),
            "phone": _text(row.get("Phone / Mobile")),
            "office_location": _text(row.get("Office / Location")),
            "availability": _text(row.get("Availability / Office Hours")),
            "notes": _text(row.get("Notes / Support Coverage")),
        }

    def prepare_preview(
        self,
        file_bytes: bytes,
        *,
        filename: str,
        company_id: int,
    ) -> list[dict[str, object]]:
        """Parse and validate one uploaded HR-contact workbook."""

        del company_id  # Validation is company-safe at import; field rules are universal.

        try:
            workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        except Exception as exc:
            raise ValueError("The uploaded file is not a readable .xlsx workbook.") from exc

        sheet_name = "HR Contact Import Template"
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                "Use the downloadable template and keep the "
                "'HR Contact Import Template' worksheet name."
            )
        sheet = workbook[sheet_name]
        headers = tuple(_text(cell.value) for cell in sheet[1])
        if headers != HR_CONTACT_IMPORT_COLUMNS:
            raise ValueError(
                "The Excel columns or their order do not match the downloadable "
                "HR Contact template."
            )

        raw_rows: list[dict[str, object]] = []
        for row_number, values in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value not in (None, "") for value in values):
                continue
            if _text(values[0]).casefold() == SAMPLE_CONTACT_NAME.casefold():
                continue
            raw_rows.append(
                {
                    "Row": row_number,
                    **dict(zip(HR_CONTACT_IMPORT_COLUMNS, values)),
                }
            )

        if not raw_rows:
            raise ValueError("The HR Contact template does not contain any contact rows.")

        preview: list[dict[str, object]] = []
        for row in raw_rows:
            errors: list[str] = []
            values = self._values_from_row(row)
            try:
                normalized = self.contact_service._normalize_values(values)
            except ValueError as exc:
                normalized = values
                errors.append(str(exc))

            preview.append(
                {
                    "Row": int(row["Row"]),
                    "Contact Name": normalized.get("name", ""),
                    "Position / Role": normalized.get("job_title", ""),
                    "Department / Team": normalized.get("team", ""),
                    "Email": normalized.get("email", ""),
                    "Phone / Mobile": normalized.get("phone", ""),
                    "Office / Location": normalized.get("office_location", ""),
                    "Availability / Office Hours": normalized.get("availability", ""),
                    "Notes / Support Coverage": normalized.get("notes", ""),
                    "Validation": "Ready" if not errors else " | ".join(dict.fromkeys(errors)),
                    "_filename": filename,
                    "_values": normalized,
                }
            )
        return preview

    def import_preview_rows(
        self,
        preview_rows: list[dict[str, object]],
        *,
        company_id: int,
    ) -> list[HRContact]:
        """Revalidate and atomically create a whole valid HR-contact batch."""

        if not preview_rows:
            raise ValueError("There are no HR contact rows ready to import.")
        if any(str(row.get("Validation", "")) != "Ready" for row in preview_rows):
            raise ValueError("Resolve every validation error before importing HR contacts.")

        normalized_rows: list[dict[str, object]] = []
        for row in preview_rows:
            row_values = dict(row.get("_values") or self._values_from_row(row))
            try:
                normalized_rows.append(self.contact_service._normalize_values(row_values))
            except ValueError as exc:
                raise ValueError(f"Row {row.get('Row')}: {exc}") from exc

        next_order = self.contact_service.repository.next_sort_order(company_id)
        contacts: list[HRContact] = []
        try:
            for offset, values in enumerate(normalized_rows):
                contact = HRContact(
                    company_id=company_id,
                    sort_order=next_order + (offset * 10),
                    is_active=True,
                    **values,
                )
                self.session.add(contact)
                contacts.append(contact)
            self.session.commit()
            for contact in contacts:
                self.session.refresh(contact)
        except Exception:
            self.session.rollback()
            raise
        return contacts
